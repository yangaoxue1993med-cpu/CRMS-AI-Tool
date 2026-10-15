"""End-to-end bilingual export and data-integrity regression checks."""

import copy
from io import BytesIO, StringIO
from pathlib import Path
import re
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import CRMS_AI_Tool as app

PATH = str(Path(__file__).resolve().parents[1] / "CRMS_AI_Tool.py")


def populated_project(template="通用"):
    p = app.make_project(template)
    p["settings"]["project_name"] = "供应商"  # A title must not become a category.
    s = app.new_site(p, 1, name="研究中心", code="CN01")
    s.update(
        pi="已完成", contact="王医生 / pi@example.test", planned_n=80, is_lead=True
    )
    s["readiness"]["F01"] = "已完成"
    s["readiness"]["F02"] = "不适用"
    p["sites"] = [s]
    p = app.upsert_enrollment(p, s["id"], "2026-01-31", 12, "需再次核实")
    p["budget"][0]["note"] = "已完成"
    p["timeline"][0]["owner"] = "研究中心"
    p["expenses"] = [
        dict(
            id="E-test",
            date="2026-01-31",
            budget_id=p["budget"][0]["id"],
            amount=123.45,
            currency="USD",
            note="=1+1",
        )
    ]
    return p


@pytest.mark.parametrize("template", ["示例项目", "通用"])
def test_english_exports_do_not_change_data_or_calculations(template):
    p = populated_project(template)
    tables = app.report_tables(p)
    before_p, before_tables = copy.deepcopy(p), copy.deepcopy(tables)
    json_before = app.project_bytes(p)
    workbook = load_workbook(BytesIO(app.excel_bytes(tables, "en")))
    assert len(workbook.sheetnames) == 13
    assert "Project Summary" in workbook.sheetnames
    assert "Monthly Enrollment - All Sites" in workbook.sheetnames
    assert all(
        len(n) <= 31 and not re.search(r"[\u4e00-\u9fff]", n)
        for n in workbook.sheetnames
    )
    for table, rows in tables.items():
        translated = app.translate_export_rows(rows, "en", table)
        # Every non-string datum, including None (missing actuals), survives exactly.
        for original, english in zip(rows, translated):
            assert len(original) == len(english)
            for old, new in zip(original.values(), english.values()):
                if not isinstance(old, str):
                    assert type(old) is type(new) and old == new
        sheet = workbook[app.export_table_name(table, "en")]
        assert all(not re.search(r"[\u4e00-\u9fff]", c.value or "") for c in sheet[1])
    summary = dict(workbook["Project Summary"].iter_rows(min_row=2, values_only=True))
    assert summary["Project"] == "供应商"
    assert summary["Leading Site"] == "研究中心"
    assert summary["Leading PI"] == "已完成"
    assert summary["Total Budget (RMB)"] == next(
        r["内容"] for r in tables["项目摘要"] if r["指标"] == "总预算RMB"
    )
    assert p == before_p and tables == before_tables
    assert app.project_bytes(p) == json_before


def test_translation_covers_presets_and_preserves_user_text():
    p = populated_project()
    tables = app.report_tables(p)
    budget = app.export_frame(tables["预算明细"], "en", "预算明细")
    assert budget.iloc[0]["Item"] == "Investigator Screening Fees"
    assert budget.iloc[0]["Notes"] == "已完成"
    assert budget.iloc[0]["Quantity Basis"] == "Enrolled Participants"
    checklist = app.export_frame(tables["FPI启动检查"], "en", "FPI启动检查")
    assert checklist.iloc[0]["Status"] == "Completed"
    assert checklist.iloc[1]["Status"] == "Not Applicable"
    assert checklist.iloc[0]["Site Name"] == "研究中心"
    assert (
        checklist.iloc[0]["Checklist Item"]
        == "Study Type and Regulatory Pathway Confirmed"
    )
    timeline = app.export_frame(tables["时间线"], "en", "时间线")
    assert timeline.iloc[0]["Owner"] == "研究中心"
    expenses = app.export_frame(tables["实际支出"], "en", "实际支出")
    assert expenses.iloc[0]["Notes"] == "'=1+1"
    snapshots = app.export_frame(tables["入组累计快照"], "en", "入组累计快照")
    assert snapshots.iloc[0]["Notes"] == "需再次核实"


@pytest.mark.parametrize("template", ["示例项目", "通用"])
def test_standard_templates_have_english_system_content(template):
    p = app.make_project(template)
    for name, rows in app.report_tables(p).items():
        translated = app.translate_export_rows(rows, "en", name)
        assert not any(
            re.search(r"[\u4e00-\u9fff]", str(v))
            for row in translated
            for v in row.values()
        )
    s = app.new_site(p, 1, name="Example Hospital")
    p["sites"] = [s]
    readiness = app.translate_export_rows(
        app.site_report_tables(p)["FPI启动检查"], "en", "FPI启动检查"
    )
    assert len(readiness) == 16
    assert not any(
        re.search(r"[\u4e00-\u9fff]", str(v)) for row in readiness for v in row.values()
    )


def test_english_csv_missing_values_and_empty_schemas():
    rows = [
        {
            "月份": "2026-01",
            "统计截至": "2026-01-31",
            "计划累计": 22.5,
            "实际累计": None,
            "纳入中心数": 2,
            "已有累计报告中心数": 1,
            "本月更新中心数": 1,
        }
    ]
    data = app.csv_bytes(rows, "en", "全中心月度入组")
    assert data.startswith(b"\xef\xbb\xbf")
    frame = pd.read_csv(StringIO(data.decode("utf-8-sig")))
    assert frame.loc[0, "Planned Cumulative Enrollment"] == 22.5
    assert pd.isna(frame.loc[0, "Actual Cumulative Enrollment"])
    assert frame.loc[0, "As of Date"] == "2026-01-31"
    empty = pd.read_csv(BytesIO(app.csv_bytes([], "en", "实际支出")))
    assert list(empty) == [
        "ID",
        "Date",
        "Budget Item ID",
        "Amount",
        "Currency",
        "Notes",
    ]
    assert empty.empty
    assert "月份" in app.csv_bytes(rows).decode("utf-8-sig")
    with pytest.raises(ValueError):
        app.csv_bytes([], "invalid", "实际支出")


def test_custom_names_in_english_generated_warnings_are_preserved():
    text = "自定义检测：无预设匹配规则，已按全周期摊分；可指定阶段。"
    result = app.english_system_message(text)
    assert result.startswith("自定义检测: no predefined allocation rule")
    assert (
        app.english_system_message("实际支出已超当前预算RMB 1,234.56。")
        == "Actual expenditure exceeds the current budget by RMB 1,234.56."
    )


def test_language_switch_invalidates_excel_cache_and_keeps_page_and_data():
    at = AppTest.from_file(PATH, default_timeout=40).run()
    at.radio(key="navigation").set_value("导出与帮助").run()
    before = copy.deepcopy(at.session_state.project)
    prepare = lambda: next(b for b in at.button if b.label == "准备Excel工作簿")
    prepare().click().run()
    zh = at.session_state.excel_export[1]
    assert "项目摘要" in load_workbook(BytesIO(zh)).sheetnames
    at.radio(key="export_language").set_value("English").run()
    assert not at.exception
    assert at.radio(key="navigation").value == "导出与帮助"
    assert at.session_state.revision == 0 and at.session_state.project == before
    # The stale Chinese download is hidden until a new English workbook is prepared.
    downloads = [x for x in at.get("download_button") if "Excel" in x.label]
    assert not downloads
    prepare().click().run()
    assert not at.exception
    assert at.session_state.excel_export[2] == "en"
    en = at.session_state.excel_export[1]
    assert "Project Summary" in load_workbook(BytesIO(en)).sheetnames
    assert at.session_state.project == before
    assert "Metric" in at.dataframe[0].value.columns
    at.radio(key="navigation").set_value("中心启动进度").run()
    assert not at.exception
    assert at.session_state.table_export_language == "English"
    at.radio(key="navigation").set_value("导出与帮助").run()
    assert at.radio(key="export_language").value == "English"
    at.radio(key="export_language").set_value("中文").run()
    assert not at.exception
    assert at.radio(key="navigation").value == "导出与帮助"
    assert at.session_state.table_export_language == "中文"
    assert at.session_state.project == before
