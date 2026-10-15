"""CRMS Clinical Study Planner 2.4 — single-file Streamlit application.

Run: streamlit run CRMS_AI_Tool.py
The calculation engine is importable without starting Streamlit.
All template figures come from the supplied V1.15 file, not new source verification.
See README.md for definitions, limitations and deployment instructions.
"""

from __future__ import annotations

import calendar
import copy
import datetime as dt
from decimal import Decimal, ROUND_HALF_UP, ROUND_FLOOR
import html
import inspect
from io import BytesIO
import json
import math
import re
from typing import Any
import uuid

from dateutil.relativedelta import relativedelta

APP_VERSION = "2.4.1"
SCHEMA_VERSION = 2
MODULES = ["研究中心", "供应商", "产品物流", "人力资源"]
BASIS = [
    "固定数量",
    "入组例数",
    "完成随访例数",
    "筛选失败例数",
    "筛选总例数",
    "中心数",
    "项目月数",
    "FTE项目",
    "FTE月",
    "FTE年",
]
CASH_MODES = ["智能匹配时间线", "示例项目预填季度", "自定义季度占比", "全周期按月均摊"]
ALLOC_MODES = ["智能规则", "阶段均摊", "阶段开始一次性", "阶段结束一次性"]
COLORS = [
    "#2357D9",
    "#17A398",
    "#7894D1",
    "#AF96D9",
    "#E8AF58",
    "#536B89",
    "#80B9AA",
    "#BF8EAC",
    "#809BA6",
    "#9FACC5",
]

EXAMPLE_PROJECT_DEFAULTS = {
    "project_name": "示例项目（预填演示）",
    "study_type": "上市前临床研究",
    "site_number": 20,  # Synopsis: up to 20 sites; site forecast also notes 12 nominated sites.
    "n_planned": 142,
    "n_lost": 0,
    "n_screen_fail": 0,
    "start_date": dt.date(2026, 3, 11),  # CIP approved
    "project_planned_duration_months": 64,  # FY26Q4 prep start (2026-02) to Study Complete (2031-05); aligns with current Timeline span
    "primary_endpoint_time_months": 30,  # CIP approved to primary endpoint database snapshot around 2028-09
}

EXAMPLE_TIMELINE_DEFAULTS = {
    # EXAMPLE Timeline is aligned to the supplied Medtronic fiscal-year slide:
    # FY26Q4 = Feb-May 2026; FY27Q1 = Jun-Aug 2026; FY27Q2 = Sep-Nov 2026.
    # Key milestones: CIP approved 11-Mar-2026, HGRAC/regulatory approval Sep-2026,
    # first site/first subject Oct-2026, last site Jan-2027, last subject Jul-2027,
    # primary endpoint database snapshot Sep-2028, 1Y CSR Jan-2029,
    # final DBL Sep-2030, final CSR Feb-2031, study complete May-2031.
    # 研究准备阶段：对应Timeline-Stretch中的“Study documents prep & PI, Site list confirmation”
    "项目立项": (True, dt.date(2026, 2, 1), dt.date(2026, 3, 11)),
    "研究中心名单提名、筛选及访视": (True, dt.date(2026, 2, 1), dt.date(2026, 5, 31)),
    "产品及人员培训": (True, dt.date(2026, 6, 1), dt.date(2026, 9, 30)),
    "供应商筛选与确认": (True, dt.date(2026, 2, 1), dt.date(2026, 5, 31)),
    "产品运输及相关管理流程准备": (True, dt.date(2026, 4, 1), dt.date(2026, 9, 30)),
    "研究者会议": (True, dt.date(2026, 3, 1), dt.date(2026, 3, 11)),
    # 研究文件准备 / CIP readiness：CIP approved = 11-Mar-2026；ICF/CRF/内部释放节点来自CIP readiness表
    "研究方案撰写及获批": (True, dt.date(2026, 2, 1), dt.date(2026, 3, 11)),
    "CRF设计及数据库建立": (True, dt.date(2026, 3, 12), dt.date(2026, 3, 25)),
    "知情同意书撰写及获批": (True, dt.date(2026, 3, 12), dt.date(2026, 3, 20)),
    "研究计划书撰写及获批（监查计划、数据管理计划、统计分析计划等）": (
        True,
        dt.date(2026, 3, 12),
        dt.date(2026, 4, 15),
    ),
    "其他相关表格及物料准备（受试者日记卡、招募海报、标签等）": (
        True,
        dt.date(2026, 3, 20),
        dt.date(2026, 5, 31),
    ),
    "文件翻译及定稿": (True, dt.date(2026, 3, 12), dt.date(2026, 4, 30)),
    # 政府部门备案：以HGRAC/监管获批启动研究为主，Sep-2026作为监管启动里程碑
    "上海药监局备案": (False, dt.date(2026, 4, 15), dt.date(2026, 9, 30)),
    "人类遗传办提交": (True, dt.date(2026, 4, 15), dt.date(2026, 5, 31)),
    "人类遗传办获批": (True, dt.date(2026, 9, 1), dt.date(2026, 9, 30)),
    # GCP/EC/合同/中心启动：Site Submission约15-Apr-2026，首中心启动Oct-2026，末中心启动Jan-2027
    "组长单位GCP资料递交": (True, dt.date(2026, 4, 15), dt.date(2026, 4, 30)),
    "组长单位EC申请及获批": (True, dt.date(2026, 4, 15), dt.date(2026, 9, 30)),
    "其他中心EC申请及获批": (True, dt.date(2026, 9, 1), dt.date(2027, 1, 31)),
    "组长单位合同起草与协商": (True, dt.date(2026, 4, 15), dt.date(2026, 9, 30)),
    "组长单位合同签字与盖章": (True, dt.date(2026, 9, 1), dt.date(2026, 9, 30)),
    "组长单位启动": (True, dt.date(2026, 10, 1), dt.date(2026, 10, 31)),
    "其他中心启动": (True, dt.date(2026, 10, 1), dt.date(2027, 1, 31)),
    # 入组/治疗：参考forecast表，入组约10–11个月；首例Oct-2026，末例Jul-2027
    "组长单位首例入组": (True, dt.date(2026, 10, 1), dt.date(2026, 10, 31)),
    "其他中心首例入组": (True, dt.date(2026, 10, 1), dt.date(2027, 1, 31)),
    "受试者入组": (True, dt.date(2026, 10, 1), dt.date(2027, 7, 31)),
    "末例入组": (True, dt.date(2027, 7, 1), dt.date(2027, 7, 31)),
    # 随访：随访窗口包含出院、30天、6个月、1年、2年、3年；主要终点DB snapshot Sep-2028
    "主要终点随访": (True, dt.date(2026, 10, 1), dt.date(2028, 9, 30)),
    "随访结束": (True, dt.date(2026, 10, 1), dt.date(2030, 9, 30)),
    "末例受试者末次访视": (True, dt.date(2030, 9, 1), dt.date(2030, 9, 30)),
    # 主要终点分析与注册递交：1Y CSR Jan-2029，商业注册递交Mar-2029
    "主要终点数据库锁定": (True, dt.date(2028, 9, 1), dt.date(2028, 9, 30)),
    "主要终点TLG/统计报告撰写": (True, dt.date(2028, 9, 1), dt.date(2029, 1, 31)),
    "主要终点临床研究报告撰写": (True, dt.date(2028, 9, 1), dt.date(2029, 1, 31)),
    "主要终点数据的监管提交": (True, dt.date(2029, 3, 1), dt.date(2029, 3, 31)),
    # 最终分析：Final DBL Sep-2030，Final CSR approved Feb-2031
    "数据库清理与锁定": (True, dt.date(2030, 9, 1), dt.date(2030, 9, 30)),
    "最终图表TLG/统计报告完成": (True, dt.date(2030, 9, 1), dt.date(2031, 2, 28)),
    "临床研究报告撰写": (True, dt.date(2030, 9, 1), dt.date(2031, 2, 28)),
    # 研究收尾：Study Complete May-2031
    "中心关闭访视": (True, dt.date(2031, 2, 1), dt.date(2031, 5, 31)),
    "临床研究报告的监管递交": (True, dt.date(2031, 2, 1), dt.date(2031, 5, 31)),
}

EXAMPLE_BUDGET_QUARTERS = [
    f"FY{fy}Q{q}" for fy in [27, 28, 29, 30, 31] for q in [1, 2, 3, 4]
]

EXAMPLE_BUDGET_MATRIX_USD = {
    "Site Cost": [
        0,
        10990,
        141470,
        72828,
        87192,
        28728,
        34776,
        21168,
        21168,
        756,
        13608,
        21168,
        21168,
        1036,
        16688,
        21168,
        21168,
        280,
        3080,
        0,
    ],
    "Product Costs": [
        0,
        4402,
        79229,
        123245,
        123245,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
    ],
    "CEC Services (Baim)-USA": [
        10141,
        11064,
        12847,
        12847,
        11517,
        8859,
        8859,
        8859,
        8859,
        8859,
        8859,
        8859,
        8859,
        8859,
        8859,
        8859,
        8859,
        3857,
        1286,
        0,
    ],
    "CTA Core Laboratory (NAMSA)-USA": [
        0,
        18957,
        56870,
        56870,
        56870,
        56870,
        56870,
        56870,
        37913,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
    ],
    "Image Transfer (Medidata Image)-USA": [
        0,
        1778,
        5333,
        5333,
        5333,
        5333,
        5333,
        5333,
        5333,
        5333,
        5333,
        5333,
        5333,
        5333,
        5333,
        5333,
        3556,
        0,
        0,
        0,
    ],
    "SMO-China": [
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        14212,
        0,
        0,
        0,
    ],
    "DM": [
        39747,
        45651,
        33809,
        21885,
        14897,
        14818,
        10696,
        7270,
        7750,
        7910,
        7722,
        7488,
        7983,
        8147,
        7954,
        7713,
        8222,
        17558,
        29176,
        6170,
    ],
    "Monitoring": [
        11200,
        16800,
        16800,
        16800,
        16800,
        16800,
        16800,
        16800,
        16800,
        14000,
        8400,
        8400,
        8400,
        8400,
        8400,
        8400,
        8400,
        11200,
        16800,
        5600,
    ],
    "Safety": [
        1426,
        2577,
        6832,
        10895,
        12990,
        14279,
        14040,
        13741,
        14411,
        14708,
        14462,
        14154,
        14844,
        15149,
        14896,
        14578,
        15289,
        14728,
        8832,
        0,
    ],
    "Other Costs*": [
        7701,
        20833,
        42430,
        33097,
        33097,
        8797.02,
        8797.02,
        8797.02,
        8797.02,
        8797.02,
        8797.02,
        8797.02,
        8797.02,
        8797.02,
        8797.02,
        8797.02,
        8797.02,
        5898.76,
        1966.25,
        0,
    ],
    "MDT Resources": [
        165880,
        160671,
        156421,
        156421,
        156421,
        156421,
        147921,
        147921,
        147921,
        147921,
        153130,
        153130,
        153130,
        153130,
        153130,
        137503,
        107094,
        116454,
        107094,
        107094,
    ],
    "MDT Travel": [
        1867,
        3733,
        5600,
        5600,
        5600,
        1867,
        0,
        0,
        3733,
        5600,
        0,
        0,
        3733,
        5600,
        3733,
        0,
        3733,
        5600,
        1867,
        0,
    ],
}

EXAMPLE_BUDGET_EXCLUDED_USD = {
    # 截图保留此行用于参考，但未纳入“In total”与本工具默认总预算。
    "MC2 (reference only; excluded from total)": [
        52373,
        65028,
        57441,
        49580,
        44687,
        45898,
        41537,
        37812,
        38961,
        36617,
        30584,
        30042,
        31226,
        31696,
        31250,
        30692,
        31911,
        43486,
        54807,
        11770,
    ],
}

EXAMPLE_BUDGET_MODULE_MAP = {
    "Site Cost": "预算模块1：研究中心花费",
    "Product Costs": "预算模块3：产品物流成本",
    "MDT Resources": "预算模块4：人力资源成本",
    "MDT Travel": "预算模块4：人力资源成本",
}

BUDGET_ALLOCATION_RULES = {
    "研究者筛选费": {
        "phases": [
            {
                "level1": "受试者入组",
                "level2": "受试者入组",
                "weight": 1.0,
                "mode": "spread",
            }
        ]
    },
    "研究者手术费": {
        "phases": [
            {
                "level1": "受试者入组",
                "level2": "受试者入组",
                "weight": 1.0,
                "mode": "spread",
            }
        ]
    },
    "筛选失败费用": {
        "phases": [
            {
                "level1": "受试者入组",
                "level2": "受试者入组",
                "weight": 1.0,
                "mode": "spread",
            }
        ]
    },
    "受试者补偿": {
        "phases": [
            {"level1": "受试者入组", "weight": 0.4},
            {"level1": "受试者随访", "weight": 0.6},
        ]
    },
    "受试者检查费用": {
        "phases": [
            {"level1": "受试者入组", "weight": 0.35},
            {"level1": "受试者随访", "weight": 0.65},
        ]
    },
    "组长单位费用": {
        "phases": [
            {
                "level1": "研究中心启动",
                "level2": "组长单位启动",
                "weight": 1.0,
                "mode": "lump_start",
            }
        ]
    },
    "EC/Clinical Institute Process费用": {
        "phases": [
            {"level1": "组长单位GCP&伦理递交", "weight": 0.8},
            {"level1": "研究中心启动", "weight": 0.2},
        ]
    },
    "中心管理费用": {
        "phases": [
            {"level1": "受试者入组", "weight": 0.45},
            {"level1": "受试者随访", "weight": 0.35},
            {"level1": "数据分析", "weight": 0.20},
        ]
    },
    "Site税费": {
        "phases": [
            {"level1": "受试者入组", "weight": 0.45},
            {"level1": "受试者随访", "weight": 0.35},
            {"level1": "数据分析", "weight": 0.20},
        ]
    },
    "CRC费用": {
        "phases": [
            {"level1": "研究中心启动", "weight": 0.10},
            {"level1": "受试者入组", "weight": 0.50},
            {"level1": "受试者随访", "weight": 0.30},
            {"level1": "数据分析", "weight": 0.10},
        ]
    },
    "PM费用": {
        "phases": [
            {"level1": "研究准备阶段", "weight": 0.10},
            {"level1": "研究文件准备", "weight": 0.15},
            {"level1": "组长单位GCP&伦理递交", "weight": 0.10},
            {"level1": "研究中心启动", "weight": 0.10},
            {"level1": "受试者入组", "weight": 0.25},
            {"level1": "受试者随访", "weight": 0.15},
            {"level1": "数据分析", "weight": 0.10},
            {"level1": "研究结束", "weight": 0.05},
        ]
    },
    "Monitor费用": {
        "phases": [
            {"level1": "研究中心启动", "weight": 0.20},
            {"level1": "受试者入组", "weight": 0.45},
            {"level1": "受试者随访", "weight": 0.25},
            {"level1": "研究结束", "weight": 0.10},
        ]
    },
    "DM费用": {
        "phases": [
            {"level1": "研究文件准备", "weight": 0.25},
            {"level1": "受试者入组", "weight": 0.15},
            {"level1": "受试者随访", "weight": 0.20},
            {"level1": "数据分析", "weight": 0.40},
        ]
    },
    "Safety费用": {
        "phases": [
            {"level1": "受试者入组", "weight": 0.30},
            {"level1": "受试者随访", "weight": 0.45},
            {"level1": "数据分析", "weight": 0.25},
        ]
    },
    "EDC系统费用": {
        "phases": [
            {"level1": "研究文件准备", "weight": 0.35},
            {"level1": "受试者入组", "weight": 0.20},
            {"level1": "受试者随访", "weight": 0.20},
            {"level1": "数据分析", "weight": 0.25},
        ]
    },
    "翻译/打印费用": {
        "phases": [
            {"level1": "研究文件准备", "weight": 0.75},
            {"level1": "研究结束", "weight": 0.25},
        ]
    },
    "保险费用": {
        "phases": [{"level1": "研究中心启动", "weight": 1.0, "mode": "lump_start"}]
    },
    "数据分析费用": {"phases": [{"level1": "数据分析", "weight": 1.0}]},
    "中心实验室费用": {
        "phases": [
            {"level1": "受试者入组", "weight": 0.25},
            {"level1": "受试者随访", "weight": 0.75},
        ]
    },
    "CEC费用": {"phases": [{"level1": "受试者随访", "weight": 1.0}]},
    "Travel费用": {
        "phases": [
            {"level1": "研究中心启动", "weight": 0.30},
            {"level1": "受试者入组", "weight": 0.40},
            {"level1": "受试者随访", "weight": 0.30},
        ]
    },
    "Recording费用": {
        "phases": [
            {"level1": "受试者入组", "weight": 0.50},
            {"level1": "受试者随访", "weight": 0.50},
        ]
    },
    "研究者会费用": {
        "phases": [
            {
                "level1": "研究准备阶段",
                "level2": "研究者会议",
                "weight": 1.0,
                "mode": "lump_start",
            }
        ]
    },
    "Vendor税费": {
        "phases": [
            {"level1": "研究准备阶段", "weight": 0.10},
            {"level1": "研究文件准备", "weight": 0.10},
            {"level1": "研究中心启动", "weight": 0.10},
            {"level1": "受试者入组", "weight": 0.30},
            {"level1": "受试者随访", "weight": 0.20},
            {"level1": "数据分析", "weight": 0.15},
            {"level1": "研究结束", "weight": 0.05},
        ]
    },
    "产品运输费用": {
        "phases": [
            {
                "level1": "研究准备阶段",
                "level2": "产品运输及相关管理流程准备",
                "weight": 0.30,
            },
            {"level1": "受试者入组", "weight": 0.70},
        ]
    },
    "仓储费用": {
        "phases": [
            {"level1": "受试者入组", "weight": 0.45},
            {"level1": "受试者随访", "weight": 0.55},
        ]
    },
    "产品成本": {"phases": [{"level1": "受试者入组", "weight": 1.0}]},
    "项目经理FTE费用": {
        "phases": [
            {"level1": "研究准备阶段", "weight": 0.12},
            {"level1": "研究文件准备", "weight": 0.12},
            {"level1": "组长单位GCP&伦理递交", "weight": 0.08},
            {"level1": "研究中心启动", "weight": 0.10},
            {"level1": "受试者入组", "weight": 0.25},
            {"level1": "受试者随访", "weight": 0.18},
            {"level1": "数据分析", "weight": 0.10},
            {"level1": "研究结束", "weight": 0.05},
        ]
    },
    "CRA费用": {
        "phases": [
            {"level1": "研究中心启动", "weight": 0.20},
            {"level1": "受试者入组", "weight": 0.40},
            {"level1": "受试者随访", "weight": 0.30},
            {"level1": "研究结束", "weight": 0.10},
        ]
    },
    "MCRS费用": {
        "phases": [
            {"level1": "受试者入组", "weight": 0.45},
            {"level1": "受试者随访", "weight": 0.45},
            {"level1": "数据分析", "weight": 0.10},
        ]
    },
    "Global team费用": {
        "phases": [
            {"level1": "研究准备阶段", "weight": 0.10},
            {"level1": "研究文件准备", "weight": 0.10},
            {"level1": "组长单位GCP&伦理递交", "weight": 0.10},
            {"level1": "研究中心启动", "weight": 0.10},
            {"level1": "受试者入组", "weight": 0.25},
            {"level1": "受试者随访", "weight": 0.15},
            {"level1": "数据分析", "weight": 0.10},
            {"level1": "研究结束", "weight": 0.10},
        ]
    },
}

TIMELINE_TEMPLATE = {
    "研究准备阶段": [
        "项目立项",
        "研究中心名单提名、筛选及访视",
        "产品及人员培训",
        "供应商筛选与确认",
        "产品运输及相关管理流程准备",
        "研究者会议",
    ],
    "研究文件准备": [
        "研究方案撰写及获批",
        "CRF设计及数据库建立",
        "知情同意书撰写及获批",
        "研究计划书撰写及获批（监查计划、数据管理计划、统计分析计划等）",
        "其他相关表格及物料准备（受试者日记卡、招募海报、标签等）",
        "文件翻译及定稿",
    ],
    "政府部门备案": [
        "上海药监局备案",
        "人类遗传办提交",
        "人类遗传办获批",
    ],
    "组长单位GCP&伦理递交": [
        "组长单位GCP资料递交",
        "组长单位EC申请及获批",
        "其他中心EC申请及获批",
    ],
    "组长单位合同签署": [
        "组长单位合同起草与协商",
        "组长单位合同签字与盖章",
    ],
    "研究中心启动": [
        "组长单位启动",
        "其他中心启动",
    ],
    "受试者入组": [
        "组长单位首例入组",
        "其他中心首例入组",
        "受试者入组",
        "末例入组",
    ],
    "受试者随访": [
        "主要终点随访",
        "随访结束",
        "末例受试者末次访视",
    ],
    "数据分析": [
        "主要终点数据库锁定",
        "主要终点TLG/统计报告撰写",
        "主要终点临床研究报告撰写",
        "主要终点数据的监管提交",
        "数据库清理与锁定",
        "最终图表TLG/统计报告完成",
        "临床研究报告撰写",
    ],
    "研究结束": [
        "中心关闭访视",
        "临床研究报告的监管递交",
    ],
}

# ---------- Pure calculation engine ----------
PHASES = list(TIMELINE_TEMPLATE)
ITEM_PHASE = {
    item: phase for phase, items in TIMELINE_TEMPLATE.items() for item in items
}
MODULE_ALIAS = {
    f"预算模块{i+1}：{old}": new
    for i, (old, new) in enumerate(
        zip(["研究中心花费", "Vendor 花费", "产品物流成本", "人力资源成本"], MODULES)
    )
}


def money(value: float | Decimal) -> float:
    return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def as_date(value: Any) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return dt.date.fromisoformat(str(value))


def add_months(value: Any, count: int) -> dt.date:
    return as_date(value) + relativedelta(months=int(count))


def month_starts(start: Any, end: Any) -> list[dt.date]:
    start, end = as_date(start), as_date(end)
    if end < start:
        raise ValueError("结束日期不能早于开始日期。")
    current, last = start.replace(day=1), end.replace(day=1)
    result = []
    while current <= last:
        result.append(current)
        current = add_months(current, 1)
    return result


def duration_months(start: Any, end: Any) -> float:
    """Inclusive calendar duration: 1 Jan–31 Jan = 1; a single day > 0."""
    start, end = as_date(start), as_date(end)
    if end < start:
        raise ValueError("结束日期不能早于开始日期。")
    end_exclusive = end + dt.timedelta(days=1)
    rd = relativedelta(end_exclusive, start)
    whole = rd.years * 12 + rd.months
    anchor = add_months(start, whole)
    next_anchor = add_months(start, whole + 1)
    return whole + (end_exclusive - anchor).days / (next_anchor - anchor).days


def split_money(total: float, weights: list[float]) -> list[float]:
    """Largest remainder allocation in cents; exactly conserves the line total."""
    if not weights or any(not math.isfinite(float(x)) or x < 0 for x in weights):
        raise ValueError("分配权重必须是有限的非负数字。")
    w = [Decimal(str(x)) for x in weights]
    if sum(w) <= 0 or total < 0 or not math.isfinite(float(total)):
        raise ValueError("分配权重合计必须大于0，金额不得为负。")
    cents = int(Decimal(str(money(total))) * 100)
    exact = [Decimal(cents) * x / sum(w) for x in w]
    units = [int(x.to_integral_value(rounding=ROUND_FLOOR)) for x in exact]
    order = sorted(range(len(w)), key=lambda i: exact[i] - units[i], reverse=True)
    for i in order[: cents - sum(units)]:
        units[i] += 1
    return [x / 100 for x in units]


def to_rmb(value: float, currency: str, fx: float) -> float:
    if currency not in ("RMB", "USD") or fx <= 0:
        raise ValueError("币种必须是RMB或USD，汇率必须大于0。")
    return money(Decimal(str(value)) * (Decimal(str(fx)) if currency == "USD" else 1))


def display_amount(rmb: float, currency: str, fx: float) -> float:
    return (
        money(Decimal(str(rmb)) / Decimal(str(fx))) if currency == "USD" else money(rmb)
    )


def fiscal_label(date: Any, start_month: int = 6, legacy: bool = False) -> str:
    d = as_date(date)
    if legacy and start_month == 6 and d.year == 2026 and d.month == 2:
        return "FY26Q4"
    offset = (d.month - start_month) % 12
    fy = d.year if start_month == 1 else d.year + int(d.month >= start_month)
    return f"FY{fy % 100:02d}Q{offset // 3 + 1}"


def quarter_months(label: str, start_month: int = 6) -> list[dt.date]:
    m = re.fullmatch(r"FY(\d{2}|20\d{2})Q([1-4])", str(label).upper())
    if not m:
        raise ValueError(f"无效财季：{label}，示例FY27Q1。")
    fy = int(m[1])
    if fy < 100:
        fy += 2000
    start = dt.date(fy if start_month == 1 else fy - 1, start_month, 1)
    first = add_months(start, (int(m[2]) - 1) * 3)
    return [add_months(first, i) for i in range(3)]


def project_span(project: dict) -> tuple[dt.date, dt.date]:
    rows = [x for x in project["timeline"] if x["enabled"]]
    if rows:
        return min(as_date(x["start"]) for x in rows), max(
            as_date(x["end"]) for x in rows
        )
    d = as_date(project["settings"]["start_date"])
    return d, d


def enrollment_curve(project: dict, rate_factor: float = 1.0) -> list[dict]:
    """Expected cohort curve; linearly ramps active sites, no subject-level data."""
    s = project["settings"]
    enrollment = next(
        (
            t
            for t in project.get("timeline", [])
            if t["enabled"] and t["task"] == "受试者入组"
        ),
        None,
    )
    first = (
        as_date(enrollment["start"])
        if enrollment
        else add_months(s["start_date"], s["startup_months"])
    )
    remaining = float(s["n_planned"])
    rows = []
    for i in range(600):
        active = s["site_number"] * min(1.0, (i + 1) / max(1, s["ramp_months"]))
        capacity = active * s["enrollment_rate"] * rate_factor
        amount = min(remaining, capacity)
        remaining = max(0.0, remaining - amount)
        rows.append(
            {
                "month": add_months(first, i).replace(day=1),
                "active_sites": active,
                "enrolled": amount,
                "cumulative": s["n_planned"] - remaining,
            }
        )
        if remaining <= 1e-8:
            return rows
    raise ValueError("当前入组效率推算超过50年，请提高入组速率或中心数量。")


def make_timeline(project: dict) -> list[dict]:
    s = project["settings"]
    if project["template"] == "示例项目":
        shift = as_date(s["start_date"]) - EXAMPLE_PROJECT_DEFAULTS["start_date"]
        dates = {
            k: (v[0], v[1] + shift, v[2] + shift)
            for k, v in EXAMPLE_TIMELINE_DEFAULTS.items()
        }
    else:
        start = as_date(s["start_date"])
        fpi = add_months(start, s["startup_months"])
        curve = enrollment_curve(project)
        lpi = add_months(fpi, len(curve)) - dt.timedelta(days=1)
        primary = add_months(lpi, s["primary_fu_months"])
        final = add_months(lpi, s["final_fu_months"])
        primary_lock = add_months(primary, s["cleaning_months"])
        final_lock = add_months(final, s["cleaning_months"])
        primary_report = add_months(primary_lock, s["report_months"])
        final_report = add_months(final_lock, s["report_months"])
        before = lambda months: min(
            add_months(start, months), fpi - dt.timedelta(days=1)
        )
        dates = {}
        for phase, items in TIMELINE_TEMPLATE.items():
            for item in items:
                a, b = start, before(2)
                if phase == "研究文件准备":
                    a, b = start, before(3)
                elif phase in (
                    "政府部门备案",
                    "组长单位GCP&伦理递交",
                    "组长单位合同签署",
                ):
                    a, b = before(2), fpi - dt.timedelta(days=1)
                elif phase == "研究中心启动":
                    a, b = fpi - dt.timedelta(days=1), add_months(
                        fpi, s["ramp_months"]
                    ) - dt.timedelta(days=1)
                elif phase == "受试者入组":
                    a, b = fpi, lpi
                    if item == "组长单位首例入组":
                        b = fpi
                    elif item == "其他中心首例入组":
                        b = min(
                            lpi,
                            add_months(fpi, s["ramp_months"]) - dt.timedelta(days=1),
                        )
                    elif item == "末例入组":
                        a = lpi
                elif phase == "受试者随访":
                    a, b = fpi, primary if item == "主要终点随访" else final
                    if item == "末例受试者末次访视":
                        a = final
                elif phase == "数据分析":
                    if item.startswith("主要终点"):
                        a, b = primary + dt.timedelta(days=1), primary_report
                        if "锁定" in item:
                            b = primary_lock
                        elif "监管提交" in item:
                            a, b = primary_report, add_months(primary_report, 1)
                        else:
                            a = primary_lock
                    else:
                        a, b = final_lock, final_report
                        if "锁定" in item:
                            a, b = final + dt.timedelta(days=1), final_lock
                elif phase == "研究结束":
                    a, b = final_report, add_months(final_report, 2)
                if item == "研究者会议":
                    a, b = before(2), before(2)
                if item == "产品及人员培训":
                    a, b = before(3), fpi - dt.timedelta(days=1)
                dates[item] = (True, a, max(a, b))
    # Only hard finish-to-start links are enforced. Parallel workflows stay parallel.
    predecessors = {
        "人类遗传办获批": "人类遗传办提交",
        "组长单位合同签字与盖章": "组长单位合同起草与协商",
        "主要终点临床研究报告撰写": "主要终点数据库锁定",
        "主要终点数据的监管提交": "主要终点临床研究报告撰写",
        "临床研究报告撰写": "数据库清理与锁定",
        "临床研究报告的监管递交": "临床研究报告撰写",
    }
    ids = {item: f"T{i+1:03d}" for i, item in enumerate(ITEM_PHASE)}
    old = {t["id"]: t for t in project.get("timeline", [])}
    rows = []
    for item, phase in ITEM_PHASE.items():
        enabled, a, b = dates[item]
        prev = old.get(ids[item], {})
        # EXAMPLE historical overlapping windows are preserved and flagged, not silently rewritten.
        rows.append(
            dict(
                id=ids[item],
                phase=phase,
                task=item,
                enabled=enabled,
                start=a.isoformat(),
                end=b.isoformat(),
                owner=prev.get("owner", ""),
                progress=prev.get("progress", 0.0),
                actual_end=prev.get("actual_end"),
                depends_on=ids.get(predecessors.get(item), ""),
            )
        )
    return rows


def cost_row(
    item: str,
    module: str,
    unit: float,
    basis: str = "固定数量",
    qty: float = 1,
    currency: str = "RMB",
    **kwargs,
) -> dict:
    return dict(
        id=kwargs.pop("id", "B" + uuid.uuid4().hex[:10]),
        item=item,
        module=module,
        enabled=True,
        unit_price=float(unit),
        basis=basis,
        quantity=float(qty),
        currency=currency,
        phase="智能匹配",
        allocation="智能规则",
        payment_lag=0,
        note="",
        **kwargs,
    )


def make_project(template: str = "示例项目") -> dict:
    if template in ("ATHENA", "EXAMPLE"):
        template = "示例项目"
    settings = dict(
        project_name=(
            "示例项目（预填演示）" if template == "示例项目" else "新建临床研究项目"
        ),
        study_type="上市前临床研究",
        site_number=20 if template == "示例项目" else 10,
        n_planned=142 if template == "示例项目" else 150,
        n_lost=0,
        n_screen_fail=0 if template == "示例项目" else 10,
        start_date=(
            "2026-03-11" if template == "示例项目" else dt.date.today().isoformat()
        ),
        primary_fu_months=12,
        final_fu_months=36 if template == "示例项目" else 24,
        enrollment_rate=0.8 if template == "示例项目" else 0.5,
        startup_months=7 if template == "示例项目" else 6,
        ramp_months=4 if template == "示例项目" else 3,
        cleaning_months=2,
        report_months=3,
        fx=7.0,
        display_currency="USD" if template == "示例项目" else "RMB",
        fy_start_month=6,
        legacy_fy=template == "示例项目",
        cash_mode="示例项目预填季度" if template == "示例项目" else "智能匹配时间线",
        site_management_pct=0.0 if template == "示例项目" else 25.0,
        site_tax_pct=0.0 if template == "示例项目" else 6.0,
        vendor_tax_pct=0.0 if template == "示例项目" else 6.0,
        tax_includes_management=True,
        reserve_pct=0.0,
        approved_budget_rmb=0.0,
    )
    p = dict(
        schema_version=SCHEMA_VERSION,
        app_version=APP_VERSION,
        template=template,
        settings=settings,
        timeline=[],
        timeline_mode="模板排期" if template == "示例项目" else "自动联动",
        budget=[],
        visits=[],
        cash_weights=[],
        expenses=[],
        baseline=None,
        sites=[],
        enrollment=[],
    )
    if template == "示例项目":
        for i, (item, vals) in enumerate(EXAMPLE_BUDGET_MATRIX_USD.items()):
            module = MODULE_ALIAS[
                EXAMPLE_BUDGET_MODULE_MAP.get(item, "预算模块2：Vendor 花费")
            ]
            p["budget"].append(
                cost_row(
                    item, module, money(sum(vals)), currency="USD", id=f"A{i+1:03d}"
                )
            )
        # Included in original lump sum Site Cost: no extra visit charges by default.
        for name, offset in [
            ("出院", 0),
            ("30天", 1),
            ("6个月", 6),
            ("1年", 12),
            ("2年", 24),
            ("3年", 36),
        ]:
            p["visits"].append(
                dict(
                    id="V" + uuid.uuid4().hex[:8],
                    enabled=False,
                    name=name,
                    offset=offset,
                    times=1.0,
                    unit_price=0.0,
                    currency="USD",
                    population="入组例数",
                )
            )
    else:
        specs = [
            ("研究者筛选费", "研究中心", 500, "入组例数", 1),
            ("研究者手术费", "研究中心", 5000, "入组例数", 1),
            ("筛选失败费用", "研究中心", 300, "筛选失败例数", 1),
            ("受试者补偿", "研究中心", 1000, "入组例数", 1),
            ("受试者检查费用", "研究中心", 2000, "入组例数", 1),
            ("组长单位费用", "研究中心", 200000, "固定数量", 1),
            ("EC/Clinical Institute Process费用", "研究中心", 10000, "中心数", 1),
            ("CRC费用", "供应商", 300, "固定数量", 200),
            ("PM费用", "供应商", 400, "固定数量", 150),
            ("Monitor费用", "供应商", 350, "固定数量", 300),
            ("DM费用", "供应商", 300, "固定数量", 150),
            ("Safety费用", "供应商", 350, "固定数量", 80),
            ("EDC系统费用", "供应商", 200000, "固定数量", 1),
            ("翻译/打印费用", "供应商", 50000, "固定数量", 1),
            ("保险费用", "供应商", 80000, "固定数量", 1),
            ("数据分析费用", "供应商", 150000, "固定数量", 1),
            ("中心实验室费用", "供应商", 100000, "固定数量", 1),
            ("CEC费用", "供应商", 120000, "固定数量", 1),
            ("Travel费用", "供应商", 2000, "固定数量", 20),
            ("Recording费用", "供应商", 200, "入组例数", 1),
            ("研究者会费用", "供应商", 150000, "固定数量", 1),
            ("产品运输费用", "产品物流", 3000, "固定数量", 10),
            ("仓储费用", "产品物流", 2000, "项目月数", 1),
            ("产品成本", "产品物流", 10000, "入组例数", 1),
            ("项目经理FTE费用", "人力资源", 300000, "FTE项目", 0.3),
            ("CRA费用", "人力资源", 250000, "FTE项目", 0.5),
            ("MCRS费用", "人力资源", 180000, "FTE项目", 0.2),
            ("Global team费用", "人力资源", 500000, "FTE项目", 0.1),
        ]
        p["budget"] = [cost_row(*x, id=f"B{i+1:03d}") for i, x in enumerate(specs)]
        for i, offset in enumerate([1, 6, 12]):
            p["visits"].append(
                dict(
                    id=f"V{i+1:03d}",
                    enabled=True,
                    name=f"{offset}个月",
                    offset=offset,
                    times=1.0,
                    unit_price=300.0,
                    currency="RMB",
                    population="入组例数",
                )
            )
    p["timeline"] = make_timeline(p)
    p["cash_weights"] = default_weights(p)
    return p


def default_weights(project: dict) -> list[dict]:
    a, b = project_span(project)
    s = project["settings"]
    labels = list(
        dict.fromkeys(
            fiscal_label(m, s["fy_start_month"], s["legacy_fy"])
            for m in month_starts(a, b)
        )
    )
    portions = split_money(100.0, [1] * len(labels))
    return [dict(quarter=q, weight=w) for q, w in zip(labels, portions)]


def enrollment_window(project: dict) -> tuple[dt.date, dt.date]:
    t = next(
        (t for t in project["timeline"] if t["enabled"] and t["task"] == "受试者入组"),
        None,
    )
    if t:
        return as_date(t["start"]), as_date(t["end"])
    return project_span(project)


def budget_lines(project: dict) -> list[dict]:
    s = project["settings"]
    a, b = project_span(project)
    months = duration_months(a, b)
    quantities = {
        "入组例数": s["n_planned"],
        "完成随访例数": s["n_planned"] - s["n_lost"],
        "筛选失败例数": s["n_screen_fail"],
        "筛选总例数": s["n_planned"] + s["n_screen_fail"],
        "中心数": s["site_number"],
        "项目月数": months,
    }
    lines = []
    for r in project["budget"]:
        if not r["enabled"]:
            continue
        qty = quantities.get(r["basis"], r["quantity"])
        if r["basis"] == "FTE月":
            qty = r["quantity"] * months
        elif r["basis"] == "FTE年":
            qty = r["quantity"] * months / 12
        native = money(Decimal(str(r["unit_price"])) * Decimal(str(qty)))
        lines.append(
            dict(
                r,
                calc_quantity=qty,
                native_amount=native,
                amount_rmb=to_rmb(native, r["currency"], s["fx"]),
                kind="direct",
            )
        )
    for v in project["visits"]:
        if not v["enabled"]:
            continue
        qty = quantities[v["population"]] * v["times"]
        native = money(Decimal(str(v["unit_price"])) * Decimal(str(qty)))
        r = cost_row(
            "研究者随访费 - " + v["name"],
            "研究中心",
            v["unit_price"],
            currency=v["currency"],
            id=v["id"],
        )
        lines.append(
            dict(
                r,
                calc_quantity=qty,
                native_amount=native,
                amount_rmb=to_rmb(native, v["currency"], s["fx"]),
                kind="visit",
                visit_offset=v["offset"],
            )
        )

    def derived(id, item, module, base, pct, kind):
        amt = money(Decimal(str(base)) * Decimal(str(pct)) / 100)
        r = cost_row(item, module, amt, id=id)
        lines.append(
            dict(
                r,
                calc_quantity=1,
                native_amount=amt,
                amount_rmb=amt,
                kind=kind,
                note=f"计费基数 RMB {base:,.2f} × {pct:g}%",
            )
        )
        return amt

    site = money(sum(x["amount_rmb"] for x in lines if x["module"] == "研究中心"))
    vendor = money(sum(x["amount_rmb"] for x in lines if x["module"] == "供应商"))
    mgmt = derived(
        "D-MGMT",
        "中心管理费用",
        "研究中心",
        site,
        s["site_management_pct"],
        "management",
    )
    derived(
        "D-SITETAX",
        "Site税费",
        "研究中心",
        site + (mgmt if s["tax_includes_management"] else 0),
        s["site_tax_pct"],
        "tax",
    )
    derived("D-VENDORTAX", "Vendor税费", "供应商", vendor, s["vendor_tax_pct"], "tax")
    derived(
        "D-RESERVE",
        "风险储备（未承诺）",
        "风险储备",
        money(sum(x["amount_rmb"] for x in lines)),
        s["reserve_pct"],
        "reserve",
    )
    return lines


def phase_range(
    project: dict, phase: str | None = None, task: str | None = None
) -> tuple[dt.date, dt.date, bool]:
    tasks = [
        t
        for t in project["timeline"]
        if t["enabled"] and (t["task"] == task if task else t["phase"] == phase)
    ]
    if tasks:
        return (
            min(as_date(t["start"]) for t in tasks),
            max(as_date(t["end"]) for t in tasks),
            True,
        )
    a, b = project_span(project)
    return a, b, False


def allocate_cashflow(
    project: dict, lines: list[dict] | None = None
) -> tuple[list[dict], list[str]]:
    lines = budget_lines(project) if lines is None else lines
    s = project["settings"]
    a, b = project_span(project)
    mode = s["cash_mode"]
    output = []
    warnings = []

    def spread(amount, start, end, line, label, weights=None):
        months = month_starts(start, end)
        parts = split_money(amount, weights or [1] * len(months))
        for m, part in zip(months, parts):
            payment = add_months(m, line.get("payment_lag", 0))
            output.append(
                dict(
                    month=payment.isoformat(),
                    accrual_month=m.isoformat(),
                    id=line["id"],
                    item=line["item"],
                    module=line["module"],
                    phase=label,
                    amount_rmb=part,
                    quarter=fiscal_label(payment, s["fy_start_month"], s["legacy_fy"]),
                )
            )

    for line in lines:
        amount = line["amount_rmb"]
        if not amount:
            continue
        if mode == "自定义季度占比":
            specs = project["cash_weights"]
            portions = split_money(amount, [r["weight"] for r in specs])
            for r, part in zip(specs, portions):
                months = quarter_months(r["quarter"], s["fy_start_month"])
                if (
                    s["legacy_fy"]
                    and s["fy_start_month"] == 6
                    and r["quarter"] == "FY26Q4"
                ):
                    months = [dt.date(2026, 2, 1)] + months
                spread(part, months[0], months[-1], line, r["quarter"])
        elif mode == "示例项目预填季度" and line["item"] in EXAMPLE_BUDGET_MATRIX_USD:
            vals = EXAMPLE_BUDGET_MATRIX_USD[line["item"]]
            portions = split_money(amount, vals)
            for q, part in zip(EXAMPLE_BUDGET_QUARTERS, portions):
                if not part:
                    continue
                months = quarter_months(q, 6)
                spread(part, months[0], months[-1], line, q)
        elif mode == "全周期按月均摊":
            spread(amount, a, b, line, "全周期按月均摊")
        elif line["allocation"] != "智能规则" or line["phase"] != "智能匹配":
            phase = None if line["phase"] == "智能匹配" else line["phase"]
            start, end, matched = phase_range(project, phase=phase)
            if not matched:
                warnings.append(f"{line['item']}：指定阶段未启用，已按全周期摊分。")
            if line["allocation"] == "阶段开始一次性":
                end = start
            elif line["allocation"] == "阶段结束一次性":
                start = end
            spread(amount, start, end, line, phase or "全项目周期")
        elif line["kind"] == "visit":
            if not any(
                t["enabled"] and t["task"] == "受试者入组" for t in project["timeline"]
            ):
                warnings.append(
                    f"{line['item']}：未启用受试者入组任务，访视费用暂按全项目入组窗口推算。"
                )
            start, end = enrollment_window(project)
            cohorts = month_starts(start, end)
            portions = split_money(amount, [1] * len(cohorts))
            for cohort, part in zip(cohorts, portions):
                date = add_months(cohort, line["visit_offset"])
                spread(part, date, date, line, f"入组后{line['visit_offset']}月访视")
        else:
            specs = BUDGET_ALLOCATION_RULES.get(line["item"], {}).get("phases", [])
            if not specs:
                spread(amount, a, b, line, "全项目周期")
                if line["kind"] == "direct":
                    warnings.append(
                        f"{line['item']}：无预设匹配规则，已按全周期摊分；可指定阶段。"
                    )
            else:
                portions = split_money(amount, [x["weight"] for x in specs])
                for spec, part in zip(specs, portions):
                    start, end, matched = phase_range(
                        project, phase=spec.get("level1"), task=spec.get("level2")
                    )
                    label = spec.get("level2") or spec.get("level1")
                    if not matched:
                        warnings.append(
                            f"{line['item']}：{label}未启用，相关金额按全周期摊分。"
                        )
                    if spec.get("mode") == "lump_start":
                        end = start
                    elif spec.get("mode") == "lump_end":
                        start = end
                    spread(
                        part,
                        start,
                        end,
                        line,
                        label if matched else "全项目周期（缺少映射阶段）",
                    )
            if mode == "示例项目预填季度":
                warnings.append(f"{line['item']}：原始矩阵无此项目，已使用智能规则。")
    return output, list(dict.fromkeys(warnings))


def quarter_totals(cash: list[dict], project: dict) -> list[dict]:
    totals = {}
    for r in cash:
        totals[r["quarter"]] = money(totals.get(r["quarter"], 0) + r["amount_rmb"])
    total = money(sum(totals.values()))
    return [
        dict(quarter=q, amount_rmb=amt, share=amt / total * 100 if total else 0)
        for q, amt in sorted(totals.items())
    ]


# ---------- Validation / portable persistence ----------
def validate_project(project: dict) -> list[str]:
    """Validate untrusted JSON and UI edits before committing any state."""
    errors = []
    if not isinstance(project, dict) or project.get("schema_version") != SCHEMA_VERSION:
        return ["仅支持本工具V2导出的项目JSON（schema_version=2）。"]
    required = [
        "settings",
        "template",
        "timeline",
        "timeline_mode",
        "budget",
        "visits",
        "cash_weights",
        "expenses",
        "baseline",
    ]
    if any(k not in project for k in required):
        return ["项目文件字段不完整。"]
    s = project["settings"]
    if not isinstance(s, dict):
        return ["settings必须为对象。"]
    expected = make_settings_keys()
    if not expected.issubset(s):
        return ["项目设置字段缺失：" + ", ".join(sorted(expected - set(s)))]

    def number(v, label, minimum=0, maximum=1e12, integer=False):
        if (
            isinstance(v, bool)
            or not isinstance(v, (float, int))
            or not math.isfinite(v)
        ):
            errors.append(f"{label}必须是有限数字。")
        elif v < minimum or v > maximum or (integer and v != int(v)):
            errors.append(
                f"{label}应为{minimum:g}–{maximum:g}之间的{'整数' if integer else '数字'}。"
            )

    def text(v, label, limit=500):
        if not isinstance(v, str) or not v.strip() or len(v) > limit:
            errors.append(f"{label}不能为空且不能超过{limit}字。")

    def date(v, label, nullable=False):
        if nullable and v is None:
            return
        try:
            d = as_date(v)
            if not dt.date(2000, 1, 1) <= d <= dt.date(2099, 12, 31):
                raise ValueError()
        except (ValueError, TypeError, OverflowError):
            errors.append(f"{label}须为2000–2099年内有效日期。")

    text(s["project_name"], "项目名称", 120)
    text(s["study_type"], "研究类型", 120)
    if project["template"] not in ("示例项目", "通用"):
        errors.append("未知项目模板。")
    if project["timeline_mode"] not in ("模板排期", "自动联动", "手工维护"):
        errors.append("未知排期模式。")
    for key in ["site_number", "n_planned"]:
        number(s[key], key, 1, 100000, True)
    for key in ["n_lost", "n_screen_fail"]:
        number(s[key], key, 0, 100000, True)
    for key in ["primary_fu_months", "final_fu_months"]:
        number(s[key], key, 0, 240, True)
    for key in ["startup_months", "ramp_months", "cleaning_months", "report_months"]:
        number(s[key], key, 1, 120, True)
    number(s["enrollment_rate"], "每中心月入组速率", 0.01, 10000)
    number(s["fx"], "汇率", 0.01, 1000)
    number(s["fy_start_month"], "财年起始月", 1, 12, True)
    for key in ["site_management_pct", "site_tax_pct", "vendor_tax_pct", "reserve_pct"]:
        number(s[key], key, 0, 100)
    number(s["approved_budget_rmb"], "获批预算", 0, 1e15)
    date(s["start_date"], "研究基准日期")
    for key in ["legacy_fy", "tax_includes_management"]:
        if not isinstance(s[key], bool):
            errors.append(f"{key}必须为布尔值。")
    if s["display_currency"] not in ("RMB", "USD"):
        errors.append("展示币种无效。")
    if s["cash_mode"] not in CASH_MODES:
        errors.append("现金流方式无效。")
    if not errors:
        if s["n_lost"] > s["n_planned"]:
            errors.append("预估失访人数不能超过计划入组人数。")
        if s["primary_fu_months"] > s["final_fu_months"]:
            errors.append("主要终点随访时长不能超过最终随访时长。")
        if s["cash_mode"] == "示例项目预填季度" and (
            project["template"] != "示例项目" or s["fy_start_month"] != 6
        ):
            errors.append(
                "示例项目原始矩阵采用6月起始的项目财季，请保持该口径或切换现金流方式。"
            )
    collections = {
        "timeline": 300,
        "budget": 1000,
        "visits": 200,
        "cash_weights": 400,
        "expenses": 5000,
    }
    for k, limit in collections.items():
        if (
            not isinstance(project[k], list)
            or len(project[k]) > limit
            or any(not isinstance(x, dict) for x in project[k])
        ):
            errors.append(f"{k}须为对象列表，最多{limit}行。")
    if errors:
        return errors
    if not any(t.get("enabled") is True for t in project["timeline"]):
        errors.append("至少启用一个时间线任务。")
    id_sets = {}
    for key in ("timeline", "budget", "visits", "expenses"):
        ids = [r.get("id") for r in project[key]]
        if any(not isinstance(i, str) or not i or len(i) > 80 for i in ids) or len(
            set(ids)
        ) != len(ids):
            errors.append(f"{key}的ID不能为空或重复。")
        id_sets[key] = set(i for i in ids if isinstance(i, str))
    if id_sets["budget"] & id_sets["visits"]:
        errors.append("预算与访视ID不能重复。")
    if (id_sets["budget"] | id_sets["visits"]) & {
        "D-MGMT",
        "D-SITETAX",
        "D-VENDORTAX",
        "D-RESERVE",
    }:
        errors.append("直接费用不能使用系统派生费用ID。")
    expected_fields = {
        "timeline": {
            "id",
            "task",
            "phase",
            "enabled",
            "start",
            "end",
            "owner",
            "progress",
            "actual_end",
            "depends_on",
        },
        "budget": {
            "id",
            "item",
            "module",
            "enabled",
            "unit_price",
            "currency",
            "basis",
            "quantity",
            "phase",
            "allocation",
            "payment_lag",
            "note",
        },
        "visits": {
            "id",
            "name",
            "enabled",
            "offset",
            "times",
            "unit_price",
            "currency",
            "population",
        },
        "expenses": {"id", "date", "budget_id", "amount", "currency", "note"},
        "cash_weights": {"quarter", "weight"},
    }
    for key, fields in expected_fields.items():
        for i, r in enumerate(project[key]):
            if not fields.issubset(r):
                errors.append(f"{key}第{i+1}行字段缺失。")
    if errors:
        return errors
    task_ids = id_sets["timeline"]
    for t in project["timeline"]:
        text(t["task"], "任务名称", 250)
        if t["phase"] not in PHASES:
            errors.append("任务阶段须从预设阶段选择。")
        if not isinstance(t["enabled"], bool):
            errors.append("任务启用状态须为布尔值。")
        if not isinstance(t["owner"], str) or len(t["owner"]) > 200:
            errors.append("负责人字段过长或无效。")
        date(t["start"], t["task"] + "开始日期")
        date(t["end"], t["task"] + "结束日期")
        date(t["actual_end"], t["task"] + "实际完成日期", True)
        number(t["progress"], "完成度", 0, 100)
        if t["depends_on"] and t["depends_on"] not in task_ids:
            errors.append(t["task"] + "的前置任务ID不存在。")
        try:
            if as_date(t["end"]) < as_date(t["start"]):
                errors.append(t["task"] + "的结束日期早于开始日期。")
        except (ValueError, TypeError):
            pass
    mapping = {t["id"]: t["depends_on"] for t in project["timeline"]}
    for origin in mapping:
        visited = set()
        node = origin
        while node and node in mapping:
            if node in visited:
                errors.append("任务依赖存在循环，请调整前置ID。")
                break
            visited.add(node)
            node = mapping[node]
        if errors and "循环" in errors[-1]:
            break
    if len({t["task"] for t in project["timeline"]}) != len(project["timeline"]):
        errors.append("任务名称不能重复。")
    for r in project["budget"]:
        text(r["item"], "费用名称", 250)
        if not isinstance(r["enabled"], bool):
            errors.append("费用启用状态无效。")
        if r["module"] not in MODULES:
            errors.append("费用模块无效。")
        if r["basis"] not in BASIS:
            errors.append("费用计费依据无效。")
        if r["currency"] not in ("RMB", "USD"):
            errors.append("费用币种无效。")
        if r["phase"] not in ["智能匹配"] + PHASES:
            errors.append("预算关联阶段无效。")
        if r["allocation"] not in ALLOC_MODES:
            errors.append("费用分配方式无效。")
        if r["allocation"] != "智能规则" and r["phase"] == "智能匹配":
            errors.append(r["item"] + "：请为手工分配方式指定阶段。")
        number(r["unit_price"], "费用单价", 0, 1e10)
        number(r["quantity"], "固定数量/FTE", 0, 1e8)
        number(r["payment_lag"], "付款滞后月数", 0, 120, True)
        if not isinstance(r["note"], str) or len(r["note"]) > 2000:
            errors.append("费用备注无效或超过2000字。")
    for v in project["visits"]:
        text(v["name"], "访视名称", 100)
        if not isinstance(v["enabled"], bool):
            errors.append("访视启用状态无效。")
        number(v["offset"], "访视月份", 0, 240, True)
        number(v["times"], "每例次数", 0, 1000)
        number(v["unit_price"], "访视单价", 0, 1e10)
        if v["currency"] not in ("RMB", "USD") or v["population"] not in (
            "入组例数",
            "完成随访例数",
        ):
            errors.append("访视币种或计费人群无效。")
    if len({r["quarter"] for r in project["cash_weights"]}) != len(
        project["cash_weights"]
    ):
        errors.append("自定义季度不能重复。")
    for r in project["cash_weights"]:
        try:
            quarter_months(r["quarter"], s["fy_start_month"])
        except (ValueError, TypeError, OverflowError):
            errors.append("自定义财季格式应为FY27Q1。")
        number(r["weight"], "季度占比", 0, 100)
    if s["cash_mode"] == "自定义季度占比" and not any("季度占比" in e for e in errors):
        if abs(sum(r["weight"] for r in project["cash_weights"]) - 100) > 0.011:
            errors.append("自定义季度占比合计须为100%；可点击归一化后再保存。")
    valid_budget_ids = (
        id_sets["budget"]
        | id_sets["visits"]
        | {"D-MGMT", "D-SITETAX", "D-VENDORTAX", "D-RESERVE"}
    )
    for e in project["expenses"]:
        date(e["date"], "支出日期")
        number(e["amount"], "实际支出", 0, 1e12)
        if e["budget_id"] not in valid_budget_ids:
            errors.append("实际支出关联的费用ID不存在。")
        if e["currency"] not in ("RMB", "USD"):
            errors.append("实际支出币种无效。")
        if not isinstance(e["note"], str) or len(e["note"]) > 2000:
            errors.append("实际支出备注无效。")
    baseline = project["baseline"]
    if baseline is not None:
        if not isinstance(baseline, dict) or not {
            "total_rmb",
            "end",
            "captured_at",
            "name",
            "timeline",
            "line_amounts",
        }.issubset(baseline):
            errors.append("基线数据不完整。")
        else:
            number(baseline["total_rmb"], "基线预算", 0, 1e16)
            date(baseline["end"], "基线结束日期")
            if (
                not isinstance(baseline["timeline"], list)
                or len(baseline["timeline"]) > 300
            ):
                errors.append("基线时间线无效。")
            elif any(
                not isinstance(t, dict)
                or not {"id", "start", "end", "task"}.issubset(t)
                for t in baseline["timeline"]
            ):
                errors.append("基线任务数据不完整。")
            else:
                for t in baseline["timeline"]:
                    date(t["start"], "基线任务开始日期")
                    date(t["end"], "基线任务结束日期")
            if not isinstance(baseline["line_amounts"], dict):
                errors.append("基线费用数据无效。")
            else:
                for amt in baseline["line_amounts"].values():
                    number(amt, "基线费用", 0, 1e16)
    if not errors:
        a, b = project_span(project)
        if (b - a).days > 366 * 25:
            errors.append("项目跨度不能超过25年。")
        if any(t["enabled"] and t["task"] == "受试者入组" for t in project["timeline"]):
            _, last = enrollment_window(project)
            if add_months(last, s["final_fu_months"]).year > 2099:
                errors.append("随访推算日期超出2099年。")
    errors.extend(validate_sites(project))
    return list(dict.fromkeys(errors))


def make_settings_keys() -> set[str]:
    return set(
        "project_name study_type site_number n_planned n_lost n_screen_fail start_date primary_fu_months final_fu_months enrollment_rate startup_months ramp_months cleaning_months report_months fx display_currency fy_start_month legacy_fy cash_mode site_management_pct site_tax_pct vendor_tax_pct tax_includes_management reserve_pct approved_budget_rmb".split()
    )


def load_project(raw: bytes) -> dict:
    if len(raw) > 5_000_000:
        raise ValueError("项目文件不能超过5 MB。")
    try:
        project = json.loads(
            raw.decode("utf-8-sig"),
            parse_constant=lambda x: (_ for _ in ()).throw(
                ValueError("不允许NaN/Infinity")
            ),
        )
    except (UnicodeError, json.JSONDecodeError, ValueError, RecursionError) as exc:
        raise ValueError("无法读取项目JSON，请选择本工具导出的有效文件。") from exc
    project = upgrade_project(project)
    try:
        errors = validate_project(project)
    except (
        TypeError,
        KeyError,
        AttributeError,
        ValueError,
        OverflowError,
        RecursionError,
    ) as exc:
        raise ValueError("项目JSON字段类型或格式错误。") from exc
    if errors:
        raise ValueError("；".join(errors[:8]))
    return project


def project_bytes(project: dict) -> bytes:
    return json.dumps(project, ensure_ascii=False, indent=2, allow_nan=False).encode(
        "utf-8"
    )


def save_baseline(project: dict, name: str = "项目基线") -> dict:
    lines = budget_lines(project)
    return dict(
        total_rmb=money(sum(r["amount_rmb"] for r in lines)),
        end=project_span(project)[1].isoformat(),
        captured_at=dt.datetime.now(dt.timezone.utc).isoformat(),
        name=name,
        timeline=[
            {k: r[k] for k in ("id", "start", "end", "task")}
            for r in project["timeline"]
            if r["enabled"]
        ],
        line_amounts={r["id"]: r["amount_rmb"] for r in lines},
    )


def checks(
    project: dict, lines: list[dict], cash: list[dict], warnings: list[str]
) -> list[dict]:
    result = []
    s = project["settings"]
    today = dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).date()

    def add(level, message):
        result.append(dict(level=level, message=message))

    by_id = {t["id"]: t for t in project["timeline"]}
    for t in project["timeline"]:
        if not t["enabled"]:
            continue
        pred = by_id.get(t["depends_on"])
        if pred and not pred["enabled"]:
            add("提示", f"{t['task']}：前置任务已禁用。")
        if pred and pred["enabled"] and as_date(t["start"]) < as_date(pred["end"]):
            add(
                "关注",
                f"{t['task']}开始早于前置任务“{pred['task']}”完成；如允许滚动开展，请移除该依赖。",
            )
        if as_date(t["end"]) < today and t["progress"] < 100 and not t["actual_end"]:
            add("关注", f"{t['task']}计划已到期，尚未登记完成。")
        if t["actual_end"] and (
            t["progress"] < 100 or as_date(t["actual_end"]) > today
        ):
            add("关注", f"{t['task']}实际完成日期与完成度或当前日期不一致。")
    enroll = next(
        (t for t in project["timeline"] if t["enabled"] and t["task"] == "受试者入组"),
        None,
    )
    if enroll:
        estimate = len(enrollment_curve(project))
        actual = len(month_starts(enroll["start"], enroll["end"]))
        if abs(estimate - actual) >= 2:
            add(
                "关注",
                f"入组速率模型估算{estimate}个月，当前排期覆盖{actual}个日历月；请复核中心启动节奏。",
            )
        lpi = as_date(enroll["end"])
        for task, months in [
            ("主要终点随访", s["primary_fu_months"]),
            ("末例受试者末次访视", s["final_fu_months"]),
        ]:
            t = next(
                (x for x in project["timeline"] if x["enabled"] and x["task"] == task),
                None,
            )
            if t and as_date(t["end"]) < add_months(lpi, months):
                add("关注", f"{task}结束早于末例入组后{months}个月。")
    if project["template"] == "示例项目":
        add(
            "提示",
            "示例数据仅用于演示；参考项MC2默认不计入总额。Site Cost已含随访费用，新增访视费前请核对是否重复。",
        )
    total = money(sum(x["amount_rmb"] for x in lines))
    cash_total = money(sum(x["amount_rmb"] for x in cash))
    if abs(total - cash_total) > 0.005:
        add("错误", f"预算与现金流不一致：差额RMB {total-cash_total:,.2f}。")
    actual = sum(
        to_rmb(e["amount"], e["currency"], s["fx"]) for e in project["expenses"]
    )
    if actual > total:
        add("关注", f"实际支出已超当前预算RMB {actual-total:,.2f}。")
    if s["approved_budget_rmb"] and total > s["approved_budget_rmb"]:
        add("关注", f"预计总预算超获批预算RMB {total-s['approved_budget_rmb']:,.2f}。")
    a, b = project_span(project)
    outside = sum(
        x["amount_rmb"]
        for x in cash
        if as_date(x["month"]) < a.replace(day=1)
        or as_date(x["month"]) > b.replace(day=1)
    )
    if outside:
        add(
            "提示",
            f"RMB {outside:,.2f}现金流在任务周期外，可能来自原始季度、访视偏移或付款滞后。",
        )
    enabled = {l["id"] for l in lines}
    if any(e["budget_id"] not in enabled for e in project["expenses"]):
        add("关注", "实际支出关联了已禁用费用，仍计入实际总支出。")
    for w in warnings:
        add("提示", w)
    return result


def safe_cell(value: Any) -> Any:
    # Keep free text inert when opened in Excel; negative numeric values remain numbers.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def safe_frame(rows):
    import pandas as pd

    return pd.DataFrame(rows).map(safe_cell)


def csv_bytes(rows, language="zh", table_name=None) -> bytes:
    return (
        export_frame(rows, language, table_name).to_csv(index=False).encode("utf-8-sig")
    )


def report_tables(project: dict) -> dict:
    s = project["settings"]
    lines = budget_lines(project)
    cash, warnings = allocate_cashflow(project, lines)
    a, b = project_span(project)
    total = money(sum(l["amount_rmb"] for l in lines))
    timeline = [
        dict(
            ID=t["id"],
            阶段=t["phase"],
            任务=t["task"],
            开始=t["start"],
            结束=t["end"],
            负责人=t["owner"],
            完成度=t["progress"],
            实际完成=t["actual_end"],
            前置ID=t["depends_on"],
        )
        for t in project["timeline"]
        if t["enabled"]
    ]
    budget = [
        dict(
            ID=l["id"],
            模块=l["module"],
            项目=l["item"],
            计费依据=l["basis"],
            原币种=l["currency"],
            原币单价=l["unit_price"],
            计算数量=l["calc_quantity"],
            原币金额=l["native_amount"],
            金额_RMB=l["amount_rmb"],
            金额_USD=display_amount(l["amount_rmb"], "USD", s["fx"]),
            备注=l["note"],
        )
        for l in lines
    ]
    product_logistics = money(
        sum(l["amount_rmb"] for l in lines if l["module"] == "产品物流")
    )
    primary_task = next(
        (
            t
            for t in project["timeline"]
            if t["enabled"] and t["task"] == "主要终点随访"
        ),
        None,
    )
    lead = leading_site(project)
    summary = [
        dict(指标=k, 内容=v)
        for k, v in [
            ("项目", s["project_name"]),
            ("研究类型", s["study_type"]),
            ("Leading site（牵头中心）", lead["name"] if lead else "未指定"),
            ("Leading PI（牵头PI）", lead["pi"] if lead and lead["pi"] else "未填写"),
            ("计划入组", s["n_planned"]),
            ("中心数", s["site_number"]),
            ("预计完成随访人数", s["n_planned"] - s["n_lost"]),
            ("主要终点随访月数", s["primary_fu_months"]),
            ("最终随访月数", s["final_fu_months"]),
            ("项目开始", a.isoformat()),
            ("项目结束", b.isoformat()),
            ("项目持续月数", round(duration_months(a, b), 2)),
            ("主要终点数据收集完成", primary_task["end"] if primary_task else "未设定"),
            ("计划访视数", len(project["visits"])),
            ("总预算RMB", total),
            ("总预算USD", display_amount(total, "USD", s["fx"])),
            ("产品物流预算RMB", product_logistics),
            ("不含产品物流预算RMB", money(total - product_logistics)),
            ("每例预算RMB", money(total / s["n_planned"])),
            (
                "每例预算不含产品物流RMB",
                money((total - product_logistics) / s["n_planned"]),
            ),
            ("每中心预算RMB", money(total / s["site_number"])),
            ("汇率USD兑RMB", s["fx"]),
            ("现金流方式", s["cash_mode"]),
            ("数据说明", "示例项目为预填演示数据；请按实际研究方案和报价更新。"),
        ]
    ]
    return {
        "项目摘要": summary,
        "预算明细": budget,
        "时间线": timeline,
        "访视设置": project["visits"],
        "月度现金流": cash,
        "季度汇总": quarter_totals(cash, project),
        "实际支出": project["expenses"],
        "检查提示": checks(project, lines, cash, warnings),
        **site_report_tables(project),
    }


def excel_bytes(tables: dict, language="zh") -> bytes:
    import pandas as pd

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, rows in tables.items():
            df = export_frame(rows, language, name)
            sheet_name = export_table_name(name, language)[:31]
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            sheet = writer.sheets[sheet_name]
            sheet.freeze_panes = "A2"
            if len(df):
                sheet.auto_filter.ref = sheet.dimensions
            from openpyxl.styles import Font, PatternFill, Alignment

            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="173767")
            for column in sheet.columns:
                width = min(
                    48,
                    max(
                        12, max(len(str(c.value or "")) for c in list(column)[:100]) + 2
                    ),
                )
                sheet.column_dimensions[column[0].column_letter].width = width
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top")
                    if isinstance(cell.value, float):
                        cell.number_format = "#,##0.00"
    return output.getvalue()


# ---------- V2.4: local, deterministic English table exports ----------
# Translate presentation copies only. Internal keys and project JSON stay stable.
EXPORT_SHEETS_EN = {
    "项目摘要": "Project Summary",
    "预算明细": "Budget Details",
    "时间线": "Timeline",
    "访视设置": "Visit Settings",
    "月度现金流": "Monthly Cash Flow",
    "季度汇总": "Quarterly Summary",
    "实际支出": "Actual Expenses",
    "检查提示": "Validation Checks",
    "中心信息与入组汇总": "Site Enrollment Summary",
    "FPI启动检查": "FPI Readiness Checklist",
    "入组累计快照": "Enrollment Snapshots",
    "入组曲线数据": "Enrollment Curve Data",
    "全中心月度入组": "Monthly Enrollment - All Sites",
}
# Schemas also give empty sheets/CSVs meaningful column headers.
EXPORT_TABLE_COLUMNS = {
    "项目摘要": "指标 内容".split(),
    "预算明细": "ID 模块 项目 计费依据 原币种 原币单价 计算数量 原币金额 金额_RMB 金额_USD 备注".split(),
    "时间线": "ID 阶段 任务 开始 结束 负责人 完成度 实际完成 前置ID".split(),
    "访视设置": "id enabled name offset times unit_price currency population".split(),
    "月度现金流": "month accrual_month id item module phase amount_rmb quarter".split(),
    "季度汇总": "quarter amount_rmb share".split(),
    "实际支出": "id date budget_id amount currency note".split(),
    "检查提示": "level message".split(),
    "中心信息与入组汇总": "site_id 中心编号 研究中心 PI PI联系方式 牵头中心 推进阶段 检查已完成 检查不适用 检查待确认 计划入组 实际累计 最新报告日期 同日计划 同日偏差 目标完成率 计划开始 计划结束 FPI已报告".split(),
    "FPI启动检查": "中心编号 研究中心 类别 检查项 状态".split(),
    "入组累计快照": "记录ID 中心编号 研究中心 报告日期 实际累计入组 备注".split(),
    "入组曲线数据": "series_id name date kind cumulative".split(),
    "全中心月度入组": "月份 统计截至 计划累计 实际累计 纳入中心数 已有累计报告中心数 本月更新中心数".split(),
}
EXPORT_HEADERS_EN = {
    "指标": "Metric",
    "内容": "Value",
    "ID": "ID",
    "id": "ID",
    "模块": "Module",
    "项目": "Item",
    "计费依据": "Quantity Basis",
    "原币种": "Original Currency",
    "原币单价": "Unit Price (Original Currency)",
    "计算数量": "Calculated Quantity",
    "原币金额": "Amount (Original Currency)",
    "金额_RMB": "Amount (RMB)",
    "金额_USD": "Amount (USD)",
    "备注": "Notes",
    "阶段": "Phase",
    "任务": "Task",
    "开始": "Start Date",
    "结束": "End Date",
    "负责人": "Owner",
    "完成度": "Progress (%)",
    "实际完成": "Actual Completion Date",
    "前置ID": "Predecessor ID",
    "enabled": "Enabled",
    "name": "Name",
    "offset": "Months After Enrollment",
    "times": "Visits per Participant",
    "unit_price": "Unit Price",
    "currency": "Currency",
    "population": "Population Basis",
    "month": "Payment Month",
    "accrual_month": "Accrual Month",
    "item": "Item",
    "module": "Module",
    "phase": "Phase",
    "amount_rmb": "Amount (RMB)",
    "quarter": "Fiscal Quarter",
    "share": "Share (%)",
    "date": "Date",
    "budget_id": "Budget Item ID",
    "amount": "Amount",
    "note": "Notes",
    "level": "Severity",
    "message": "Message",
    "site_id": "Site ID",
    "中心编号": "Site Code",
    "研究中心": "Site Name",
    "PI": "PI",
    "PI联系方式": "PI Contact",
    "牵头中心": "Leading Site",
    "推进阶段": "Startup Stage",
    "检查已完成": "Checks Completed",
    "检查不适用": "Checks Not Applicable",
    "检查待确认": "Checks Pending",
    "计划入组": "Planned Enrollment",
    "实际累计": "Actual Cumulative Enrollment",
    "最新报告日期": "Latest Report Date",
    "同日计划": "Planned as of Report Date",
    "同日偏差": "Actual minus Planned",
    "目标完成率": "Target Completion Rate",
    "计划开始": "Planned Start Date",
    "计划结束": "Planned End Date",
    "FPI已报告": "FPI Reported",
    "类别": "Category",
    "检查项": "Checklist Item",
    "状态": "Status",
    "记录ID": "Record ID",
    "报告日期": "Report Date",
    "实际累计入组": "Actual Cumulative Enrollment",
    "series_id": "Series ID",
    "kind": "Series Type",
    "cumulative": "Cumulative Enrollment",
    "月份": "Month",
    "统计截至": "As of Date",
    "计划累计": "Planned Cumulative Enrollment",
    "纳入中心数": "Included Sites",
    "已有累计报告中心数": "Sites with Cumulative Reports",
    "本月更新中心数": "Sites Updated This Month",
}
EXPORT_METRICS_EN = {
    "项目": "Project",
    "研究类型": "Study Type",
    "Leading site（牵头中心）": "Leading Site",
    "Leading PI（牵头PI）": "Leading PI",
    "计划入组": "Planned Enrollment",
    "中心数": "Number of Sites",
    "预计完成随访人数": "Expected Follow-up Completers",
    "主要终点随访月数": "Primary Endpoint Follow-up (Months)",
    "最终随访月数": "Final Follow-up (Months)",
    "项目开始": "Project Start Date",
    "项目结束": "Project End Date",
    "项目持续月数": "Project Duration (Months)",
    "主要终点数据收集完成": "Primary Endpoint Data Collection Complete",
    "计划访视数": "Number of Planned Visits",
    "总预算RMB": "Total Budget (RMB)",
    "总预算USD": "Total Budget (USD)",
    "产品物流预算RMB": "Product and Logistics Budget (RMB)",
    "不含产品物流预算RMB": "Budget Excluding Product and Logistics (RMB)",
    "每例预算RMB": "Budget per Participant (RMB)",
    "每例预算不含产品物流RMB": "Budget per Participant Excl. Product and Logistics (RMB)",
    "每中心预算RMB": "Budget per Site (RMB)",
    "汇率USD兑RMB": "Exchange Rate (RMB per USD)",
    "现金流方式": "Cash Flow Method",
    "数据说明": "Data Note",
}
EXPORT_VALUES_EN = {
    "研究中心": "Study Sites",
    "供应商": "Vendors",
    "产品物流": "Product and Logistics",
    "人力资源": "Human Resources",
    "风险储备": "Contingency Reserve",
    "固定数量": "Fixed Quantity",
    "入组例数": "Enrolled Participants",
    "完成随访例数": "Follow-up Completers",
    "筛选失败例数": "Screen Failures",
    "筛选总例数": "Total Screened",
    "中心数": "Number of Sites",
    "项目月数": "Project Months",
    "FTE项目": "FTE per Project",
    "FTE月": "FTE-Months",
    "FTE年": "FTE-Years",
    "智能匹配时间线": "Rule-based Timeline Allocation",
    "示例项目预填季度": "Example Project Quarterly Profile",
    "自定义季度占比": "Custom Quarterly Weights",
    "全周期按月均摊": "Even Monthly Allocation",
    "全项目周期": "Full Project Period",
    "全项目周期（缺少映射阶段）": "Full Project Period (Mapped Phase Unavailable)",
    "智能规则": "Rule-based Allocation",
    "智能匹配": "Automatic Matching",
    "阶段均摊": "Even Allocation within Phase",
    "阶段开始一次性": "Lump Sum at Phase Start",
    "阶段结束一次性": "Lump Sum at Phase End",
    "上市前临床研究": "Premarket Clinical Study",
    "上市后临床研究": "Postmarket Clinical Study",
    "回顾性分析": "Retrospective Analysis",
    "其他": "Other",
    "示例项目": "Example Project",
    "新建临床研究项目": "New Clinical Study Project",
    "示例项目（预填演示）": "Example Project (Prefilled Demo)",
    "未指定": "Not Assigned",
    "未填写": "Not Entered",
    "未设定": "Not Set",
    "示例项目为预填演示数据；请按实际研究方案和报价更新。": "The example project contains prefilled demonstration data. Update it using the actual study protocol and quotations.",
    "研究准备阶段": "Study Preparation",
    "研究文件准备": "Study Documentation",
    "政府部门备案": "Regulatory Filings",
    "组长单位GCP&伦理递交": "Lead Site GCP and Ethics Submission",
    "组长单位合同签署": "Lead Site Contract Execution",
    "研究中心启动": "Site Initiation",
    "受试者入组": "Participant Enrollment",
    "受试者随访": "Participant Follow-up",
    "数据分析": "Data Analysis",
    "研究结束": "Study Closeout",
    "项目立项": "Project Initiation",
    "研究中心名单提名、筛选及访视": "Site Nomination, Selection and Visits",
    "产品及人员培训": "Product and Staff Training",
    "供应商筛选与确认": "Vendor Selection and Confirmation",
    "产品运输及相关管理流程准备": "Product Shipping and Logistics Preparation",
    "研究者会议": "Investigator Meeting",
    "研究方案撰写及获批": "Protocol Development and Approval",
    "CRF设计及数据库建立": "CRF Design and Database Setup",
    "知情同意书撰写及获批": "Informed Consent Form Development and Approval",
    "研究计划书撰写及获批（监查计划、数据管理计划、统计分析计划等）": "Study Plan Development and Approval (Monitoring, Data Management, Statistical Analysis, etc.)",
    "其他相关表格及物料准备（受试者日记卡、招募海报、标签等）": "Other Forms and Materials (Participant Diaries, Recruitment Posters, Labels, etc.)",
    "文件翻译及定稿": "Document Translation and Finalization",
    "上海药监局备案": "Shanghai Drug Administration Filing",
    "人类遗传办提交": "Human Genetic Resources Submission",
    "人类遗传办获批": "Human Genetic Resources Approval",
    "组长单位GCP资料递交": "Lead Site GCP Document Submission",
    "组长单位EC申请及获批": "Lead Site Ethics Submission and Approval",
    "其他中心EC申请及获批": "Other Sites Ethics Submission and Approval",
    "组长单位合同起草与协商": "Lead Site Contract Drafting and Negotiation",
    "组长单位合同签字与盖章": "Lead Site Contract Signing and Sealing",
    "组长单位启动": "Lead Site Initiation",
    "其他中心启动": "Other Sites Initiation",
    "组长单位首例入组": "Lead Site First Participant In",
    "其他中心首例入组": "Other Sites First Participant In",
    "末例入组": "Last Participant In",
    "主要终点随访": "Primary Endpoint Follow-up",
    "随访结束": "Follow-up Completion",
    "末例受试者末次访视": "Last Participant Last Visit",
    "主要终点数据库锁定": "Primary Endpoint Database Lock",
    "主要终点TLG/统计报告撰写": "Primary Endpoint TLG and Statistical Report Preparation",
    "主要终点临床研究报告撰写": "Primary Endpoint Clinical Study Report Preparation",
    "主要终点数据的监管提交": "Primary Endpoint Regulatory Submission",
    "数据库清理与锁定": "Database Cleaning and Lock",
    "最终图表TLG/统计报告完成": "Final TLG and Statistical Report Completion",
    "临床研究报告撰写": "Clinical Study Report Preparation",
    "中心关闭访视": "Site Closeout Visits",
    "临床研究报告的监管递交": "Clinical Study Report Regulatory Submission",
    "研究者筛选费": "Investigator Screening Fees",
    "研究者手术费": "Investigator Procedure Fees",
    "筛选失败费用": "Screen Failure Costs",
    "受试者补偿": "Participant Compensation",
    "受试者检查费用": "Participant Assessment Costs",
    "组长单位费用": "Lead Site Fees",
    "EC/Clinical Institute Process费用": "Ethics Committee and Institution Processing Fees",
    "CRC费用": "CRC Fees",
    "PM费用": "Project Management Fees",
    "Monitor费用": "Monitoring Fees",
    "DM费用": "Data Management Fees",
    "Safety费用": "Safety Services Fees",
    "EDC系统费用": "EDC System Fees",
    "翻译/打印费用": "Translation and Printing Costs",
    "保险费用": "Insurance Costs",
    "数据分析费用": "Data Analysis Fees",
    "中心实验室费用": "Central Laboratory Fees",
    "CEC费用": "CEC Fees",
    "Travel费用": "Travel Costs",
    "Recording费用": "Recording Costs",
    "研究者会费用": "Investigator Meeting Costs",
    "产品运输费用": "Product Shipping Costs",
    "仓储费用": "Storage Costs",
    "产品成本": "Product Costs",
    "项目经理FTE费用": "Project Manager FTE Costs",
    "CRA费用": "CRA Costs",
    "MCRS费用": "MCRS Costs",
    "Global team费用": "Global Team Costs",
    "中心管理费用": "Site Management Fees",
    "Site税费": "Site Taxes",
    "Vendor税费": "Vendor Taxes",
    "风险储备（未承诺）": "Uncommitted Contingency Reserve",
    "出院": "Discharge",
    "30天": "Day 30",
    "1个月": "Month 1",
    "6个月": "Month 6",
    "12个月": "Month 12",
    "1年": "Year 1",
    "2年": "Year 2",
    "3年": "Year 3",
    "未开始": "Not Started",
    "待确认": "Pending Confirmation",
    "已完成": "Completed",
    "不适用": "Not Applicable",
    "是": "Yes",
    "否": "No",
    "提示": "Info",
    "关注": "Attention",
    "错误": "Error",
    "研究合规": "Study Compliance",
    "合同与财务": "Contracts and Finance",
    "器械与临床资源": "Devices and Clinical Resources",
    "人员与数据": "Personnel and Data",
    "招募与中心激活": "Recruitment and Site Activation",
    "研究类型及监管路径已确认": "Study Type and Regulatory Pathway Confirmed",
    "机构/科研管理审批完成": "Institutional and Research Administration Approvals Completed",
    "本中心伦理批准有效": "Valid Site Ethics Approval in Place",
    "适用的备案、登记、专项许可已完成": "Applicable Filings, Registrations and Special Permits Completed",
    "CTA及必要服务协议生效": "CTA and Required Service Agreements Effective",
    "经费支付及受试者补偿流程就绪": "Payment and Participant Compensation Processes Ready",
    "受试者损害处理和保障安排明确": "Participant Injury Management and Protection Arrangements Confirmed",
    "试验器械与耗材可用且验收合格": "Study Devices and Consumables Available and Accepted",
    "术者、导管室、配套设备就绪": "Operators, Catheterization Laboratory and Supporting Equipment Ready",
    "围手术期和应急预案确认": "Perioperative and Emergency Plans Confirmed",
    "PI及研究团队完成授权与培训": "PI and Study Team Delegation and Training Completed",
    "EDC、CTMS及相关系统可用": "EDC, CTMS and Related Systems Available",
    "源文件、ISF及安全性报告流程就绪": "Source Documents, ISF and Safety Reporting Processes Ready",
    "合规的受试者招募计划已建立": "Compliant Participant Recruitment Plan Established",
    "SIV遗留关键问题全部关闭": "All Critical SIV Action Items Closed",
    "已发出正式中心激活确认": "Formal Site Activation Confirmation Issued",
    "计划入组": "Planned Enrollment",
    "实际入组": "Actual Enrollment",
    "全研究合计": "All Sites Total",
    "示例数据仅用于演示；参考项MC2默认不计入总额。Site Cost已含随访费用，新增访视费前请核对是否重复。": "Example data is for demonstration only. MC2 is excluded from the total by default. Site Cost includes follow-up fees; check for duplication before adding visit fees.",
    "实际支出关联了已禁用费用，仍计入实际总支出。": "Actual expenses linked to disabled budget items remain included in total actual expenditure.",
}


def english_system_value(value):
    if not isinstance(value, str):
        return value
    if value in EXPORT_VALUES_EN:
        return EXPORT_VALUES_EN[value]
    if value.startswith("研究者随访费 - "):
        return "Investigator Follow-up Fee - " + english_system_value(
            value.split(" - ", 1)[1]
        )
    match = re.fullmatch(r"入组后([\d.]+)月访视", value)
    if match:
        return f"Visit {match[1]} Month(s) After Enrollment"
    return value  # Custom task/item names are preserved, never guessed.


def english_system_message(value):
    """Translate only known generated messages; captured custom names stay verbatim."""
    if not isinstance(value, str) or value in EXPORT_VALUES_EN:
        return english_system_value(value)
    patterns = [
        (r"(.+)：前置任务已禁用。", "{0}: the predecessor task is disabled."),
        (
            r"(.+)开始早于前置任务“(.+)”完成；如允许滚动开展，请移除该依赖。",
            "{0} starts before predecessor '{1}' finishes. Remove the dependency if overlapping work is intended.",
        ),
        (
            r"(.+)计划已到期，尚未登记完成。",
            "{0} is past its planned end date and has not been marked complete.",
        ),
        (
            r"(.+)实际完成日期与完成度或当前日期不一致。",
            "{0}: the actual completion date conflicts with progress or today's date.",
        ),
        (
            r"入组速率模型估算(\d+)个月，当前排期覆盖(\d+)个日历月；请复核中心启动节奏。",
            "The enrollment model estimates {0} months; the schedule covers {1} calendar months. Review site activation timing.",
        ),
        (
            r"(.+)结束早于末例入组后([\d.]+)个月。",
            "{0} ends earlier than {1} months after last participant enrollment.",
        ),
        (
            r"预算与现金流不一致：差额RMB ([\d,.-]+)。",
            "Budget and cash flow differ by RMB {0}.",
        ),
        (
            r"实际支出已超当前预算RMB ([\d,.]+)。",
            "Actual expenditure exceeds the current budget by RMB {0}.",
        ),
        (
            r"预计总预算超获批预算RMB ([\d,.]+)。",
            "The estimated budget exceeds the approved budget by RMB {0}.",
        ),
        (
            r"RMB ([\d,.]+)现金流在任务周期外，可能来自原始季度、访视偏移或付款滞后。",
            "RMB {0} of cash flow falls outside the task period, possibly due to the original quarterly profile, visit offsets or payment lag.",
        ),
        (
            r"(.+)：指定阶段未启用，已按全周期摊分。",
            "{0}: the specified phase is disabled; allocation uses the full project period.",
        ),
        (
            r"(.+)：未启用受试者入组任务，访视费用暂按全项目入组窗口推算。",
            "{0}: the enrollment task is disabled; visit costs use the project-wide enrollment window.",
        ),
        (
            r"(.+)：无预设匹配规则，已按全周期摊分；可指定阶段。",
            "{0}: no predefined allocation rule; allocation uses the full project period. A phase can be assigned.",
        ),
        (
            r"(.+)：(.+)未启用，相关金额按全周期摊分。",
            "{0}: {1} is disabled; the related amount is allocated across the full project period.",
        ),
        (
            r"(.+)：原始矩阵无此项目，已使用智能规则。",
            "{0}: this item is absent from the original matrix; rule-based allocation is used.",
        ),
    ]
    for pattern, template in patterns:
        match = re.fullmatch(pattern, value)
        if match:
            return template.format(*(english_system_value(g) for g in match.groups()))
    return value


def export_table_name(name, language="zh"):
    if language not in ("zh", "en"):
        raise ValueError("导出语言仅支持中文或English。")
    return EXPORT_SHEETS_EN.get(name, name) if language == "en" else name


def translate_export_rows(rows, language="zh", table_name=None):
    """Return a separate display copy; names, notes, numbers, IDs and dates survive."""
    export_table_name(table_name, language)  # Validate even for an empty table.
    if language == "zh":
        return copy.deepcopy(rows)
    value_columns = {
        "预算明细": {"模块", "项目", "计费依据"},
        "时间线": {"阶段", "任务"},
        "访视设置": {"name", "population"},
        "月度现金流": {"item", "module", "phase"},
        "检查提示": {"level"},
        "中心信息与入组汇总": {"牵头中心", "推进阶段", "FPI已报告"},
        "FPI启动检查": {"类别", "检查项", "状态"},
        "入组曲线数据": {"kind"},
    }.get(table_name, set())
    translated = []
    for row in rows:
        result = {}
        for key, value in row.items():
            if table_name == "项目摘要":
                metric = row.get("指标")
                if key == "指标":
                    value = EXPORT_METRICS_EN.get(value, value)
                elif key == "内容" and metric in {
                    "项目",
                    "研究类型",
                    "现金流方式",
                    "数据说明",
                    "主要终点数据收集完成",
                    "Leading site（牵头中心）",
                    "Leading PI（牵头PI）",
                }:
                    # Only sentinel values are translated in the two leading-name fields.
                    leading = metric.startswith("Leading")
                    if metric == "项目":
                        if value in (
                            "示例项目",
                            "示例项目（预填演示）",
                            "新建临床研究项目",
                        ):
                            value = english_system_value(value)
                    elif not leading or value in ("未指定", "未填写"):
                        value = english_system_value(value)
            elif key in value_columns:
                value = english_system_value(value)
            elif table_name == "检查提示" and key == "message":
                value = english_system_message(value)
            elif (
                table_name == "预算明细"
                and key == "备注"
                and row.get("ID")
                in {
                    "D-MGMT",
                    "D-SITETAX",
                    "D-VENDORTAX",
                    "D-RESERVE",
                }
            ):
                value = value.replace("计费基数 RMB ", "Calculation Base: RMB ", 1)
            elif (
                table_name == "入组曲线数据"
                and key == "name"
                and row.get("series_id") == "ALL"
            ):
                value = "All Sites Total"
            result[EXPORT_HEADERS_EN.get(key, key)] = copy.deepcopy(value)
        translated.append(result)
    return translated


def export_frame(rows, language="zh", table_name=None):
    translated = translate_export_rows(rows, language, table_name)
    frame = safe_frame(translated)
    if not translated and table_name in EXPORT_TABLE_COLUMNS:
        columns = EXPORT_TABLE_COLUMNS[table_name]
        if language == "en":
            columns = [EXPORT_HEADERS_EN.get(c, c) for c in columns]
        frame = frame.reindex(columns=columns)
    return frame


def current_export_language():
    return "en" if st.session_state.get("table_export_language") == "English" else "zh"


def remember_export_language():
    # Retain the preference when Streamlit cleans up the export-page widget.
    st.session_state.table_export_language = st.session_state.export_language


def export_csv_filename(stem):
    return stem + ("_EN" if current_export_language() == "en" else "") + ".csv"


# ---------- Streamlit interface (lazy imports keep engine testable) ----------
CSS = """
<style>
:root {--navy:#122D56;--ink:#152D4D;--teal:#168F89;--muted:#6B7D96;}
.stApp {background:#F5F7FB;}
.block-container {max-width:1460px;padding:2rem 2.6rem 3rem;}
[data-testid="stSidebar"] {background:#FFFFFF;border-right:1px solid #E3E9F2;}
[data-testid="stSidebar"] .block-container {padding-top:1.2rem;}
h1,h2,h3 {letter-spacing:-.035em;color:var(--ink);}
h1 {font-size:2rem !important;} h2 {font-size:1.4rem !important;} h3 {font-size:1.12rem !important;}
[data-testid="stMetric"] {background:white;border:1px solid #E3E9F2;border-radius:14px;padding:18px 20px;min-height:115px;}
[data-testid="stMetricValue"] {font-size:1.65rem;letter-spacing:-.04em;color:var(--ink);}
[data-testid="stMetricLabel"] {color:var(--muted);}
[data-testid="stForm"] {background:white;border:1px solid #E3E9F2;border-radius:16px;padding:1.35rem;}
[data-testid="stDataFrame"] {border-radius:12px;overflow:hidden;border:1px solid #E3E9F2;}
[data-testid="stExpander"] {background:white;border-radius:12px;}
.stButton>button,.stDownloadButton>button {border-radius:9px;min-height:2.65rem;}
.stButton>button[kind="primary"],.stFormSubmitButton>button[kind="primary"] {background:#2357D9;border:0;}
.hero {background:linear-gradient(110deg,#122D56 0%,#1B4474 70%,#176E78 100%);border-radius:19px;padding:26px 30px;margin-bottom:22px;color:white;}
.hero .eyebrow {font-size:11px;letter-spacing:.18em;color:#9FC7EF;font-weight:700;margin-bottom:10px;}
.hero h1 {color:white;margin:0;padding:0;font-size:27px !important;line-height:1.45;overflow-wrap:anywhere;}
.hero p {color:#BED1E7;margin:9px 0 0;font-size:13px;line-height:1.8;}
.brand {font-size:27px;font-weight:800;letter-spacing:-1px;color:#142F56;line-height:1.1;margin-top:12px;}
.brand span {color:#168F89;}
.section-note {padding:12px 16px;border-left:3px solid #168F89;background:#EAF5F5;border-radius:0 8px 8px 0;color:#34516B;font-size:13px;margin:10px 0 18px;}
@media(max-width:800px){.block-container{padding:1.2rem 1rem;}.hero{padding:20px;}.hero h1{font-size:23px !important;}}

/* Inputs remain native controls; only editable surfaces get a blue outline. */

/* Stable test IDs cover the current React Aria controls. BaseWeb selectors below
   also support earlier compatible Streamlit control rendering. */
[data-testid="stTextInputRootElement"]:not(:has(input:disabled)),
[data-testid="stNumberInputContainer"]:not(:has(input:disabled)),
[data-testid="stTextAreaRootElement"]:not(:has(textarea:disabled)),
[data-testid="stDateInputField"]:not([data-disabled]),
[data-testid="stSelectbox"] button[aria-haspopup="listbox"]:not(:disabled),
[data-testid="stMultiSelect"] [role="combobox"]:not(:disabled) {
    background:#FFFFFF !important;border:1px solid #8EBBE8 !important;border-radius:8px !important;
}
[data-testid="stTextInputRootElement"]:has(input:disabled),
[data-testid="stNumberInputContainer"]:has(input:disabled),
[data-testid="stTextAreaRootElement"]:has(textarea:disabled),
[data-testid="stDateInputField"][data-disabled] {
    background:#DFE6EF !important;border:1px solid #BECAD9 !important;color:#52647B !important;
}

[data-testid="stTextInput"] input:not(:disabled),
[data-testid="stNumberInput"] input:not(:disabled),
[data-testid="stTextArea"] textarea:not(:disabled),
[data-testid="stDateInput"] input:not(:disabled) {
    background:#FFFFFF !important; color:#17365C !important;
}
[data-testid="stTextInput"] [data-baseweb="input"]:not(:has(input:disabled)),
[data-testid="stNumberInput"] [data-baseweb="input"]:not(:has(input:disabled)),
[data-testid="stTextArea"] [data-baseweb="textarea"]:not(:has(textarea:disabled)),
[data-testid="stDateInput"] [data-baseweb="input"]:not(:has(input:disabled)),
[data-testid="stSelectbox"] [data-baseweb="select"] > div:first-child,
[data-testid="stMultiSelect"] [data-baseweb="select"] > div:first-child {
    background:#FFFFFF !important; border:1px solid #8EBBE8 !important; border-radius:8px !important;
}
[data-testid="stSelectbox"] [aria-disabled="true"] > div:first-child,
[data-testid="stMultiSelect"] [aria-disabled="true"] > div:first-child,
input:disabled, textarea:disabled,
[data-baseweb="input"]:has(input:disabled) {
    background:#DFE6EF !important; color:#52647B !important; border-color:#BECAD9 !important;
}
[data-testid="stMetric"] {background:#EDF1F6; border-color:#D6DFEA;}
.field-legend {display:flex;flex-wrap:wrap;gap:12px 28px;align-items:center;margin:0 0 20px;padding:12px 16px;
    background:#FFFFFF;border:1px solid #DFE6F0;border-radius:10px;color:#536881;font-size:12px;}
.field-legend span {display:flex;align-items:center;gap:7px;flex-wrap:wrap;}
.field-legend i {display:inline-block;width:18px;height:18px;border-radius:4px;flex-shrink:0;}
.editable-swatch {background:#FFFFFF;border:2px solid #8EBBE8;}
.readonly-swatch {background:#DFE6EF;border:1px solid #BECAD9;}

</style>
"""


# Table colors are applied through supported pandas Styler metadata: CSS cannot
# recolor individual canvas cells. Keep data names/types unchanged for persistence.
EDIT_FIELD_BG = "#FFFFFF"
READONLY_FIELD_BG = "#DFE6EF"


def editable_table(data, **kwargs):
    frame = data.copy() if isinstance(data, pd.DataFrame) else pd.DataFrame(data)
    config = copy.deepcopy(kwargs.pop("column_config", {}))
    disabled = kwargs.get("disabled", False)
    readonly = []
    for column in frame.columns:
        if config.get(column, "absent") is None:
            continue
        cfg = config.get(column, {})
        if not isinstance(cfg, dict):
            cfg = {"label": str(cfg)}
        locked = (
            disabled is True
            or (isinstance(disabled, (list, tuple, set)) and column in disabled)
            or cfg.get("disabled", False)
        )
        if locked:
            readonly.append(column)
        label = cfg.get("label") or str(column)
        cfg["label"] = label + (" 🔒" if locked else " ✎")
        hint = (
            "只读：由系统生成或从其他页面关联。"
            if locked
            else "可填写：单击单元格修改，完成后点击本表保存按钮。"
        )
        cfg["help"] = hint + (" " + cfg["help"] if cfg.get("help") else "")
        config[column] = cfg
    style = frame.style.set_uuid("crms_editor")
    if readonly:
        style = style.set_properties(
            subset=readonly,
            **{"background-color": READONLY_FIELD_BG, "color": "#42536B"},
        )
    return st.data_editor(style, column_config=config, **kwargs)


def readonly_table(data, **kwargs):
    frame = data.copy() if isinstance(data, pd.DataFrame) else pd.DataFrame(data)
    style = frame.style.set_uuid("crms_readonly").set_properties(
        **{"background-color": READONLY_FIELD_BG, "color": "#42536B"}
    )
    return st.dataframe(style, **kwargs)


def render_field_legend():
    st.markdown(
        '<div class="field-legend" role="note">'
        '<span><i class="editable-swatch"></i><b>✎ 可填写 / 修改</b>：蓝框输入区、表格浅色单元格</span>'
        '<span><i class="readonly-swatch"></i><b>🔒 只读 / 自动计算</b>：灰底字段</span>'
        "</div>",
        unsafe_allow_html=True,
    )


def apply_change(project: dict, message: str, *, replace_project=False) -> None:
    errors = validate_project(project)
    if not errors:
        try:
            enrollment_curve(project)
        except ValueError as exc:
            errors.append(str(exc))
    if errors:
        for error in errors[:12]:
            st.error(error)
        return
    st.session_state.active_page = st.session_state.get("navigation", "项目概览")
    if replace_project:
        st.session_state.ui_generation = st.session_state.get("ui_generation", 0) + 1
    st.session_state.project = project
    st.session_state.revision += 1
    st.session_state.notice = message
    st.rerun()


def chart_style(fig, height=360):
    fig.update_layout(
        template="plotly_white",
        height=height,
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        margin=dict(l=10, r=20, t=35, b=20),
        font=dict(
            family="Arial, Microsoft YaHei, sans-serif", color="#314866", size=12
        ),
        colorway=COLORS,
        legend=dict(orientation="h", y=-0.2),
        title_font_size=15,
    )
    return fig


def metric_values(project, lines):
    s = project["settings"]
    total = money(sum(l["amount_rmb"] for l in lines))
    return total, lambda x: f"{display_amount(x,s['display_currency'],s['fx']):,.0f}"


def draw_timeline(project, detail=False, overlay=False):
    rows = [
        dict(
            阶段=t["phase"],
            任务=t["task"],
            开始=as_date(t["start"]),
            结束=as_date(t["end"]) + dt.timedelta(days=1),
            原结束=t["end"],
            完成度=t["progress"],
        )
        for t in project["timeline"]
        if t["enabled"]
    ]
    if not rows:
        st.info("请启用时间线任务。")
        return
    df = pd.DataFrame(rows)
    y = "任务" if detail else "阶段"
    if not detail:
        df = df.groupby("阶段", sort=False, as_index=False).agg(
            开始=("开始", "min"), 结束=("结束", "max")
        )
        df["原结束"] = (pd.to_datetime(df["结束"]) - pd.Timedelta(days=1)).dt.strftime(
            "%Y-%m-%d"
        )
    df["开始"] = pd.to_datetime(df["开始"])
    df["结束"] = pd.to_datetime(df["结束"])
    fig = px.timeline(
        df,
        x_start="开始",
        x_end="结束",
        y=y,
        color="阶段",
        color_discrete_sequence=COLORS,
        hover_data={"原结束": True, "结束": False},
        category_orders={y: df[y].tolist()},
    )
    fig.update_yaxes(
        categoryorder="array",
        categoryarray=df[y].tolist(),
        autorange="reversed",
        title=None,
    )
    fig.update_xaxes(title=None, gridcolor="#E5ECF5")
    today = dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).date().isoformat()
    fig.add_shape(
        type="line",
        x0=today,
        x1=today,
        y0=0,
        y1=1,
        yref="paper",
        line=dict(color="#1A988D", width=1.5, dash="dot"),
    )
    fig.add_annotation(
        x=today,
        y=1.04,
        yref="paper",
        text="今天",
        showarrow=False,
        font=dict(color="#168F89"),
    )
    if overlay and project["baseline"] and detail:
        current = {t["id"] for t in project["timeline"] if t["enabled"]}
        baseline = [t for t in project["baseline"]["timeline"] if t["id"] in current]
        if baseline:
            bdf = pd.DataFrame(baseline)
            bdf["start"] = pd.to_datetime(bdf["start"])
            bdf["end"] = pd.to_datetime(bdf["end"]) + pd.Timedelta(days=1)
            bfig = px.timeline(bdf, x_start="start", x_end="end", y="task")
            for trace in bfig.data:
                trace.update(
                    marker_color="rgba(130,145,165,.32)",
                    name="保存的基线",
                    showlegend=True,
                    hovertemplate="基线：%{y}<extra></extra>",
                    width=0.25,
                )
                fig.add_trace(trace)
    chart_style(fig, max(370, len(df) * 30 + 110) if detail else 410)
    fig.update_layout(showlegend=detail, barmode="overlay")
    st.plotly_chart(fig, width="stretch")
    st.caption("日期包含起止当日；同日里程碑显示为1天。竖线为北京时间今天。")


def sidebar(project):
    s = project["settings"]
    with st.sidebar:
        st.markdown(
            '<div class="brand">CRMS<span> / </span>PLANNER</div>',
            unsafe_allow_html=True,
        )
        st.caption("临床研究 · 预算与时间线工作台")
        st.divider()
        pages = [
            "项目概览",
            "项目设置",
            "预算编制",
            "时间线管理",
            "中心启动进度",
            "入组进度",
            "现金流计划",
            "执行跟踪",
            "情景测算",
            "导出与帮助",
        ]
        selected_page = st.session_state.get("active_page", "项目概览")
        page = st.radio(
            "工作空间",
            [
                "项目概览",
                "项目设置",
                "预算编制",
                "时间线管理",
                "中心启动进度",
                "入组进度",
                "现金流计划",
                "执行跟踪",
                "情景测算",
                "导出与帮助",
            ],
            label_visibility="collapsed",
            key="navigation",
            index=pages.index(selected_page) if selected_page in pages else 0,
            on_change=remember_navigation,
        )
        st.session_state.active_page = page
        st.divider()
        st.caption("当前项目")
        st.write(s["project_name"])
        st.caption(
            f"{project['template']}模板 · 本会话第 {st.session_state.revision} 次保存"
        )
        st.download_button(
            "保存项目 JSON",
            project_bytes(project),
            file_name="CRMS_project.json",
            mime="application/json",
            width="stretch",
            on_click="ignore",
        )
        st.caption(
            "表格导出："
            + ("English" if current_export_language() == "en" else "中文")
            + "（在“导出与帮助”切换）"
        )
        with stable_expander("打开 / 新建项目"):
            st.caption("先导出当前项目；打开或新建会替换当前会话。")
            upload = st.file_uploader(
                "打开V2项目JSON", type=["json"], key="import_file"
            )
            if st.button("载入文件", disabled=upload is None, width="stretch"):
                try:
                    candidate = load_project(upload.getvalue())
                except (
                    ValueError,
                    TypeError,
                    KeyError,
                    OverflowError,
                    RecursionError,
                ) as exc:
                    st.error(str(exc))
                else:
                    apply_change(candidate, "项目已载入。", replace_project=True)
            template = st.selectbox("新建项目模板", ["示例项目", "通用"])
            if st.button("新建并载入模板", width="stretch"):
                apply_change(
                    make_project(template), "已载入模板。", replace_project=True
                )
        st.caption("表单修改后请点击保存。刷新浏览器或服务重启前，请导出项目JSON备份。")
        st.caption(f"CRMS {APP_VERSION} · 不需要外部AI API")
    return page


def render_overview(p, lines, cash, warnings):
    s = p["settings"]
    a, b = project_span(p)
    total, fmt = metric_values(p, lines)
    actual = money(
        sum(to_rmb(e["amount"], e["currency"], s["fx"]) for e in p["expenses"])
    )
    lead = leading_site(p)
    lead_cols = st.columns(2)
    with lead_cols[0]:
        with st.container(border=True):
            st.caption("Leading site · 牵头中心")
            st.write(
                lead["name"] if lead else "尚未指定，请在中心启动进度中勾选牵头中心"
            )
    with lead_cols[1]:
        with st.container(border=True):
            st.caption("Leading PI · 牵头PI")
            st.write(
                lead["pi"] if lead and lead["pi"] else "待填写（自动关联牵头中心的PI）"
            )
    c = st.columns(4)
    c[0].metric("预计总预算 · " + s["display_currency"], fmt(total))
    c[1].metric("每例预算 · " + s["display_currency"], fmt(total / s["n_planned"]))
    c[2].metric("项目持续时间", f"{duration_months(a,b):.1f} 月")
    c[3].metric("实际支出占预算", f"{actual/total*100:.1f}%" if total else "—")
    st.caption(
        f"计划入组 {s['n_planned']} 例  ·  {s['site_number']} 家中心  ·  {a:%Y.%m.%d} — {b:%Y.%m.%d}"
    )
    render_overview_enrollment(p)
    left, right = st.columns([1.6, 1], gap="large")
    with left:
        st.subheader("项目阶段总览")
        draw_timeline(p)
    with right:
        st.subheader("预算构成")
        df = pd.DataFrame(lines).groupby("module", as_index=False)["amount_rmb"].sum()
        df = df[df.amount_rmb > 0]
        if not df.empty:
            fig = px.pie(
                df,
                values="amount_rmb",
                names="module",
                hole=0.72,
                color_discrete_sequence=COLORS,
            )
            fig.update_traces(textposition="outside", textinfo="label+percent")
            chart_style(fig, 355)
            fig.update_layout(showlegend=False)
            st.plotly_chart(fig, width="stretch")
        else:
            st.info("尚无有效预算。")
        st.caption("风险储备默认关闭，启用后单独列示。")
    st.subheader("当前关注事项")
    findings = checks(p, lines, cash, warnings)
    important = [x for x in findings if x["level"] in ("错误", "关注")]
    if important:
        for x in important[:4]:
            st.warning(x["message"])
        if len(important) > 4:
            st.caption(f"另有{len(important)-4}条关注事项，可在“执行跟踪”查看。")
    else:
        st.success("当前未检出需关注的关联问题。")
    st.markdown(
        '<div class="section-note">建议顺序：项目设置 → 预算编制 → 时间线管理 → 现金流计划。执行中记录实际支出，并保存基线用于比较变更。</div>',
        unsafe_allow_html=True,
    )


def render_settings(p):
    s = p["settings"]
    rev = form_epoch("settings", [s, p["timeline_mode"]])
    st.subheader("项目设置")
    st.caption("按研究规模、随访和入组效率生成排期；所有单价按原币种保存。")
    with st.form(f"settings_{rev}"):
        c1, c2 = st.columns(2, gap="large")
        with c1:
            name = st.text_input("项目名称", s["project_name"], max_chars=120)
            types = ["上市前临床研究", "上市后临床研究", "ERP", "回顾性分析", "其他"]
            study = st.selectbox(
                "研究类型",
                types,
                index=types.index(s["study_type"]) if s["study_type"] in types else 4,
            )
            n = st.number_input("计划入组人数", 1, 100000, int(s["n_planned"]))
            sites = st.number_input("计划中心数", 1, 100000, int(s["site_number"]))
            lost = st.number_input(
                "预估失访人数",
                0,
                100000,
                int(s["n_lost"]),
                help="仅影响计费依据为“完成随访例数”的费用；不自动扣减手术、筛选和已发生的随访费用。",
            )
            failed = st.number_input(
                "预估筛选失败人数", 0, 100000, int(s["n_screen_fail"])
            )
            start = st.date_input(
                "研究基准日期（示例项目为CIP批准日）",
                as_date(s["start_date"]),
                min_value=dt.date(2000, 1, 1),
                max_value=dt.date(2099, 12, 31),
            )
        with c2:
            primary = st.number_input(
                "主要终点：每例随访月数",
                0,
                240,
                int(s["primary_fu_months"]),
                help="例如12个月随访终点。末例主要终点完成日期＝末例入组日期＋12个月。",
            )
            final = st.number_input(
                "最终随访：每例随访月数", 0, 240, int(s["final_fu_months"])
            )
            rate = st.number_input(
                "每中心月入组速率（例/月）",
                0.01,
                10000.0,
                float(s["enrollment_rate"]),
                step=0.1,
                format="%.2f",
            )
            startup = st.number_input(
                "从基准日到首例入组（月）", 1, 120, int(s["startup_months"])
            )
            ramp = st.number_input(
                "中心逐步启动周期（月）",
                1,
                120,
                int(s["ramp_months"]),
                help="按月线性增加活跃中心；这是期望值模型，不是中心级排期。",
            )
            cleaning = st.number_input(
                "末次访视后数据清理（月）", 1, 120, int(s["cleaning_months"])
            )
            report = st.number_input(
                "锁库后报告周期（月）", 1, 120, int(s["report_months"])
            )
        st.divider()
        c1, c2, c3 = st.columns(3)
        currency = c1.selectbox(
            "展示币种",
            ["RMB", "USD"],
            index=["RMB", "USD"].index(s["display_currency"]),
        )
        fx = c2.number_input(
            "汇率：1 USD = RMB", 0.01, 1000.0, float(s["fx"]), step=0.01, format="%.4f"
        )
        fy = c3.number_input(
            "项目财年起始月",
            1,
            12,
            int(s["fy_start_month"]),
            help="6月起始：FY27Q1=2026年6–8月。采用月份口径，不模拟公司逐周财务日历。",
        )
        legacy = st.checkbox(
            "保留示例项目历史展示例外：2026年2月记为FY26Q4",
            value=s["legacy_fy"],
            help="沿用原文件项目约定，仅在6月起始财年时有效。",
        )
        schedule = st.selectbox(
            "排期维护方式",
            ["模板排期", "自动联动", "手工维护"],
            index=["模板排期", "自动联动", "手工维护"].index(p["timeline_mode"]),
            help="示例项目模板/自动联动采用原有固定日期并随基准日整体平移；通用模板自动联动重建排期。手工维护保留当前日期。重建会替换自定义任务。",
        )
        approved = st.number_input(
            "获批预算上限（RMB，0表示未设置）",
            min_value=0.0,
            value=float(s["approved_budget_rmb"]),
            step=100000.0,
        )
        saved = st.form_submit_button("保存项目设置", type="primary")
    if saved:
        candidate = copy.deepcopy(p)
        candidate["settings"].update(
            project_name=name,
            study_type=study,
            n_planned=n,
            site_number=sites,
            n_lost=lost,
            n_screen_fail=failed,
            start_date=start.isoformat(),
            primary_fu_months=primary,
            final_fu_months=final,
            enrollment_rate=rate,
            startup_months=startup,
            ramp_months=ramp,
            cleaning_months=cleaning,
            report_months=report,
            display_currency=currency,
            fx=fx,
            fy_start_month=fy,
            legacy_fy=legacy,
            approved_budget_rmb=approved,
        )
        candidate["timeline_mode"] = schedule
        errors = validate_project(candidate)
        if errors:
            for e in errors[:8]:
                st.error(e)
        else:
            try:
                if schedule != "手工维护":
                    candidate["timeline"] = make_timeline(candidate)
            except (ValueError, OverflowError) as exc:
                st.error(str(exc))
            else:
                apply_change(candidate, "项目设置已保存，自动关联数量已更新。")
    if p["template"] == "示例项目":
        st.info(
            "示例项目采用预填排期。修改入组速率仅影响情景测算；修改基准日期会在模板排期模式下整体平移日期。"
        )


def ui_records(df, reverse, base=None):
    """Convert editor values to plain JSON; preserve only explicitly defined fields."""
    records = []
    for _, r in df.iterrows():
        record = {}
        for column, key in reverse.items():
            value = r.get(column)
            if pd.isna(value):
                value = None
            if isinstance(value, (dt.date, dt.datetime, pd.Timestamp)):
                value = as_date(value).isoformat()
            elif hasattr(value, "item"):
                value = value.item()
            record[key] = value
        records.append(record)
    return records


def render_budget(p, lines):
    s = p["settings"]
    rev = form_epoch(
        "budget",
        [
            p["budget"],
            {
                k: s[k]
                for k in (
                    "site_management_pct",
                    "site_tax_pct",
                    "vendor_tax_pct",
                    "reserve_pct",
                    "tax_includes_management",
                )
            },
        ],
    )
    st.subheader("预算编制")
    st.caption(
        "可像Excel一样连续填写、粘贴和新增行；完成后统一保存。每项费用独立选择币种。"
    )
    total, fmt = metric_values(p, lines)
    c = st.columns(3)
    c[0].metric("当前总预算 · " + s["display_currency"], fmt(total))
    c[1].metric(
        "研究中心费用",
        fmt(sum(l["amount_rmb"] for l in lines if l["module"] == "研究中心")),
    )
    c[2].metric(
        "产品物流费用",
        fmt(sum(l["amount_rmb"] for l in lines if l["module"] == "产品物流")),
    )
    tab1, tab2, tab3 = stable_tabs(
        ["费用明细编辑", "访视计划与计费", "计算结果与构成"], key="tabs_budget"
    )
    with tab1:
        st.caption(
            "自动依据（人数/中心/月数）忽略“固定数量/FTE”列；FTE项目=FTE×项目单价，FTE月/年按实际项目周期折算。"
        )
        names = {
            "id": "费用ID",
            "enabled": "启用",
            "module": "预算模块",
            "item": "费用项目",
            "currency": "币种",
            "unit_price": "单价",
            "basis": "计费依据",
            "quantity": "固定数量/FTE",
            "phase": "匹配阶段",
            "allocation": "分配方式",
            "payment_lag": "付款滞后(月)",
            "note": "备注",
        }
        df = pd.DataFrame(p["budget"], columns=names).rename(columns=names)
        with st.form(f"budget_{rev}"):
            edited = editable_table(
                df,
                hide_index=True,
                width="stretch",
                num_rows="dynamic",
                height=510,
                key=f"budget_grid_{rev}",
                column_config={
                    "费用ID": st.column_config.TextColumn(disabled=True),
                    "启用": st.column_config.CheckboxColumn(default=True),
                    "预算模块": st.column_config.SelectboxColumn(
                        options=MODULES, required=True
                    ),
                    "费用项目": st.column_config.TextColumn(
                        required=True, width="large"
                    ),
                    "币种": st.column_config.SelectboxColumn(
                        options=["RMB", "USD"], default="RMB", required=True
                    ),
                    "单价": st.column_config.NumberColumn(
                        min_value=0.0, format="%.2f", required=True
                    ),
                    "计费依据": st.column_config.SelectboxColumn(
                        options=BASIS, default="固定数量", required=True
                    ),
                    "固定数量/FTE": st.column_config.NumberColumn(
                        min_value=0.0, default=1.0, format="%.3f"
                    ),
                    "匹配阶段": st.column_config.SelectboxColumn(
                        options=["智能匹配"] + PHASES, default="智能匹配", required=True
                    ),
                    "分配方式": st.column_config.SelectboxColumn(
                        options=ALLOC_MODES, default="智能规则", required=True
                    ),
                    "付款滞后(月)": st.column_config.NumberColumn(
                        min_value=0, max_value=120, default=0, step=1
                    ),
                },
            )
            st.markdown("**管理费、税费和储备**")
            a, b, c, d = st.columns(4)
            mgmt = a.number_input(
                "中心管理费率（%）", 0.0, 100.0, float(s["site_management_pct"])
            )
            tax = b.number_input("Site税费（%）", 0.0, 100.0, float(s["site_tax_pct"]))
            vendor = c.number_input(
                "Vendor税费（%）", 0.0, 100.0, float(s["vendor_tax_pct"])
            )
            reserve = d.number_input(
                "风险储备（%）", 0.0, 100.0, float(s["reserve_pct"])
            )
            incl = st.checkbox(
                "Site计税基数包含中心管理费", value=s["tax_includes_management"]
            )
            st.caption(
                "中心管理费基数＝研究中心直接费用＋启用的访视费用；Vendor税基数＝供应商直接费用。示例项目预算按总额预填，默认不再叠加管理费或税费。"
            )
            submitted = st.form_submit_button("保存预算明细", type="primary")
        if submitted:
            candidate = copy.deepcopy(p)
            rows = ui_records(edited, {v: k for k, v in names.items()})
            defaults = dict(
                enabled=True,
                currency="RMB",
                basis="固定数量",
                quantity=1.0,
                phase="智能匹配",
                allocation="智能规则",
                payment_lag=0,
                note="",
            )
            for r in rows:
                for k, v in defaults.items():
                    if r[k] is None:
                        r[k] = v
                if not r["id"]:
                    r["id"] = "B" + uuid.uuid4().hex[:10]
            candidate["budget"] = rows
            candidate["settings"].update(
                site_management_pct=mgmt,
                site_tax_pct=tax,
                vendor_tax_pct=vendor,
                reserve_pct=reserve,
                tax_includes_management=incl,
            )
            apply_change(candidate, "预算已保存；明细、现金流和摘要已同步。")
    with tab2:
        vrev = form_epoch("visits", p["visits"])
        st.caption(
            "每个访视的费用按“入组月份＋访视间隔”匹配现金流；默认假设入组窗口内每月均匀入组。"
        )
        if p["template"] == "示例项目":
            st.info(
                "原始Site Cost已含随访费用，以下访视默认不额外计费。仅在确认需要单独追加费用时启用。"
            )
        names = {
            "id": "访视ID",
            "enabled": "启用计费",
            "name": "访视名称",
            "offset": "入组后(月)",
            "times": "每例次数",
            "unit_price": "单价/次",
            "currency": "币种",
            "population": "计费人群",
        }
        df = pd.DataFrame(p["visits"], columns=names).rename(columns=names)
        with st.form(f"visits_{vrev}"):
            edited = editable_table(
                df,
                hide_index=True,
                width="stretch",
                num_rows="dynamic",
                key=f"visits_grid_{vrev}",
                column_config={
                    "访视ID": st.column_config.TextColumn(disabled=True),
                    "启用计费": st.column_config.CheckboxColumn(default=True),
                    "访视名称": st.column_config.TextColumn(required=True),
                    "入组后(月)": st.column_config.NumberColumn(
                        min_value=0, max_value=240, step=1, required=True
                    ),
                    "每例次数": st.column_config.NumberColumn(
                        min_value=0.0, default=1.0
                    ),
                    "单价/次": st.column_config.NumberColumn(
                        min_value=0.0, required=True, format="%.2f"
                    ),
                    "币种": st.column_config.SelectboxColumn(
                        options=["RMB", "USD"], default="RMB", required=True
                    ),
                    "计费人群": st.column_config.SelectboxColumn(
                        options=["入组例数", "完成随访例数"],
                        default="入组例数",
                        required=True,
                    ),
                },
            )
            submitted = st.form_submit_button("保存访视计划", type="primary")
        if submitted:
            candidate = copy.deepcopy(p)
            rows = ui_records(edited, {v: k for k, v in names.items()})
            for r in rows:
                if not r["id"]:
                    r["id"] = "V" + uuid.uuid4().hex[:10]
                for key, value in dict(
                    enabled=True, times=1.0, currency="RMB", population="入组例数"
                ).items():
                    if r[key] is None:
                        r[key] = value
            candidate["visits"] = rows
            apply_change(candidate, "访视计划和费用已保存。")
        a, b = enrollment_window(p)
        windows = [
            dict(
                访视=v["name"],
                首例访视=add_months(a, v["offset"]),
                末例访视=add_months(b, v["offset"]),
                额外计费="是" if v["enabled"] else "否",
            )
            for v in p["visits"]
        ]
        readonly_table(pd.DataFrame(windows), hide_index=True, width="stretch")
    with tab3:
        readonly_table(
            pd.DataFrame(report_tables(p)["预算明细"]),
            hide_index=True,
            width="stretch",
            height=480,
        )
        detail = pd.DataFrame(lines)
        if detail.amount_rmb.sum() > 0:
            module = st.selectbox(
                "查看模块明细", MODULES + ["风险储备"], key="module_detail"
            )
            detail = detail[(detail.module == module) & (detail.amount_rmb > 0)].copy()
            if not detail.empty:
                detail["显示金额"] = detail.amount_rmb.apply(
                    lambda x: display_amount(x, s["display_currency"], s["fx"])
                )
                fig = px.bar(
                    detail.sort_values("显示金额"),
                    x="显示金额",
                    y="item",
                    orientation="h",
                    color_discrete_sequence=[COLORS[0]],
                )
                fig.update_layout(yaxis_title=None, xaxis_title=s["display_currency"])
                st.plotly_chart(
                    chart_style(fig, max(330, len(detail) * 32)), width="stretch"
                )


def render_timeline(p):
    rev = form_epoch("timeline", p["timeline"])
    st.subheader("时间线管理")
    st.caption(
        f"当前：{p['timeline_mode']}。手工保存后切换为“手工维护”，避免下一次改参数覆盖你的排期。"
    )
    detailed = st.checkbox("展开任务明细", value=False)
    overlay = st.checkbox(
        "叠加保存的基线",
        value=p["baseline"] is not None,
        disabled=p["baseline"] is None,
    )
    draw_timeline(p, detail=detailed, overlay=overlay)
    names = {
        "id": "任务ID",
        "enabled": "启用",
        "phase": "阶段",
        "task": "任务名称",
        "start": "开始日期",
        "end": "结束日期",
        "owner": "负责人",
        "progress": "完成度(%)",
        "actual_end": "实际完成日期",
        "depends_on": "前置任务ID",
    }
    df = pd.DataFrame(p["timeline"], columns=names).rename(columns=names)
    for col in ["开始日期", "结束日期", "实际完成日期"]:
        df[col] = pd.to_datetime(df[col]).dt.date
    with st.form(f"timeline_{rev}"):
        edited = editable_table(
            df,
            hide_index=True,
            width="stretch",
            num_rows="dynamic",
            height=540,
            key=f"timeline_grid_{rev}",
            column_config={
                "任务ID": st.column_config.TextColumn(disabled=True),
                "启用": st.column_config.CheckboxColumn(default=True),
                "阶段": st.column_config.SelectboxColumn(options=PHASES, required=True),
                "任务名称": st.column_config.TextColumn(required=True, width="large"),
                "开始日期": st.column_config.DateColumn(
                    format="YYYY-MM-DD", required=True
                ),
                "结束日期": st.column_config.DateColumn(
                    format="YYYY-MM-DD", required=True
                ),
                "实际完成日期": st.column_config.DateColumn(format="YYYY-MM-DD"),
                "完成度(%)": st.column_config.NumberColumn(
                    min_value=0.0, max_value=100.0, default=0.0, step=5.0
                ),
                "前置任务ID": st.column_config.TextColumn(
                    help="填写任务ID，如T001；表示本任务应在前置任务完成后开始。依赖冲突会提示，不自动改期。"
                ),
            },
        )
        submitted = st.form_submit_button("保存时间线", type="primary")
    if submitted:
        candidate = copy.deepcopy(p)
        rows = ui_records(edited, {v: k for k, v in names.items()})
        for r in rows:
            if not r["id"]:
                r["id"] = "T" + uuid.uuid4().hex[:10]
            for key, value in dict(
                enabled=True, owner="", progress=0.0, depends_on=""
            ).items():
                if r[key] is None:
                    r[key] = value
        candidate["timeline"] = rows
        candidate["timeline_mode"] = "手工维护"
        apply_change(candidate, "时间线已保存，现金流已根据当前分配方式重新计算。")
    with stable_expander("重新生成模板 / 自动排期"):
        st.caption(
            "保留标准任务已登记的负责人和完成度，重新生成日期；自定义任务会被移除。先导出JSON可保留完整版本。"
        )
        if st.button("重建排期", key="rebuild_timeline"):
            candidate = copy.deepcopy(p)
            try:
                candidate["timeline"] = make_timeline(candidate)
            except (ValueError, OverflowError) as exc:
                st.error(str(exc))
            else:
                candidate["timeline_mode"] = (
                    "模板排期" if p["template"] == "示例项目" else "自动联动"
                )
                apply_change(candidate, "排期已重新生成。")


def render_cash(p, lines, cash, warnings):
    s = p["settings"]
    rev = form_epoch("cash", [s["cash_mode"], p["cash_weights"]])
    st.subheader("现金流计划")
    st.caption(
        "每项预算按分摊规则或原始季度匹配到月份，再叠加付款滞后；预算与现金流按分核对。"
    )
    with stable_expander("分配方式与自定义季度", expanded=False):
        with st.form(f"cash_{rev}"):
            mode = st.selectbox(
                "现金流分配方式", CASH_MODES, index=CASH_MODES.index(s["cash_mode"])
            )
            st.caption(
                "示例项目预填季度：逐项按原矩阵比例缩放到当前金额；新增项使用智能规则。智能匹配：使用原费用规则＋当前任务日期。"
            )
            df = pd.DataFrame(p["cash_weights"], columns=["quarter", "weight"]).rename(
                columns={"quarter": "财季", "weight": "占比(%)"}
            )
            edited = editable_table(
                df,
                hide_index=True,
                width="stretch",
                num_rows="dynamic",
                height=280,
                key=f"weights_grid_{rev}",
                column_config={
                    "财季": st.column_config.TextColumn(
                        required=True, help="示例FY27Q1；FY用结束年度命名"
                    ),
                    "占比(%)": st.column_config.NumberColumn(
                        min_value=0.0, max_value=100.0, required=True, format="%.2f"
                    ),
                },
            )
            c1, c2 = st.columns(2)
            submit = c1.form_submit_button("保存现金流设置", type="primary")
            normalize = c2.form_submit_button("归一化占比并保存")
        if submit or normalize:
            candidate = copy.deepcopy(p)
            weights = ui_records(edited, {"财季": "quarter", "占比(%)": "weight"})
            if normalize:
                try:
                    portions = split_money(100, [r["weight"] for r in weights])
                    weights = [dict(r, weight=w) for r, w in zip(weights, portions)]
                except (ValueError, TypeError) as exc:
                    st.error(str(exc))
                    return
            for row in weights:
                if isinstance(row["quarter"], str):
                    row["quarter"] = row["quarter"].strip().upper()
            candidate["cash_weights"] = weights
            candidate["settings"]["cash_mode"] = mode
            apply_change(candidate, "现金流分配设置已保存。")
        if st.button("按当前时间线重建均分占比", key="rebuild_weights"):
            candidate = copy.deepcopy(p)
            candidate["cash_weights"] = default_weights(candidate)
            apply_change(candidate, "季度占比表已按当前时间线重建。")
    if not cash:
        st.info("预算为0，暂无现金流。")
        return
    total, fmt = metric_values(p, lines)
    quarters = quarter_totals(cash, p)
    peak = max(quarters, key=lambda x: x["amount_rmb"])
    c = st.columns(3)
    c[0].metric(
        "现金流总计 · " + s["display_currency"], fmt(sum(r["amount_rmb"] for r in cash))
    )
    c[1].metric("资金需求峰值季度", peak["quarter"])
    c[2].metric("峰值季度资金 · " + s["display_currency"], fmt(peak["amount_rmb"]))
    df = pd.DataFrame(cash)
    df["显示金额"] = df.amount_rmb.apply(
        lambda x: display_amount(x, s["display_currency"], s["fx"])
    )
    quarterly = df.groupby(["quarter", "module"], as_index=False)["显示金额"].sum()
    monthly = (
        df.groupby("month", as_index=False)["amount_rmb"].sum().sort_values("month")
    )
    # Reindex so zero-spend months remain visible and cumulative curves are continuous.
    full = pd.date_range(monthly.month.min(), monthly.month.max(), freq="MS").strftime(
        "%Y-%m-%d"
    )
    monthly = (
        monthly.set_index("month")
        .reindex(full, fill_value=0)
        .rename_axis("month")
        .reset_index()
    )
    monthly["显示金额"] = monthly.amount_rmb.apply(
        lambda x: display_amount(x, s["display_currency"], s["fx"])
    )
    monthly["累计"] = monthly.amount_rmb.cumsum().apply(
        lambda x: display_amount(x, s["display_currency"], s["fx"])
    )
    left, right = st.columns([1.4, 1])
    with left:
        st.markdown("**季度资金需求**")
        fig = px.bar(
            quarterly,
            x="quarter",
            y="显示金额",
            color="module",
            color_discrete_sequence=COLORS,
            category_orders={"quarter": sorted(quarterly.quarter.unique())},
        )
        fig.update_layout(
            xaxis_title=None, yaxis_title=s["display_currency"], legend_title=None
        )
        st.plotly_chart(chart_style(fig, 390), width="stretch")
    with right:
        st.markdown("**累计计划支出**")
        fig = px.area(monthly, x="month", y="累计", color_discrete_sequence=[COLORS[0]])
        fig.update_layout(xaxis_title=None, yaxis_title=s["display_currency"])
        st.plotly_chart(chart_style(fig, 390), width="stretch")
    st.markdown("**财年汇总**")
    qdf = pd.DataFrame(quarters)
    qdf["财年"] = qdf.quarter.str.split("Q").str[0]
    annual = qdf.groupby("财年", as_index=False).amount_rmb.sum()
    annual[s["display_currency"]] = annual.amount_rmb.apply(
        lambda x: display_amount(x, s["display_currency"], s["fx"])
    )
    readonly_table(
        annual[["财年", s["display_currency"]]], hide_index=True, width="stretch"
    )
    with stable_expander("查看季度和月度明细 / 匹配规则"):
        readonly_table(pd.DataFrame(quarters), hide_index=True, width="stretch")
        readonly_table(
            df.drop(columns="显示金额"), hide_index=True, width="stretch", height=300
        )
        st.caption(
            "金额_RMB为对账基准；均摊只是费用归属或付款计划估算，合同付款节点可用分配方式及滞后月数调整。"
        )
        for w in warnings:
            st.caption("• " + w)
    if p["template"] == "示例项目":
        with stable_expander("原始示例项目预算矩阵（USD）"):
            matrix = pd.DataFrame(
                EXAMPLE_BUDGET_MATRIX_USD, index=EXAMPLE_BUDGET_QUARTERS
            ).T
            matrix["总额"] = matrix.sum(axis=1)
            readonly_table(matrix, width="stretch")
            st.caption(
                "原始矩阵仅供溯源。当前现金流以已启用费用及编辑后的金额为准；MC2不计入默认总额。"
            )


def render_tracking(p, lines, cash, warnings):
    s = p["settings"]
    rev = form_epoch(
        "expenses", [p["expenses"], [(r["id"], r["item"]) for r in p["budget"]]]
    )
    total, fmt = metric_values(p, lines)
    st.subheader("执行跟踪")
    actual = money(
        sum(to_rmb(e["amount"], e["currency"], s["fx"]) for e in p["expenses"])
    )
    c = st.columns(3)
    c[0].metric("已登记实际支出 · " + s["display_currency"], fmt(actual))
    c[1].metric("剩余预算 · " + s["display_currency"], fmt(total - actual))
    baseline = p["baseline"]
    c[2].metric(
        "相对基线预算变动 · " + s["display_currency"],
        fmt(total - baseline["total_rmb"]) if baseline else "未保存基线",
    )
    tab1, tab2, tab3 = stable_tabs(
        ["实际支出台账", "基线与变更", "检查与提醒"], key="tabs_tracking"
    )
    with tab1:
        st.caption(
            "这里只登记已支付金额；剩余预算＝当前预算－已支付，不代表完工尚需成本。外币支出按当前规划汇率重估。"
        )
        lookup = {l["id"]: l["item"] for l in lines}
        lookup.update({r["id"]: r["item"] for r in p["budget"]})
        lookup.update({v["id"]: "访视：" + v["name"] for v in p["visits"]})
        lookup.update(
            {
                "D-MGMT": "中心管理费用",
                "D-SITETAX": "Site税费",
                "D-VENDORTAX": "Vendor税费",
                "D-RESERVE": "风险储备",
            }
        )
        with stable_expander("费用ID对照", expanded=not bool(p["expenses"])):
            readonly_table(
                pd.DataFrame([dict(费用ID=k, 项目=v) for k, v in lookup.items()]),
                hide_index=True,
                width="stretch",
                height=200,
            )
        names = {
            "id": "记录ID",
            "date": "支付日期",
            "budget_id": "费用ID",
            "amount": "支付金额",
            "currency": "币种",
            "note": "备注",
        }
        df = pd.DataFrame(p["expenses"], columns=names).rename(columns=names)
        df["支付日期"] = pd.to_datetime(df["支付日期"]).dt.date
        with st.form(f"expenses_{rev}"):
            edited = editable_table(
                df,
                hide_index=True,
                width="stretch",
                num_rows="dynamic",
                key=f"expenses_grid_{rev}",
                column_config={
                    "记录ID": st.column_config.TextColumn(disabled=True),
                    "支付日期": st.column_config.DateColumn(
                        format="YYYY-MM-DD", required=True
                    ),
                    "费用ID": st.column_config.SelectboxColumn(
                        options=list(lookup), required=True
                    ),
                    "支付金额": st.column_config.NumberColumn(
                        min_value=0.0, required=True, format="%.2f"
                    ),
                    "币种": st.column_config.SelectboxColumn(
                        options=["RMB", "USD"], default="RMB", required=True
                    ),
                },
            )
            submit = st.form_submit_button("保存实际支出", type="primary")
        if submit:
            candidate = copy.deepcopy(p)
            rows = ui_records(edited, {v: k for k, v in names.items()})
            for r in rows:
                if not r["id"]:
                    r["id"] = "E" + uuid.uuid4().hex[:10]
                if not r["currency"]:
                    r["currency"] = "RMB"
                if r["note"] is None:
                    r["note"] = ""
            candidate["expenses"] = rows
            apply_change(candidate, "实际支出台账已保存。")
        if p["expenses"]:
            planned = {l["id"]: l["amount_rmb"] for l in lines}
            paid = {}
            for e in p["expenses"]:
                paid[e["budget_id"]] = money(
                    paid.get(e["budget_id"], 0)
                    + to_rmb(e["amount"], e["currency"], s["fx"])
                )
            rows = [
                dict(
                    费用ID=k,
                    项目=lookup.get(k, k),
                    当前预算_RMB=planned.get(k, 0),
                    已付_RMB=v,
                    剩余_RMB=money(planned.get(k, 0) - v),
                )
                for k, v in paid.items()
            ]
            readonly_table(pd.DataFrame(rows), hide_index=True, width="stretch")
    with tab2:
        st.caption(
            "基线冻结保存时的预算、任务日期和汇率折算结果，用于跟踪后续变更；导出JSON会包含基线。"
        )
        name = st.text_input(
            "基线名称", value="项目预算与排期基线", max_chars=120, key="baseline_name"
        )
        if st.button(
            "保存当前基线" if baseline is None else "用当前数据替换基线", type="primary"
        ):
            candidate = copy.deepcopy(p)
            candidate["baseline"] = save_baseline(candidate, name)
            apply_change(candidate, "项目基线已保存。")
        if baseline:
            st.write("当前基线：", baseline["name"])
            st.caption("保存于 " + baseline["captured_at"])
            days = (project_span(p)[1] - as_date(baseline["end"])).days
            st.metric("项目结束日期相对基线", f"{days:+d} 天")
            baseline_tasks = {t["id"]: t for t in baseline["timeline"]}
            rows = []
            for t in p["timeline"]:
                old = baseline_tasks.get(t["id"])
                if old and t["enabled"]:
                    diff = (as_date(t["end"]) - as_date(old["end"])).days
                    if diff:
                        rows.append(
                            dict(
                                任务=t["task"],
                                基线结束=old["end"],
                                当前结束=t["end"],
                                变动天数=diff,
                            )
                        )
            if rows:
                readonly_table(pd.DataFrame(rows), hide_index=True, width="stretch")
            else:
                st.success("已存在任务的结束日期未偏离基线。")
            current = {l["id"]: l for l in lines}
            amounts = baseline["line_amounts"]
            changes = [
                dict(
                    费用ID=k,
                    费用项目=current.get(k, {}).get("item", k),
                    基线_RMB=amounts.get(k, 0),
                    当前_RMB=current.get(k, {}).get("amount_rmb", 0),
                    变化_RMB=money(
                        current.get(k, {}).get("amount_rmb", 0) - amounts.get(k, 0)
                    ),
                )
                for k in sorted(set(current) | set(amounts))
            ]
            readonly_table(
                pd.DataFrame([r for r in changes if r["变化_RMB"]]),
                hide_index=True,
                width="stretch",
            )
    with tab3:
        findings = checks(p, lines, cash, warnings)
        for level in ("错误", "关注", "提示"):
            filtered = [x for x in findings if x["level"] == level]
            if filtered:
                with stable_expander(
                    f"{level} · {len(filtered)}条", expanded=level != "提示"
                ):
                    for x in filtered:
                        if level == "错误":
                            st.error(x["message"])
                        elif level == "关注":
                            st.warning(x["message"])
                        else:
                            st.info(x["message"])
        if not findings:
            st.success("未检出明显关联问题。")


def scenario_project(project: dict, rate_factor: float, n: int, sites: int) -> dict:
    """What-if only: scale existing enrollment duration; shift downstream windows."""
    candidate = copy.deepcopy(project)
    candidate["settings"].update(n_planned=n, site_number=sites)
    candidate["settings"]["n_lost"] = min(candidate["settings"]["n_lost"], n)
    base = len(enrollment_curve(project))
    target = len(enrollment_curve(candidate, rate_factor))
    old_start, old_end = enrollment_window(project)
    old_months = len(month_starts(old_start, old_end))
    new_months = max(1, math.ceil(old_months * target / base))
    new_end = add_months(old_start, new_months) - dt.timedelta(days=1)
    delta = new_end - old_end
    for t in candidate["timeline"]:
        if not t["enabled"]:
            continue
        if t["task"] == "受试者入组":
            t["end"] = new_end.isoformat()
        elif t["task"] == "末例入组":
            t["start"] = (as_date(t["start"]) + delta).isoformat()
            t["end"] = (as_date(t["end"]) + delta).isoformat()
        elif t["phase"] in ("受试者随访", "数据分析", "研究结束"):
            a, b = as_date(t["start"]), as_date(t["end"])
            if a >= old_end:
                a += delta
            b += delta
            t["start"] = a.isoformat()
            t["end"] = max(a, b).isoformat()
    return candidate


def render_scenario(p):
    st.subheader("入组与预算情景测算")
    s = p["settings"]
    st.caption("对比入组速率变化对当前排期和工期相关费用的影响；此页面不会修改项目。")
    with st.form("scenario"):
        a, b = st.columns(2)
        n = a.number_input("情景计划入组人数", 1, 100000, int(s["n_planned"]))
        sites = b.number_input("情景中心数", 1, 100000, int(s["site_number"]))
        st.form_submit_button("更新情景对比", type="primary")
    rows = []
    curves = []
    try:
        for label, factor in [
            ("保守：速率−25%", 0.75),
            ("参考速率", 1.0),
            ("积极：速率+25%", 1.25),
        ]:
            candidate = scenario_project(p, factor, n, sites)
            lines = budget_lines(candidate)
            e_start, e_end = enrollment_window(candidate)
            rows.append(
                dict(
                    情景=label,
                    模型入组月数=len(enrollment_curve(candidate, factor)),
                    当前窗口调整后末例入组=e_end.isoformat(),
                    推算项目结束=project_span(candidate)[1].isoformat(),
                    预算_RMB=money(sum(l["amount_rmb"] for l in lines)),
                )
            )
            for point in enrollment_curve(candidate, factor):
                curves.append(
                    dict(
                        情景=label,
                        月份=point["month"],
                        预计累计入组=round(point["cumulative"], 2),
                    )
                )
    except (ValueError, OverflowError) as exc:
        st.error(str(exc))
        return
    readonly_table(pd.DataFrame(rows), hide_index=True, width="stretch")
    fig = px.line(
        pd.DataFrame(curves),
        x="月份",
        y="预计累计入组",
        color="情景",
        color_discrete_sequence=["#8295B4", COLORS[0], COLORS[1]],
    )
    st.plotly_chart(chart_style(fig, 390), width="stretch")
    st.info(
        "模型按中心线性启动和恒定单中心速率预测。调整后的入组窗口按模型月数比例缩放，后续任务整体顺延；按固定总额录入的费用不自动涨价。仅人数、中心数、项目月数或FTE月/年等关联费用重算；额外监查、续保等费用需手工补充。"
    )


def render_export(p):
    st.subheader("导出与帮助")
    st.caption("项目JSON用于继续编辑；Excel/CSV用于汇报与对账。下载包含已保存的数据。")
    tables = report_tables(p)
    st.radio(
        "表格导出语言",
        ["中文", "English"],
        index=1 if current_export_language() == "en" else 0,
        key="export_language",
        on_change=remember_export_language,
        horizontal=True,
        help="统一用于全部Excel及各页面CSV下载；界面和项目JSON保持原语言。",
    )
    language = current_export_language()
    st.info(
        "当前表格导出语言："
        + ("English" if language == "en" else "中文")
        + "。下方预览与各页面表格下载采用相同语言。"
    )
    st.caption(
        "英文模式转换13张工作表的表名、表头、系统预设费用/任务、启动检查和系统提示。"
        "医院、PI、联系人、自定义费用/任务和自由备注保留原文；不调用外部翻译服务。"
    )
    file_stub = (
        re.sub(r"[^\w\u4e00-\u9fff-]+", "_", p["settings"]["project_name"])[:65]
        or "CRMS"
    )
    a, b = st.columns(2)
    a.download_button(
        "下载完整项目 JSON",
        project_bytes(p),
        file_name=file_stub + "_project.json",
        mime="application/json",
        width="stretch",
        on_click="ignore",
    )
    if b.button("准备Excel工作簿", width="stretch"):
        st.session_state.excel_export = (
            st.session_state.revision,
            excel_bytes(tables, language),
            language,
        )
    prepared = st.session_state.get("excel_export")
    if (
        prepared
        and len(prepared) == 3
        and prepared[0] == st.session_state.revision
        and prepared[2] == language
    ):
        b.download_button(
            "下载Excel工作簿（全部数据表）",
            prepared[1],
            file_name=file_stub
            + ("_report_EN.xlsx" if language == "en" else "_report.xlsx"),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            width="stretch",
            on_click="ignore",
        )
    elif prepared:
        st.caption("数据或导出语言已更改，请重新点击“准备Excel工作簿”。")
    selected = st.selectbox(
        "选择单表CSV",
        list(tables),
        key="export_table",
        format_func=lambda name: export_table_name(name, language),
    )
    st.download_button(
        "下载CSV",
        csv_bytes(tables[selected], language, selected),
        file_name=export_csv_filename(
            file_stub + "_" + export_table_name(selected, language).replace(" ", "_")
        ),
        mime="text/csv",
        on_click="ignore",
    )
    st.caption("当前单表预览（最多显示前300行；下载包含全部已保存记录）")
    preview = export_frame(tables[selected], language, selected).head(300)
    if selected == "项目摘要":
        preview = preview.astype(
            str
        )  # Mixed metric values need a display-only string column.
    readonly_table(
        preview,
        hide_index=True,
        width="stretch",
    )
    with stable_expander("填写与计算说明", expanded=True):
        st.markdown("""
1. **项目设置**：填写人数、中心和每例随访月数；保存后计算关联数量。通用自动排期允许准备阶段并行。
2. **预算编制**：逐行选择币种、单价和计费依据；自动数量随人数/中心/月数更新，固定数量按填值计算。小时费用可用固定数量录入工时。
3. **时间线管理**：批量编辑日期、责任人和完成度。开始/结束日期含当日；前置ID用于冲突提示，不自动重排。
4. **现金流计划**：智能匹配使用规则与当前日期；示例项目预填季度保持原始季度比例，金额随预算变化。自定义季度必须合计100%。付款滞后在费用行设置。
5. **执行跟踪**：记录实际支付，保存预算/时间线基线。表单中的编辑须点击保存；当前版本无数据库或跨用户共享编辑。
6. **中心启动与入组**：先建立中心和PI清单，用阶段指针与16项检查追踪启动；再填写各中心计划和带日期的实际累计入组，导出虚线/实线曲线。
7. **保存**：导出JSON后可随时重新载入。浏览器刷新、会话中断或服务重启可能清空当前会话，部署代码不会自动保存用户填报数据。
8. **英文表格**：在本页将“表格导出语言”切换为English；全部Excel和各页面CSV使用英文系统标签。医院/PI等专有名称、自定义文字及备注保留原文。项目JSON用于恢复原始数据，不作翻译；图表图片和HTML仍沿用中文界面标签。
""")
    with stable_expander("范围与模板口径"):
        st.write(
            "示例项目提供预填的预算与排期，仅用于演示功能。开始正式填报前，请依据实际研究方案、报价和合同核对数据。项目财季采用月份口径，可在项目设置中调整。"
        )
        st.write(
            "智能计算由明确规则驱动，未调用大模型；入组预测和费用分配为规划估算。当前版本支持单项目/单会话；账号权限、多人协作、审批及数据库可作为后续正式产品扩展。"
        )


def main():
    global st, pd, px
    import streamlit as st
    import pandas as pd
    import plotly.express as px

    st.set_page_config(
        page_title="CRMS · 临床研究工作台",
        page_icon="📊",
        layout="wide",
        initial_sidebar_state="expanded",
    )
    st.markdown(CSS, unsafe_allow_html=True)
    if "project" not in st.session_state:
        st.session_state.project = make_project("示例项目")
        st.session_state.revision = 0
    if (
        "sites" not in st.session_state.project
        or "enrollment" not in st.session_state.project
        or st.session_state.project.get("app_version") != APP_VERSION
    ):
        st.session_state.project = upgrade_project(st.session_state.project)
    p = st.session_state.project
    page = sidebar(p)
    s = p["settings"]
    st.markdown(
        f'<div class="hero"><div class="eyebrow">CLINICAL RESEARCH MANAGEMENT · {html.escape(page)}</div>'
        f'<h1>{html.escape(s["project_name"])}</h1><p>{html.escape(s["study_type"])} &nbsp; / &nbsp; '
        f'{s["site_number"]} 家中心 &nbsp; / &nbsp; {s["n_planned"]} 例计划入组 &nbsp; / &nbsp; '
        f'{html.escape(s["cash_mode"])}<br>让研究计划、预算与执行进度在同一处衔接。</p></div>',
        unsafe_allow_html=True,
    )
    if s["project_name"] in ("示例项目", "示例项目（预填演示）"):
        st.info(
            "当前显示示例项目的预填信息，用于演示预算和时间线匹配。请先修改项目名称，再根据实际研究核对并填写数据。"
        )
    render_field_legend()
    notice_slot = st.empty()
    notice = st.session_state.pop("notice", None)
    if notice:
        notice_slot.success(notice)
    errors = validate_project(p)
    if errors:
        for err in errors[:10]:
            st.error(err)
        return
    try:
        lines = budget_lines(p)
        cash, warnings = allocate_cashflow(p, lines)
    except (ValueError, OverflowError) as exc:
        st.error("无法计算：" + str(exc))
        return
    if page == "项目概览":
        render_overview(p, lines, cash, warnings)
    elif page == "项目设置":
        render_settings(p)
    elif page == "预算编制":
        render_budget(p, lines)
    elif page == "时间线管理":
        render_timeline(p)
    elif page == "中心启动进度":
        render_sites(p)
    elif page == "入组进度":
        render_enrollment(p)
    elif page == "现金流计划":
        render_cash(p, lines, cash, warnings)
    elif page == "执行跟踪":
        render_tracking(p, lines, cash, warnings)
    elif page == "情景测算":
        render_scenario(p)
    elif page == "导出与帮助":
        render_export(p)


# ---------- V2.1: site startup and dated enrollment tracking ----------
# The 16 checklist labels below are transcribed from the user's reference image.
FPI_CHECKLIST = {
    "研究合规": [
        ("F01", "研究类型及监管路径已确认"),
        ("F02", "机构/科研管理审批完成"),
        ("F03", "本中心伦理批准有效"),
        ("F04", "适用的备案、登记、专项许可已完成"),
    ],
    "合同与财务": [
        ("F05", "CTA及必要服务协议生效"),
        ("F06", "经费支付及受试者补偿流程就绪"),
        ("F07", "受试者损害处理和保障安排明确"),
    ],
    "器械与临床资源": [
        ("F08", "试验器械与耗材可用且验收合格"),
        ("F09", "术者、导管室、配套设备就绪"),
        ("F10", "围手术期和应急预案确认"),
    ],
    "人员与数据": [
        ("F11", "PI及研究团队完成授权与培训"),
        ("F12", "EDC、CTMS及相关系统可用"),
        ("F13", "源文件、ISF及安全性报告流程就绪"),
    ],
    "招募与中心激活": [
        ("F14", "合规的受试者招募计划已建立"),
        ("F15", "SIV遗留关键问题全部关闭"),
        ("F16", "已发出正式中心激活确认"),
    ],
}
FPI_ITEM_IDS = [key for rows in FPI_CHECKLIST.values() for key, _ in rows]
SITE_STAGE_LABELS = ["未开始", *FPI_CHECKLIST]
READINESS_STATUSES = ["待确认", "已完成", "不适用"]


def china_today() -> dt.date:
    return dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).date()


def upgrade_project(project):
    """Add optional V2.1 collections; old V2 project files remain readable."""
    if isinstance(project, dict) and project.get("schema_version") == SCHEMA_VERSION:
        project = copy.deepcopy(project)
        project.setdefault("sites", [])
        project.setdefault("enrollment", [])
        if project.get("template") in ("ATHENA", "EXAMPLE"):
            project["template"] = "示例项目"
        settings = project.get("settings")
        if isinstance(settings, dict):
            if settings.get("project_name") in (
                "ATHENA Heli-FX China Trial",
                "EXAMPLE Heli-FX China Trial",
            ):
                settings["project_name"] = "示例项目（预填演示）"
            if settings.get("cash_mode") in ("ATHENA原始季度", "EXAMPLE原始季度"):
                settings["cash_mode"] = "示例项目预填季度"
        for site in project.get("sites", []):
            if isinstance(site, dict):
                site.setdefault("is_lead", False)
        project["app_version"] = APP_VERSION
    return project


def new_site(project, index=1, *, name="", code="") -> dict:
    a, b = enrollment_window(project)
    return dict(
        id="S" + uuid.uuid4().hex[:10],
        code=code or f"S{index:02d}",
        name=name or f"待填写中心{index:02d}",
        pi="",
        contact="",
        note="",
        planned_n=0,
        plan_start=a.isoformat(),
        plan_end=b.isoformat(),
        stage=0,
        readiness={key: "待确认" for key in FPI_ITEM_IDS},
        is_lead=False,
    )


def validate_sites(project) -> list[str]:
    errors = []
    sites, records = project.get("sites", []), project.get("enrollment", [])
    if (
        not isinstance(sites, list)
        or len(sites) > 500
        or any(not isinstance(r, dict) for r in sites)
    ):
        return ["中心清单必须为对象列表，最多500个中心。"]
    if (
        not isinstance(records, list)
        or len(records) > 20000
        or any(not isinstance(r, dict) for r in records)
    ):
        return ["入组记录必须为对象列表，最多20000条。"]
    site_fields = {
        "id",
        "code",
        "name",
        "pi",
        "contact",
        "note",
        "planned_n",
        "plan_start",
        "plan_end",
        "stage",
        "readiness",
    }
    record_fields = {"id", "site_id", "date", "cumulative", "note"}
    if any(not site_fields.issubset(r) for r in sites):
        return ["中心信息字段不完整。"]
    if any(not record_fields.issubset(r) for r in records):
        return ["入组快照字段不完整。"]
    for collection, label in ((sites, "中心"), (records, "入组快照")):
        ids = [r["id"] for r in collection]
        if any(not isinstance(v, str) or not v or len(v) > 80 for v in ids):
            errors.append(label + "ID无效。")
        elif len(set(ids)) != len(ids):
            errors.append(label + "ID重复。")
    if errors:
        return errors
    codes = []
    names = []

    def integer(value, low, high):
        return (
            isinstance(value, (int, float))
            and not isinstance(value, bool)
            and math.isfinite(value)
            and low <= value <= high
            and value == int(value)
        )

    for site in sites:
        for key, limit in (
            ("code", 40),
            ("name", 160),
            ("pi", 120),
            ("contact", 240),
            ("note", 2000),
        ):
            value = site[key]
            if (
                not isinstance(value, str)
                or len(value) > limit
                or (key in ("code", "name") and not value.strip())
            ):
                errors.append(f"中心{site['id']}的{key}字段无效；名称和编号不能为空。")
        if isinstance(site["code"], str):
            codes.append(site["code"].strip().casefold())
        if isinstance(site["name"], str):
            names.append(site["name"].strip().casefold())
        if not integer(site["planned_n"], 0, 100000):
            errors.append("中心计划入组数应为0–100000之间的整数。")
        if not integer(site["stage"], 0, 5):
            errors.append("中心启动阶段必须在0–5之间。")
        ready = site["readiness"]
        if (
            not isinstance(ready, dict)
            or set(ready) != set(FPI_ITEM_IDS)
            or any(v not in READINESS_STATUSES for v in ready.values())
        ):
            errors.append("启动检查须包含截图中的16项，状态为待确认/已完成/不适用。")
        elif all(v == "不适用" for v in ready.values()):
            errors.append("不能将一个中心的全部启动检查都设为不适用。")
        try:
            a, b = as_date(site["plan_start"]), as_date(site["plan_end"])
            if (
                a < dt.date(2000, 1, 1)
                or b > dt.date(2099, 12, 31)
                or b < a
                or (b - a).days > 366 * 25
            ):
                raise ValueError()
        except (TypeError, ValueError, OverflowError):
            errors.append(
                "中心计划起止日期无效（须在2000–2099年内、开始不晚于结束且跨度不超过25年）。"
            )
    if any(not isinstance(site.get("is_lead", False), bool) for site in sites):
        errors.append("牵头中心标记必须为布尔值。")
    if sum(site.get("is_lead", False) is True for site in sites) > 1:
        errors.append("只能勾选一个牵头中心，请取消其他中心的牵头标记。")
    if len(set(codes)) != len(codes):
        errors.append("中心编号不能重复（不区分大小写）。")
    if len(set(names)) != len(names):
        errors.append("中心名称不能重复；分院/科室请在名称中区分。")
    ids = {site["id"] for site in sites}
    snapshots = {}
    seen = set()
    for r in records:
        if not isinstance(r["site_id"], str) or r["site_id"] not in ids:
            errors.append(
                "入组记录关联的中心不存在；已有入组记录的中心请先处理记录再删除。"
            )
            continue
        try:
            day = as_date(r["date"])
            if day < dt.date(2000, 1, 1) or day > china_today():
                raise ValueError()
        except (TypeError, ValueError, OverflowError):
            errors.append("实际入组报告日期须为2000年起且不晚于今天的有效日期。")
            continue
        key = (r["site_id"], day)
        if key in seen:
            errors.append("同一中心同一天只保留一条累计快照，请编辑已有记录。")
        seen.add(key)
        if not integer(r["cumulative"], 0, 1000000):
            errors.append("实际累计入组须为非负整数。")
        else:
            snapshots.setdefault(r["site_id"], []).append((day, int(r["cumulative"])))
        if not isinstance(r["note"], str) or len(r["note"]) > 2000:
            errors.append("入组备注须为不超过2000字的文本。")
    for site_id, rows in snapshots.items():
        ordered = sorted(rows)
        if any(after[1] < before[1] for before, after in zip(ordered, ordered[1:])):
            errors.append(
                f"中心{site_id}累计入组不能随时间减少；更正误报时请同时修正相应历史快照。"
            )
    return list(dict.fromkeys(errors))


def site_readiness(site):
    values = list(site["readiness"].values())
    done = values.count("已完成")
    na = values.count("不适用")
    pending = values.count("待确认")
    applicable = len(values) - na
    return dict(
        done=done,
        na=na,
        pending=pending,
        ratio=done / applicable if applicable else 0,
        ready=bool(applicable and not pending),
    )


def planned_cumulative(site, date):
    a, b = as_date(site["plan_start"]), as_date(site["plan_end"])
    day = as_date(date)
    if day < a:
        return 0.0
    if day >= b:
        return float(site["planned_n"])
    return site["planned_n"] * ((day - a).days + 1) / ((b - a).days + 1)


def site_snapshots(project, site_id, as_of=None):
    as_of = as_date(as_of or china_today())
    return sorted(
        (
            r
            for r in project.get("enrollment", [])
            if r["site_id"] == site_id and as_date(r["date"]) <= as_of
        ),
        key=lambda r: r["date"],
    )


def site_enrollment_summary(project, as_of=None):
    as_of = as_date(as_of or china_today())
    result = []
    for site in project.get("sites", []):
        snapshots = site_snapshots(project, site["id"], as_of)
        last = snapshots[-1] if snapshots else None
        actual = last["cumulative"] if last else None
        status = site_readiness(site)
        result.append(
            dict(
                site_id=site["id"],
                中心编号=site["code"],
                研究中心=site["name"],
                PI=site["pi"],
                PI联系方式=site["contact"],
                牵头中心="是" if site.get("is_lead", False) else "否",
                推进阶段=SITE_STAGE_LABELS[site["stage"]],
                检查已完成=status["done"],
                检查不适用=status["na"],
                检查待确认=status["pending"],
                计划入组=site["planned_n"],
                实际累计=actual,
                最新报告日期=last["date"] if last else None,
                同日计划=(
                    round(planned_cumulative(site, last["date"]), 2) if last else None
                ),
                同日偏差=(
                    round(actual - planned_cumulative(site, last["date"]), 2)
                    if last
                    else None
                ),
                目标完成率=(
                    round(actual / site["planned_n"] * 100, 1)
                    if last and site["planned_n"]
                    else None
                ),
                计划开始=site["plan_start"],
                计划结束=site["plan_end"],
                FPI已报告="是" if last and actual > 0 else "否",
            )
        )
    return result


def upsert_enrollment(project, site_id, date, cumulative, note=""):
    if (
        isinstance(cumulative, bool)
        or not isinstance(cumulative, (int, float))
        or not math.isfinite(cumulative)
        or cumulative < 0
        or cumulative != int(cumulative)
    ):
        raise ValueError("累计入组必须是非负整数，不能自动截断小数。")
    candidate = copy.deepcopy(project)
    day = as_date(date).isoformat()
    row = next(
        (
            r
            for r in candidate["enrollment"]
            if r["site_id"] == site_id and r["date"] == day
        ),
        None,
    )
    if row:
        row.update(cumulative=int(cumulative), note=note)
    else:
        candidate["enrollment"].append(
            dict(
                id="E-S" + uuid.uuid4().hex[:10],
                site_id=site_id,
                date=day,
                cumulative=int(cumulative),
                note=note,
            )
        )
    return candidate


def distribute_site_targets(project):
    candidate = copy.deepcopy(project)
    sites = candidate["sites"]
    if not sites:
        raise ValueError("请先建立研究中心清单。")
    n = candidate["settings"]["n_planned"]
    base, extra = divmod(int(n), len(sites))
    for i, site in enumerate(sites):
        site["planned_n"] = base + int(i < extra)
    return candidate


def enrollment_chart_rows(project, site_ids=None, as_of=None, aggregate=True):
    """Plan: linear daily accrual. Actual: stepwise dated reports, no future extrapolation.

    A total actual trace is limited to the common reporting interval of all selected sites;
    unreported sites are never silently treated as zero. A reported first zero can establish
    earlier zeros because validated cumulative enrollment cannot decrease.
    """
    from bisect import bisect_right

    as_of = as_date(as_of or china_today())
    sites = [
        s for s in project.get("sites", []) if site_ids is None or s["id"] in site_ids
    ]
    if not sites:
        return []
    first = min(
        min(as_date(s["plan_start"]) for s in sites) - dt.timedelta(days=1), as_of
    )
    last = max(max(as_date(s["plan_end"]) for s in sites), as_of)
    snapshots = {s["id"]: site_snapshots(project, s["id"], as_of) for s in sites}
    for records in snapshots.values():
        if records:
            first = min(first, as_date(records[0]["date"]))
    dates = {first, last, as_of}
    for m in month_starts(first, last):
        dates.add(m)
        dates.add(min(last, add_months(m, 1) - dt.timedelta(days=1)))
    calendar_dates = {day for day in dates if first <= day <= last}
    for s in sites:
        dates.update(
            [
                as_date(s["plan_start"]) - dt.timedelta(days=1),
                as_date(s["plan_start"]),
                as_date(s["plan_end"]),
            ]
        )
    for rows in snapshots.values():
        dates.update(as_date(r["date"]) for r in rows)
    dates = {day for day in dates if first <= day <= last}
    groups = (
        [("ALL", "全研究合计", sites)]
        if aggregate
        else [(s["id"], f"{s['code']} · {s['name']}", [s]) for s in sites]
    )
    output = []
    for group_id, label, members in groups:
        group_dates = (
            dates
            if aggregate
            else calendar_dates
            | {
                as_date(r["date"])
                for member in members
                for r in snapshots[member["id"]]
            }
            | {
                day
                for member in members
                for day in (
                    as_date(member["plan_start"]) - dt.timedelta(days=1),
                    as_date(member["plan_start"]),
                    as_date(member["plan_end"]),
                )
            }
        )
        for day in sorted(group_dates):
            output.append(
                dict(
                    series_id=group_id,
                    name=label,
                    date=day.isoformat(),
                    kind="计划入组",
                    cumulative=sum(planned_cumulative(s, day) for s in members),
                )
            )
        if any(not snapshots[s["id"]] for s in members):
            continue
        lower = max(
            (
                first
                if snapshots[s["id"]][0]["cumulative"] == 0
                else as_date(snapshots[s["id"]][0]["date"])
            )
            for s in members
        )
        upper = min(as_date(snapshots[s["id"]][-1]["date"]) for s in members)
        if upper < lower:
            continue
        actual_dates = sorted(
            {lower, upper} | {d for d in dates if lower <= d <= upper}
        )
        indexed = {
            s["id"]: (
                [as_date(r["date"]) for r in snapshots[s["id"]]],
                [r["cumulative"] for r in snapshots[s["id"]]],
            )
            for s in members
        }
        for day in actual_dates:
            actual = 0
            for s in members:
                ds, ns = indexed[s["id"]]
                i = bisect_right(ds, day) - 1
                actual += ns[i] if i >= 0 else 0
            output.append(
                dict(
                    series_id=group_id,
                    name=label,
                    date=day.isoformat(),
                    kind="实际入组",
                    cumulative=int(actual),
                )
            )
    return output


def make_enrollment_figure(rows, title="入组进度", project_name=""):
    import plotly.graph_objects as go

    fig = go.Figure()
    labels = list(dict.fromkeys((r["series_id"], r["name"]) for r in rows))
    for i, (group_id, label) in enumerate(labels):
        color = COLORS[i % len(COLORS)]
        for kind, dash, shape in [
            ("计划入组", "dash", "linear"),
            ("实际入组", "solid", "hv"),
        ]:
            points = sorted(
                (r for r in rows if r["series_id"] == group_id and r["kind"] == kind),
                key=lambda r: r["date"],
            )
            if not points:
                continue
            fig.add_trace(
                go.Scatter(
                    x=[r["date"] for r in points],
                    y=[r["cumulative"] for r in points],
                    mode="lines+markers" if kind == "实际入组" else "lines",
                    name=html.escape(label) + " · " + kind,
                    legendgroup=group_id,
                    line=dict(color=color, width=2.5, dash=dash, shape=shape),
                    marker=dict(size=4),
                    hovertemplate="%{x|%Y-%m-%d}<br>累计入组：%{y:.2f}例<extra>%{fullData.name}</extra>",
                )
            )
    chart_style(fig, 520)
    fig.update_layout(
        title=html.escape(title),
        xaxis_title="日期",
        yaxis_title="累计入组（例）",
        hovermode="x unified",
        legend=dict(orientation="h", y=-0.23),
        margin=dict(b=110),
        title_font_size=18,
    )
    fig.update_yaxes(rangemode="tozero", gridcolor="#E5ECF5")
    fig.update_xaxes(type="date", tickformat="%Y-%m", gridcolor="#EDF1F7")
    return fig


def enrollment_svg_bytes(rows, title="入组进度"):
    """Standalone SVG of exactly the same series; no Chrome/Kaleido server dependency."""
    if not rows:
        raise ValueError("暂无曲线数据。")
    w, h = 1600, 900
    left, right, top, bottom = 100, 70, 100, 160
    labels = list(dict.fromkeys((r["series_id"], r["name"]) for r in rows))
    bottom = max(bottom, 70 + math.ceil(len(labels) / 3) * 28)
    plot_w = w - left - right
    plot_h = h - top - bottom
    dates = [as_date(r["date"]) for r in rows]
    first, last = min(dates), max(dates)
    days = max(1, (last - first).days)
    peak = max(1, max(r["cumulative"] for r in rows)) * 1.1
    xx = lambda date: left + (as_date(date) - first).days / days * plot_w
    yy = lambda n: top + plot_h - float(n) / peak * plot_h
    esc = lambda value: html.escape(str(value), quote=True)
    elements = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}">',
        '<rect width="100%" height="100%" fill="#FFFFFF"/>',
        '<g font-family="Microsoft YaHei, PingFang SC, Noto Sans CJK SC, Arial, sans-serif" fill="#163456">',
        f'<text x="{left}" y="46" font-size="27" font-weight="700">{esc(title)}</text>',
        f'<text x="{left}" y="78" font-size="16">计划：虚线  /  实际：实线（按报告日期阶梯显示）</text>',
    ]
    for i in range(6):
        value = peak * i / 5
        y = yy(value)
        elements.append(
            f'<line x1="{left}" y1="{y:.2f}" x2="{w-right}" y2="{y:.2f}" stroke="#E3EAF3"/>'
        )
        elements.append(
            f'<text x="{left-14}" y="{y+6:.2f}" text-anchor="end" font-size="16">{value:,.0f}</text>'
        )
    for i in range(7):
        day = first + dt.timedelta(days=round(days * i / 6))
        x = xx(day)
        elements.append(
            f'<text x="{x:.2f}" y="{top+plot_h+28}" text-anchor="middle" font-size="16">{day:%Y-%m-%d}</text>'
        )
    for i, (group_id, label) in enumerate(labels):
        color = COLORS[i % len(COLORS)]
        for kind in ["计划入组", "实际入组"]:
            points = sorted(
                (r for r in rows if r["series_id"] == group_id and r["kind"] == kind),
                key=lambda r: r["date"],
            )
            if not points:
                continue
            x0, y0 = xx(points[0]["date"]), yy(points[0]["cumulative"])
            path = f"M{x0:.2f},{y0:.2f}"
            for point in points[1:]:
                x, y = xx(point["date"]), yy(point["cumulative"])
                path += (
                    f"H{x:.2f}V{y:.2f}" if kind == "实际入组" else f"L{x:.2f},{y:.2f}"
                )
            dash = ' stroke-dasharray="10 7"' if kind == "计划入组" else ""
            elements.append(
                f'<path d="{path}" fill="none" stroke="{color}" stroke-width="3"{dash}/>'
            )
            if kind == "实际入组":
                for point in points:
                    elements.append(
                        f'<circle cx="{xx(point["date"]):.2f}" cy="{yy(point["cumulative"]):.2f}" r="3" fill="{color}"/>'
                    )
        x = left + (i % 3) * 470
        y = top + plot_h + 83 + (i // 3) * 28
        label = label if len(label) <= 28 else label[:27] + "…"
        elements.append(
            f'<line x1="{x}" y1="{y-6}" x2="{x+26}" y2="{y-6}" stroke="{color}" stroke-width="4"/>'
        )
        elements.append(f'<text x="{x+36}" y="{y}" font-size="16">{esc(label)}</text>')
    elements.extend(
        [
            f'<text x="{left}" y="{h-20}" font-size="14" fill="#6C7F96">实际曲线截止已报告日期；未报告不按0处理。计划按各中心起止日期均匀分配。</text>',
            "</g></svg>",
        ]
    )
    return "".join(elements).encode("utf-8")


def site_report_tables(project):
    sites = project.get("sites", [])
    lookup = {s["id"]: s for s in sites}
    readiness = [
        dict(
            中心编号=s["code"],
            研究中心=s["name"],
            类别=group,
            检查项=label,
            状态=s["readiness"][key],
        )
        for s in sites
        for group, checks_ in FPI_CHECKLIST.items()
        for key, label in checks_
    ]
    records = [
        dict(
            记录ID=r["id"],
            中心编号=lookup[r["site_id"]]["code"],
            研究中心=lookup[r["site_id"]]["name"],
            报告日期=r["date"],
            实际累计入组=r["cumulative"],
            备注=r["note"],
        )
        for r in sorted(
            project.get("enrollment", []), key=lambda r: (r["site_id"], r["date"])
        )
    ]
    return {
        "中心信息与入组汇总": site_enrollment_summary(project),
        "FPI启动检查": readiness,
        "入组累计快照": records,
        "入组曲线数据": enrollment_chart_rows(project, aggregate=False),
        "全中心月度入组": monthly_enrollment(project),
    }


def persist_site_stage(site_id, widget_key):
    """Slider updates are saved immediately; the readiness checklist stays independent."""
    candidate = copy.deepcopy(st.session_state.project)
    site = next((s for s in candidate["sites"] if s["id"] == site_id), None)
    if site is None:
        return
    site["stage"] = SITE_STAGE_LABELS.index(st.session_state[widget_key])
    st.session_state.project = candidate
    st.session_state.revision += 1


def render_sites(p):
    rev = form_epoch(
        "sites",
        [
            {
                k: s.get(k, False if k == "is_lead" else "")
                for k in ("id", "code", "name", "pi", "contact", "note", "is_lead")
            }
            for s in p.get("sites", [])
        ],
    )
    sites = p.get("sites", [])
    st.subheader("中心启动进度")
    st.caption(
        "先建立中心与PI清单；每个中心可拖动阶段指针，并独立维护截图中的5类16项FPI准备检查。"
    )
    a, b, c = st.columns(3)
    a.metric("已建中心 / 项目计划", f"{len(sites)} / {p['settings']['site_number']}")
    b.metric("检查均已确认的中心", sum(site_readiness(s)["ready"] for s in sites))
    c.metric(
        "已报告入组的中心",
        sum(
            bool(
                site_snapshots(p, s["id"])
                and site_snapshots(p, s["id"])[-1]["cumulative"] > 0
            )
            for s in sites
        ),
    )
    tab1, tab2 = stable_tabs(["中心与PI信息", "启动阶段与检查清单"], key="tabs_sites")
    with tab1:
        st.caption(
            "中心ID固定，重命名不会丢失入组历史。勾选唯一牵头中心后，概览自动关联该中心及PI；联系方式可填电话或邮箱。"
        )
        names = {
            "id": "中心ID",
            "is_lead": "牵头中心",
            "code": "中心编号",
            "name": "研究中心名称",
            "pi": "PI姓名",
            "contact": "PI联系方式",
            "note": "备注",
        }
        frame = pd.DataFrame(sites, columns=names).rename(columns=names)
        with st.form(f"sites_{rev}"):
            edited = editable_table(
                frame,
                hide_index=True,
                width="stretch",
                num_rows="dynamic",
                height=350,
                key=f"sites_grid_{rev}",
                column_config={
                    "中心ID": st.column_config.TextColumn(disabled=True),
                    "牵头中心": st.column_config.CheckboxColumn(
                        default=False,
                        help="仅勾选一个牵头中心；概览中的Leading PI自动关联该中心PI。",
                    ),
                    "中心编号": st.column_config.TextColumn(required=True),
                    "研究中心名称": st.column_config.TextColumn(
                        required=True, width="large"
                    ),
                    "PI姓名": st.column_config.TextColumn(),
                    "PI联系方式": st.column_config.TextColumn(width="medium"),
                    "备注": st.column_config.TextColumn(width="large"),
                },
            )
            submit = st.form_submit_button("保存中心信息", type="primary")
        if submit:
            candidate = copy.deepcopy(p)
            old = {s["id"]: s for s in sites}
            output = []
            for i, r in enumerate(ui_records(edited, {v: k for k, v in names.items()})):
                base = (
                    copy.deepcopy(old.get(r["id"]))
                    if r["id"] in old
                    else new_site(p, i + 1)
                )
                base["is_lead"] = bool(r.get("is_lead", False))
                for key in ["code", "name", "pi", "contact", "note"]:
                    value = r[key]
                    if key in ["pi", "contact", "note"] and value is None:
                        value = ""
                    base[key] = value.strip() if isinstance(value, str) else value
                output.append(base)
            candidate["sites"] = output
            apply_change(candidate, "中心与PI信息已保存，入组页面已同步。")
        if not sites:
            st.info(
                "可以在上表添加真实中心；也可先按项目中心数建立待填清单，再批量补充名称和PI。"
            )
            if st.button("按项目中心数生成待填清单", key="generate_site_roster"):
                if p["settings"]["site_number"] > 500:
                    st.error("中心追踪页面最多支持500个中心，请先调整项目中心数。")
                else:
                    candidate = copy.deepcopy(p)
                    candidate["sites"] = [
                        new_site(p, i + 1) for i in range(p["settings"]["site_number"])
                    ]
                    apply_change(
                        candidate,
                        "待填中心清单已生成；请补充真实名称，并到入组页面填写计划。",
                    )
        else:
            st.download_button(
                "下载中心与PI清单 CSV",
                csv_bytes(
                    site_enrollment_summary(p),
                    current_export_language(),
                    "中心信息与入组汇总",
                ),
                file_name=export_csv_filename("CRMS_sites"),
                mime="text/csv",
                on_click="ignore",
            )
    with tab2:
        if not sites:
            st.info("请先在“中心与PI信息”建立中心清单。")
            return
        st.caption(
            "阶段指针拖动后自动保存，表示当前工作推进位置；下方准备完成度由逐项检查计算，两者独立维护。"
        )
        query = st.text_input("查找中心或PI", key="site_search")
        shown = [
            s
            for s in sites
            if not query
            or query.casefold() in f"{s['code']} {s['name']} {s['pi']}".casefold()
        ]
        columns = st.columns(2, gap="medium")
        for i, site in enumerate(shown):
            with columns[i % 2]:
                with st.container(border=True):
                    st.write(f"**{site['code']} · {site['name']}**")
                    st.caption(
                        f"PI：{site['pi'] or '待填写'}　｜　联系方式：{site['contact'] or '待填写'}"
                    )
                    key = f"site_stage_{site['id']}_{st.session_state.get('ui_generation',0)}"
                    st.select_slider(
                        "当前推进环节",
                        options=SITE_STAGE_LABELS,
                        value=SITE_STAGE_LABELS[site["stage"]],
                        key=key,
                        on_change=persist_site_stage,
                        args=(site["id"], key),
                    )
                    ready = site_readiness(site)
                    st.progress(
                        ready["ratio"],
                        text=f"准备检查：{ready['done']}/ {16-ready['na']} 项已完成；{ready['na']}项不适用",
                    )
                    if ready["ready"]:
                        st.caption("16项均已确认（已完成或不适用）。")
                    else:
                        st.caption(f"仍有 {ready['pending']} 项待确认。")
        st.markdown("**FPI Readiness Checklist · 逐中心确认**")
        lookup = {s["id"]: s for s in sites}
        selected = st.selectbox(
            "选择中心维护16项检查",
            list(lookup),
            format_func=lambda k: f"{lookup[k]['code']} · {lookup[k]['name']}",
            key="readiness_site",
        )
        site = lookup[selected]
        rrev = form_epoch("readiness:" + selected, site["readiness"])
        rows = [
            dict(检查ID=key, 类别=group, 检查项目=label, 状态=site["readiness"][key])
            for group, items in FPI_CHECKLIST.items()
            for key, label in items
        ]
        with st.form(f"readiness_{selected}_{rrev}"):
            edited = editable_table(
                pd.DataFrame(rows),
                hide_index=True,
                width="stretch",
                height=610,
                key=f"readiness_grid_{selected}_{rrev}",
                column_config={
                    "检查ID": st.column_config.TextColumn(disabled=True),
                    "类别": st.column_config.TextColumn(disabled=True),
                    "检查项目": st.column_config.TextColumn(
                        disabled=True, width="large"
                    ),
                    "状态": st.column_config.SelectboxColumn(
                        options=READINESS_STATUSES, required=True
                    ),
                },
            )
            save = st.form_submit_button("保存FPI准备检查", type="primary")
        if save:
            candidate = copy.deepcopy(p)
            updated = next(s for s in candidate["sites"] if s["id"] == selected)
            updated["readiness"] = {
                r["检查ID"]: r["状态"] for _, r in edited.iterrows()
            }
            apply_change(candidate, "本中心FPI准备检查已保存。")
        st.caption(
            "截图检查项可以并行开展；不适用项可明确标记。阶段指针不会自动勾选检查，也不会自动生成实际入组记录。"
        )


def render_enrollment(p):
    sites = p.get("sites", [])
    rev = form_epoch(
        "site_plan",
        [
            {
                k: s[k]
                for k in (
                    "id",
                    "code",
                    "name",
                    "pi",
                    "planned_n",
                    "plan_start",
                    "plan_end",
                )
            }
            for s in p.get("sites", [])
        ],
    )
    st.subheader("入组进度")
    if not sites:
        st.info("请先到“中心启动进度”建立中心清单，本页会自动使用相同的中心名称和PI。")
        return
    summary = site_enrollment_summary(p)
    reported = [r for r in summary if r["实际累计"] is not None]
    total_plan = sum(s["planned_n"] for s in sites)
    total_actual = sum(r["实际累计"] for r in reported)
    cols = st.columns(4)
    cols[0].metric("分中心计划合计", f"{total_plan:,} 例")
    cols[1].metric("已报告实际累计", f"{total_actual:,} 例")
    cols[2].metric("已报告 / 全部中心", f"{len(reported)} / {len(sites)}")
    cols[3].metric(
        "目标完成率",
        (
            f"{total_actual/total_plan*100:.1f}%"
            if total_plan and len(reported) == len(sites)
            else "待补计划/报告"
        ),
    )
    st.caption(
        "已报告实际合计取各中心最近快照，报告日期可能不同；同日比较请查看曲线的共同报告区间。"
    )
    if len(reported) < len(sites):
        st.info(
            f"有{len(sites)-len(reported)}个中心尚无实际报告；已报告合计不代表全部中心的完整实际数。无入组也请明确填写0例。"
        )
    if (
        total_plan != p["settings"]["n_planned"]
        or len(sites) != p["settings"]["site_number"]
    ):
        st.caption(
            f"项目设置：{p['settings']['site_number']}中心、{p['settings']['n_planned']}例；本页：{len(sites)}中心、{total_plan}例。保存分中心计划不会自动改写预算。"
        )
    tab1, tab2, tab3 = stable_tabs(
        ["分中心入组计划", "实际入组填报", "曲线与导出"], key="tabs_enrollment"
    )
    with tab1:
        st.caption(
            "计划曲线按每个中心的计划开始/结束日期均匀分配，累计至目标例数；可为不同中心设置不同的启动日期。"
        )
        names = {
            "id": "中心ID",
            "code": "中心编号",
            "name": "研究中心",
            "pi": "PI",
            "planned_n": "计划入组(例)",
            "plan_start": "计划首例日期",
            "plan_end": "计划末例日期",
        }
        frame = pd.DataFrame(sites, columns=names).rename(columns=names)
        for key in ["计划首例日期", "计划末例日期"]:
            frame[key] = pd.to_datetime(frame[key]).dt.date
        with st.form(f"site_plan_{rev}"):
            edited = editable_table(
                frame,
                hide_index=True,
                width="stretch",
                height=380,
                key=f"site_plan_grid_{rev}",
                column_config={
                    **{
                        key: st.column_config.TextColumn(disabled=True)
                        for key in ["中心ID", "中心编号", "研究中心", "PI"]
                    },
                    "计划入组(例)": st.column_config.NumberColumn(
                        min_value=0, max_value=100000, step=1, required=True
                    ),
                    "计划首例日期": st.column_config.DateColumn(
                        format="YYYY-MM-DD", required=True
                    ),
                    "计划末例日期": st.column_config.DateColumn(
                        format="YYYY-MM-DD", required=True
                    ),
                },
            )
            submit = st.form_submit_button("保存分中心入组计划", type="primary")
        if submit:
            candidate = copy.deepcopy(p)
            lookup = {s["id"]: s for s in candidate["sites"]}
            for r in ui_records(edited, {v: k for k, v in names.items()}):
                lookup[r["id"]].update(
                    planned_n=r["planned_n"],
                    plan_start=r["plan_start"],
                    plan_end=r["plan_end"],
                )
            apply_change(candidate, "分中心计划已保存，计划曲线已同步更新。")
        a, b = st.columns(2)
        if a.button(
            "将项目计划例数均分到当前中心", key="split_targets", width="stretch"
        ):
            apply_change(
                distribute_site_targets(p),
                "项目计划已按整数均分到各中心；可以继续手工调整。",
            )
        if b.button(
            "用分中心计划同步项目人数与中心数", key="sync_site_targets", width="stretch"
        ):
            if total_plan <= 0:
                st.error("请先填写至少1例的分中心计划。")
            else:
                candidate = copy.deepcopy(p)
                candidate["settings"].update(
                    n_planned=int(total_plan), site_number=len(sites)
                )
                if candidate["timeline_mode"] != "手工维护":
                    try:
                        candidate["timeline"] = make_timeline(candidate)
                    except (ValueError, OverflowError) as exc:
                        st.error(str(exc))
                        return
                apply_change(
                    candidate,
                    "项目人数、中心数及相关预算已同步；分中心计划日期保持填写值。",
                )
        st.caption(
            "均分会覆盖当前各中心目标；同步至项目设置会重算按人数/中心计费的预算，并按当前维护方式更新项目排期。"
        )
        readonly_table(
            pd.DataFrame(summary).drop(columns=["site_id", "PI联系方式"]),
            hide_index=True,
            width="stretch",
            column_config={
                "目标完成率": st.column_config.ProgressColumn(
                    min_value=0, max_value=100, format="%.1f%%"
                ),
                "实际累计": st.column_config.NumberColumn(
                    format="%d", help="空白表示尚无实际报告，不是0。"
                ),
            },
        )
    with tab2:
        st.caption(
            "填写“截至某日累计已入组”，不是当天新增人数。空白表示未填报，0表示已确认尚无入组。"
        )
        day = st.date_input(
            "填报日期（截至当天累计）",
            value=china_today(),
            min_value=dt.date(2000, 1, 1),
            max_value=china_today(),
            key="enrollment_report_date",
        )
        rows = []
        for site in sites:
            existing = next(
                (
                    r
                    for r in p["enrollment"]
                    if r["site_id"] == site["id"] and r["date"] == day.isoformat()
                ),
                None,
            )
            prior = site_snapshots(p, site["id"], day)
            rows.append(
                dict(
                    中心ID=site["id"],
                    中心编号=site["code"],
                    研究中心=site["name"],
                    PI=site["pi"],
                    最近已报累计=prior[-1]["cumulative"] if prior else None,
                    本日累计入组=existing["cumulative"] if existing else None,
                    备注=existing["note"] if existing else "",
                )
            )
        frame = pd.DataFrame(rows)
        # Explicit numeric dtypes keep an initially all-empty editable column writable.
        frame["本日累计入组"] = pd.array(frame["本日累计入组"], dtype="Int64")
        drev = form_epoch("enrollment_daily:" + day.isoformat(), rows)
        with st.form(f"enrollment_daily_{day}_{drev}"):
            edited = editable_table(
                frame,
                hide_index=True,
                width="stretch",
                height=370,
                key=f"enrollment_daily_grid_{day}_{drev}",
                column_config={
                    **{
                        key: st.column_config.TextColumn(disabled=True)
                        for key in ["中心ID", "中心编号", "研究中心", "PI"]
                    },
                    "最近已报累计": st.column_config.NumberColumn(
                        disabled=True, format="%d"
                    ),
                    "本日累计入组": st.column_config.NumberColumn(
                        min_value=0, max_value=1000000, step=1, format="%d"
                    ),
                    "备注": st.column_config.TextColumn(),
                },
            )
            submit = st.form_submit_button("保存本日入组快照", type="primary")
        if submit:
            candidate = copy.deepcopy(p)
            count = 0
            for r in ui_records(
                edited,
                {"中心ID": "site_id", "本日累计入组": "cumulative", "备注": "note"},
            ):
                if r["cumulative"] is None:
                    continue
                try:
                    candidate = upsert_enrollment(
                        candidate, r["site_id"], day, r["cumulative"], r["note"] or ""
                    )
                except (ValueError, TypeError) as exc:
                    st.error(str(exc))
                    return
                count += 1
            if count:
                apply_change(
                    candidate,
                    f"已保存{count}个中心截至{day}的入组快照；同日记录已更新。",
                )
            else:
                st.info("尚未填写任何累计例数，未新增记录。")
        with stable_expander("历史快照：更正、补录或删除"):
            lookup = {s["id"]: s for s in sites}
            codes = {s["code"]: s["id"] for s in sites}
            rows = [
                dict(
                    记录ID=r["id"],
                    中心编号=lookup[r["site_id"]]["code"],
                    报告日期=as_date(r["date"]),
                    累计入组=r["cumulative"],
                    备注=r["note"],
                )
                for r in sorted(
                    p["enrollment"], key=lambda r: (r["site_id"], r["date"])
                )
            ]
            frame = pd.DataFrame(
                rows, columns=["记录ID", "中心编号", "报告日期", "累计入组", "备注"]
            )
            frame["报告日期"] = pd.to_datetime(frame["报告日期"]).dt.date
            frame["累计入组"] = pd.array(frame["累计入组"], dtype="Int64")
            hrev = form_epoch("enrollment_history", rows)
            with st.form(f"enrollment_history_{hrev}"):
                edited = editable_table(
                    frame,
                    hide_index=True,
                    width="stretch",
                    num_rows="dynamic",
                    height=320,
                    key=f"enrollment_history_grid_{hrev}",
                    column_config={
                        "记录ID": st.column_config.TextColumn(disabled=True),
                        "中心编号": st.column_config.SelectboxColumn(
                            options=list(codes), required=True
                        ),
                        "报告日期": st.column_config.DateColumn(
                            format="YYYY-MM-DD", required=True
                        ),
                        "累计入组": st.column_config.NumberColumn(
                            min_value=0, max_value=1000000, required=True, step=1
                        ),
                        "备注": st.column_config.TextColumn(),
                    },
                )
                submit = st.form_submit_button("保存历史入组记录")
            if submit:
                candidate = copy.deepcopy(p)
                output = []
                for r in ui_records(
                    edited,
                    {
                        "记录ID": "id",
                        "中心编号": "code",
                        "报告日期": "date",
                        "累计入组": "cumulative",
                        "备注": "note",
                    },
                ):
                    output.append(
                        dict(
                            id=r["id"] or "E-S" + uuid.uuid4().hex[:10],
                            site_id=codes.get(r["code"], ""),
                            date=r["date"],
                            cumulative=r["cumulative"],
                            note=r["note"] or "",
                        )
                    )
                candidate["enrollment"] = output
                apply_change(candidate, "历史入组快照已保存。")
            st.caption(
                "同一中心同日不能重复；累计入组应随时间不减少。更正误报后，请检查后续累计数是否仍一致。"
            )
    with tab3:
        render_enrollment_curves(p)


def render_enrollment_curves(p):
    sites = p["sites"]
    lookup = {s["id"]: s for s in sites}
    mode = st.radio(
        "曲线视图",
        ["全研究合计", "分中心对比", "单中心"],
        horizontal=True,
        key="enrollment_view",
    )
    if mode == "单中心":
        selected = [
            st.selectbox(
                "选择入组曲线中心",
                list(lookup),
                format_func=lambda k: f"{lookup[k]['code']} · {lookup[k]['name']}",
                key="enrollment_single_site",
            )
        ]
    else:
        selected = st.multiselect(
            "纳入曲线的中心",
            list(lookup),
            default=list(lookup),
            format_func=lambda k: f"{lookup[k]['code']} · {lookup[k]['name']}",
            key="enrollment_curve_sites",
        )
    if not selected:
        st.info("请选择至少一个中心。")
        return
    as_of = st.date_input(
        "实际数据截至日期",
        value=china_today(),
        min_value=dt.date(2000, 1, 1),
        max_value=china_today(),
        key="enrollment_curve_asof",
    )
    rows = enrollment_chart_rows(p, selected, as_of, aggregate=mode == "全研究合计")
    title = f"{p['settings']['project_name']} · {mode}入组进度"
    fig = make_enrollment_figure(rows, title)
    st.plotly_chart(
        fig,
        width="stretch",
        key="enrollment_progress_chart",
        theme=None,
        config={
            "displaylogo": False,
            "toImageButtonOptions": {
                "format": "png",
                "filename": "CRMS_enrollment_progress",
                "width": 1600,
                "height": 900,
                "scale": 2,
            },
        },
    )
    st.caption(
        "虚线＝均匀分配的累计计划；实线＝累计实际报告，按报告日期阶梯连接。实际曲线停在已报告日期，不延伸至未来。"
    )
    selected_sites = [s for s in sites if s["id"] in selected]
    missing = [
        s["name"] for s in selected_sites if not site_snapshots(p, s["id"], as_of)
    ]
    if missing:
        st.info(
            "尚无实际报告的中心："
            + "、".join(missing)
            + "。未报告不按0补齐；全研究合计实际线需所有所选中心都有报告。"
        )
    elif mode == "全研究合计":
        actual = [r for r in rows if r["kind"] == "实际入组"]
        if actual:
            st.caption(
                f"合计实际线采用共同报告区间：{actual[0]['date']} 至 {actual[-1]['date']}；区间内沿用各中心最近已报累计数。"
            )
        else:
            st.info(
                "各中心报告日期尚无共同覆盖区间，暂不绘制合计实际线。可增加同期报告或查看分中心曲线。"
            )
    over = [
        s
        for s in selected_sites
        if site_snapshots(p, s["id"], as_of)
        and site_snapshots(p, s["id"], as_of)[-1]["cumulative"] > s["planned_n"]
    ]
    if over:
        st.warning(
            "以下中心实际已超过当前目标（数据仍保留）："
            + "、".join(s["name"] for s in over)
        )
    cols = st.columns(3)
    cols[0].download_button(
        "下载曲线 SVG",
        enrollment_svg_bytes(rows, title),
        file_name="CRMS_enrollment_progress.svg",
        mime="image/svg+xml",
        width="stretch",
        on_click="ignore",
    )
    cols[1].download_button(
        "下载交互曲线 HTML",
        fig.to_html(full_html=True, include_plotlyjs=True).encode("utf-8"),
        file_name="CRMS_enrollment_progress.html",
        mime="text/html",
        width="stretch",
        on_click="ignore",
    )
    cols[2].download_button(
        "下载曲线数据 CSV",
        csv_bytes(rows, current_export_language(), "入组曲线数据"),
        file_name=export_csv_filename("CRMS_enrollment_progress"),
        mime="text/csv",
        width="stretch",
        on_click="ignore",
    )
    st.caption(
        "PNG图片可点击图表右上角相机导出；SVG是可放入PPT的矢量图；HTML可离线打开交互。"
    )


# ---------- V2.3: leadership, monthly totals and stable UI identity ----------
def leading_site(project):
    return next(
        (site for site in project.get("sites", []) if site.get("is_lead", False)), None
    )


def remember_navigation():
    st.session_state.active_page = st.session_state.navigation


def form_epoch(scope, content):
    """Unrelated saves keep the same form/widget identity and buffered edits."""
    signature = json.dumps(
        [st.session_state.get("ui_generation", 0), content],
        ensure_ascii=False,
        sort_keys=True,
        default=str,
    )
    epochs = st.session_state.setdefault("form_epochs", {})
    old = epochs.get(scope)
    if old is None:
        epochs[scope] = dict(signature=signature, version=0)
    elif old["signature"] != signature:
        epochs[scope] = dict(signature=signature, version=old["version"] + 1)
    return epochs[scope]["version"]


def stable_expander(label, **kwargs):
    """Use stable keys when supported, and the native older API otherwise."""
    if "key" in inspect.signature(st.expander).parameters:
        kwargs.setdefault(
            "key", "exp_" + uuid.uuid5(uuid.NAMESPACE_URL, "crms:" + label).hex[:16]
        )
    else:
        kwargs.pop("key", None)
    return st.expander(label, **kwargs)


def stable_tabs(labels, *, key):
    # Streamlit 1.50 accepts no tab key; newer versions use it to retain identity.
    if "key" in inspect.signature(st.tabs).parameters:
        return st.tabs(labels, key=key)
    return st.tabs(labels)


def monthly_enrollment(project, site_ids=None, as_of=None):
    """One all-site cumulative observation per calendar month.

    Past months use month end; the current month uses the selected cutoff.
    Actual values use each site's latest dated report by that cutoff, never
    interpolate enrollment or assume an unreported center has zero participants.
    No actual points are extended beyond the latest month with any report.
    """
    from bisect import bisect_right

    as_of = as_date(as_of or china_today())
    sites = [
        s for s in project.get("sites", []) if site_ids is None or s["id"] in site_ids
    ]
    if not sites:
        return []
    snapshots = {s["id"]: site_snapshots(project, s["id"], as_of) for s in sites}
    first = min(as_date(s["plan_start"]) for s in sites)
    reported_dates = [
        as_date(r["date"]) for records in snapshots.values() for r in records
    ]
    if reported_dates:
        first = min(first, min(reported_dates))
    last = max(as_of, max(as_date(s["plan_end"]) for s in sites))
    last_actual_month = max(reported_dates).replace(day=1) if reported_dates else None
    indexed = {
        sid: ([as_date(r["date"]) for r in records], [r["cumulative"] for r in records])
        for sid, records in snapshots.items()
    }
    output = []
    for month in month_starts(first, last):
        month_end = add_months(month, 1) - dt.timedelta(days=1)
        current = month.year == as_of.year and month.month == as_of.month
        cutoff = as_of if current else month_end
        plan = sum(planned_cumulative(site, cutoff) for site in sites)
        actual = 0
        known = 0
        fresh = 0
        eligible = (
            month <= as_of.replace(day=1)
            and last_actual_month is not None
            and month <= last_actual_month
        )
        for site in sites:
            dates, counts = indexed[site["id"]]
            i = bisect_right(dates, cutoff) - 1
            if i >= 0:
                known += 1
                actual += counts[i]
                fresh += int(
                    dates[i].year == month.year and dates[i].month == month.month
                )
            elif counts and counts[0] == 0:
                # A declared first zero and nondecreasing totals establish earlier zero.
                known += 1
        output.append(
            dict(
                月份=month.strftime("%Y-%m"),
                统计截至=cutoff.isoformat(),
                计划累计=round(plan, 2),
                实际累计=int(actual) if eligible and known == len(sites) else None,
                纳入中心数=len(sites),
                已有累计报告中心数=known if eligible else 0,
                本月更新中心数=fresh if eligible else 0,
            )
        )
    return output


def monthly_curve_rows(monthly):
    return [
        dict(
            series_id="ALL",
            name="全中心合计",
            date=row["月份"] + "-01",
            kind=kind,
            cumulative=row[field],
        )
        for row in monthly
        for kind, field in [("计划入组", "计划累计"), ("实际入组", "实际累计")]
    ]


def make_monthly_enrollment_figure(monthly, title="全中心月度累计入组"):
    import plotly.graph_objects as go

    fig = go.Figure()
    for field, label, dash, color in [
        ("计划累计", "计划入组（虚线）", "dash", COLORS[0]),
        ("实际累计", "实际入组（实线）", "solid", COLORS[1]),
    ]:
        fig.add_trace(
            go.Scatter(
                x=[r["月份"] + "-01" for r in monthly],
                y=[r[field] for r in monthly],
                mode="lines+markers",
                name=label,
                connectgaps=False,
                line=dict(color=color, width=3, dash=dash, shape="linear"),
                marker=dict(size=7),
                customdata=[
                    [r["统计截至"], r["本月更新中心数"], r["纳入中心数"]]
                    for r in monthly
                ],
                hovertemplate="%{x|%Y-%m}<br>累计入组：%{y:.2f}例<br>统计截至：%{customdata[0]}<br>本月更新中心：%{customdata[1]}/%{customdata[2]}<extra>%{fullData.name}</extra>",
            )
        )
    chart_style(fig, 430)
    fig.update_layout(
        title=html.escape(title),
        hovermode="x unified",
        xaxis_title="月份",
        yaxis_title="累计入组（例）",
        legend=dict(orientation="h", y=-0.28),
        margin=dict(b=100),
    )
    fig.update_xaxes(
        type="date",
        tickformat="%Y-%m",
        dtick="M" + str(max(1, math.ceil(len(monthly) / 24))),
        tickangle=-45,
        gridcolor="#E5ECF5",
    )
    fig.update_yaxes(rangemode="tozero", gridcolor="#E5ECF5")
    return fig


def monthly_enrollment_svg(monthly, title="全中心月度累计入组"):
    """Monthly standalone vector chart, with gaps retained and straight connections."""
    w, h = 1600, 850
    left, right, top, bottom = 100, 60, 105, 130
    pw, ph = w - left - right, h - top - bottom
    n = len(monthly)
    peak = (
        max(1, max((r[k] or 0) for r in monthly for k in ("计划累计", "实际累计")))
        * 1.12
    )
    x = lambda i: left + (i / (n - 1) if n > 1 else 0.5) * pw
    y = lambda v: top + ph - float(v) / peak * ph
    esc = lambda v: html.escape(str(v), quote=True)
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}"><rect width="100%" height="100%" fill="#FFFFFF"/>',
        '<g fill="#163456" font-family="Microsoft YaHei, PingFang SC, Arial, sans-serif">',
        f'<text x="{left}" y="48" font-size="28" font-weight="700">{esc(title)}</text>',
        f'<text x="{left}" y="80" font-size="17">计划：蓝色虚线  /  实际：青色实线  /  每月一个累计点，直线连接</text>',
    ]
    for i in range(6):
        value = peak * i / 5
        yy = y(value)
        parts.extend(
            [
                f'<line x1="{left}" y1="{yy:.2f}" x2="{w-right}" y2="{yy:.2f}" stroke="#E3EAF3"/>',
                f'<text x="{left-14}" y="{yy+5:.2f}" text-anchor="end" font-size="16">{value:,.0f}</text>',
            ]
        )
    step = max(1, math.ceil(n / 18))
    for i, row in enumerate(monthly):
        if i % step == 0 or i == n - 1:
            parts.append(
                f'<text x="{x(i):.2f}" y="{top+ph+35}" text-anchor="middle" font-size="15">{row["月份"]}</text>'
            )
    for field, color, dash in [
        ("计划累计", COLORS[0], 'stroke-dasharray="10 7"'),
        ("实际累计", COLORS[1], ""),
    ]:
        path = ""
        connected = False
        for i, row in enumerate(monthly):
            value = row[field]
            if value is None:
                connected = False
                continue
            path += ("L" if connected else "M") + f"{x(i):.2f},{y(value):.2f}"
            connected = True
            parts.append(
                f'<circle cx="{x(i):.2f}" cy="{y(value):.2f}" r="5" fill="{color}"/>'
            )
        parts.append(
            f'<path d="{path}" fill="none" stroke="{color}" stroke-width="3" {dash}/>'
        )
    parts.extend(
        [
            f'<text x="{left}" y="{h-45}" font-size="16">历史月份取月末，当月取统计截至日。实际汇总各中心最近已报累计；缺报不按0补齐。</text>',
            f'<text x="{left}" y="{h-18}" font-size="14" fill="#63758B">实际点不延伸到未来或全部中心都尚未更新的后续月份。</text>',
            "</g></svg>",
        ]
    )
    return "".join(parts).encode("utf-8")


def render_overview_enrollment(p):
    st.subheader("全中心月度入组进度")
    if not p.get("sites"):
        st.info(
            "请先在中心启动进度中建立中心清单，并在入组进度中填写各中心计划及实际累计；概览将在此汇总所有中心。"
        )
        return
    monthly = monthly_enrollment(p)
    fig = make_monthly_enrollment_figure(monthly)
    st.plotly_chart(
        fig,
        width="stretch",
        key="overview_monthly_enrollment",
        theme=None,
        config={
            "displaylogo": False,
            "toImageButtonOptions": {
                "format": "png",
                "filename": "CRMS_monthly_enrollment",
                "width": 1600,
                "height": 850,
                "scale": 2,
            },
        },
    )
    st.caption(
        "汇总所有已登记中心，每月一个累计点：计划为虚线，实际为实线，均以直线连接。历史月份取月末，当月取统计截至日。"
    )
    st.caption(
        "实际取各中心截至统计日最近一条累计报告；本月未更新的中心沿用此前已报数，未报告不按0补齐，未来月份仅显示计划。"
    )
    current = next(
        (r for r in monthly if r["月份"] == china_today().strftime("%Y-%m")), None
    )
    if (
        current
        and current["实际累计"] is not None
        and current["本月更新中心数"] < current["纳入中心数"]
    ):
        st.info(
            f"本月{current['本月更新中心数']}/{current['纳入中心数']}个中心已更新，其余沿用最近已报累计；请及时补充报告。"
        )
    if not any(r["实际累计"] is not None for r in monthly):
        st.info(
            "各中心同期实际报告尚不完整，暂不绘制合计实际线；无入组的中心也请明确填报0例。"
        )
    total = sum(s["planned_n"] for s in p["sites"])
    if (
        total != p["settings"]["n_planned"]
        or len(p["sites"]) != p["settings"]["site_number"]
    ):
        st.caption(
            f"当前纳入{len(p['sites'])}家已登记中心，分中心目标合计{total}例；项目设置为{p['settings']['site_number']}中心、{p['settings']['n_planned']}例。"
        )
    a, b = st.columns(2)
    a.download_button(
        "下载月度入组曲线 SVG",
        monthly_enrollment_svg(monthly, p["settings"]["project_name"] + " · 月度入组"),
        file_name="CRMS_monthly_enrollment.svg",
        mime="image/svg+xml",
        on_click="ignore",
        width="stretch",
    )
    b.download_button(
        "下载月度入组数据 CSV",
        csv_bytes(monthly, current_export_language(), "全中心月度入组"),
        file_name=export_csv_filename("CRMS_monthly_enrollment"),
        mime="text/csv",
        on_click="ignore",
        width="stretch",
    )


if __name__ == "__main__":
    main()
